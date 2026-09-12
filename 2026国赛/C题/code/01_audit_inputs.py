"""Read-only source audit. Run in conda environment math_modeling."""
from pathlib import Path
from datetime import datetime, time
import hashlib
import json
import sys

import numpy as np
import pandas as pd
import openpyxl
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'results' / 'audit'
OUT.mkdir(parents=True, exist_ok=True)


def rows(name, sheet=0):
    wb = openpyxl.load_workbook(ROOT / '附件' / name, data_only=True, read_only=True)
    ws = wb.worksheets[sheet]
    data = list(ws.values)
    wb.close()
    return data


def minutes(value):
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    s = str(value)
    h, m = s.replace('+1', '').split(':')[:2]
    return int(h) * 60 + int(m) + (1440 if '+1' in s else 0)


def stats(a):
    a = np.asarray(a, dtype=float)
    finite = a[np.isfinite(a)]
    return dict(shape=list(a.shape), count=int(a.size), missing=int(np.isnan(a).sum()),
                nonfinite=int((~np.isfinite(a)).sum()), negative=int((a < 0).sum()),
                zero=int((a == 0).sum()), minimum=float(finite.min()),
                maximum=float(finite.max()), mean=float(finite.mean()))


def annual(name, sheet=0):
    data = rows(name, sheet)
    ticks = [minutes(v) for v in data[0][1:]]
    assert ticks == list(range(10, 1441, 10)), (name, 'unexpected time grid')
    dates = pd.DatetimeIndex([r[0] for r in data[1:]])
    assert dates.equals(pd.date_range('2025-01-01', '2025-12-31'))
    values = np.array([r[1:] for r in data[1:]], dtype=float)
    assert values.shape == (365, 144)
    return dates, ticks, values


def main():
    assert Path(sys.prefix).name == 'math_modeling', sys.prefix
    sources = [ROOT / 'C题.pdf'] + sorted(p for p in (ROOT / '附件').rglob('*.xlsx')
                                               if not p.name.startswith('~$'))
    checksums = {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in sources}
    reader = PdfReader(ROOT / 'C题.pdf')
    extracted = '\n\n'.join(f'## 第{i + 1}页\n{p.extract_text()}' for i, p in enumerate(reader.pages))
    (ROOT / 'reports/数据与题意/题目原文提取.txt').write_text(extracted, encoding='utf-8')
    source = rows('附件1.xlsx')
    assert [minutes(r[0]) for r in source[1:]] == list(range(10, 1441, 10))
    q1 = pd.DataFrame([r[1:] for r in source[1:]], columns=['price_yuan_kWh', 'load_kW', 'pv_kW'])
    q1.insert(0, 'end_minute', range(10, 1441, 10))
    q1.insert(0, 'start_minute', range(0, 1440, 10))
    q1.to_csv(OUT / 'q1_normalized.csv', index=False, encoding='utf-8-sig')
    dates, ticks, load = annual('附件2.xlsx')
    _, _, pv = annual('附件2.xlsx', 1)
    _, _, price = annual('附件4.xlsx')
    end = pd.DatetimeIndex([d + pd.Timedelta(minutes=t) for d in dates for t in ticks])
    assert end.is_unique and ((end[1:] - end[:-1]) == pd.Timedelta(minutes=10)).all()
    actual = pd.DataFrame(dict(start_time=end-pd.Timedelta(minutes=10), end_time=end,
                              load_kW=load.ravel(), pv_kW=pv.ravel(), price_yuan_kWh=price.ravel()))
    actual.to_csv(OUT / 'actual_2025_normalized.csv', index=False, encoding='utf-8-sig')
    forecast_rows = []
    current_date = None
    for r in rows('附件3.xlsx')[1:]:
        if r[0] not in (None, ''):
            current_date = pd.Timestamp(r[0])
        issue = current_date + pd.Timedelta(minutes=minutes(r[1]))
        for lead, value in enumerate(r[2:], 1):
            forecast_rows.append((issue, issue + pd.Timedelta(hours=lead), lead, float(value)))
    forecast = pd.DataFrame(forecast_rows, columns=['issue_time', 'target_time', 'lead_hours', 'pv_forecast_kW'])
    assert len(forecast) == 365 * 4 * 24
    assert not forecast.duplicated(['issue_time', 'target_time']).any()
    assert sorted(forecast.issue_time.dt.hour.unique()) == [0, 6, 12, 18]
    forecast.to_csv(OUT / 'pv_forecast_versions.csv', index=False, encoding='utf-8-sig')
    matched = forecast.merge(actual[['end_time', 'pv_kW']], left_on='target_time', right_on='end_time', how='left')
    errors = []
    for lo, hi in [(1, 6), (7, 12), (13, 18), (19, 24)]:
        sub = matched[matched.lead_hours.between(lo, hi) & (matched.pv_kW > 100)]
        errors.append(dict(lead=f'{lo}-{hi}', n=len(sub),
                           wape=float((sub.pv_forecast_kW-sub.pv_kW).abs().sum()/sub.pv_kW.sum())))
    templates = {}
    for p in sources:
        if p.parent.name != '附件5':
            continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        templates[p.name] = {ws.title: dict(rows=ws.max_row, columns=ws.max_column,
                               first_rows=[[str(v) if v is not None else None for v in r]
                                           for r in ws.iter_rows(min_row=1, max_row=3, values_only=True)])
                            for ws in wb.worksheets}
        wb.close()
    summary = dict(environment=dict(python=sys.version, executable=sys.executable,
                                   numpy=np.__version__, pandas=pd.__version__, openpyxl=openpyxl.__version__),
                   source_sha256=checksums,
                   q1={c: stats(q1[c]) for c in q1.columns[2:]},
                   annual_load=stats(load), annual_pv=stats(pv), annual_price=stats(price),
                   forecast=stats(forecast.pv_forecast_kW),
                   first_end_time=str(end[0]), last_end_time=str(end[-1]),
                   forecast_unmatched=int(matched.pv_kW.isna().sum()),
                   duplicate_load_days=int(pd.DataFrame(load).duplicated().sum()),
                   duplicate_pv_days=int(pd.DataFrame(pv).duplicated().sum()),
                   duplicate_price_days=int(pd.DataFrame(price).duplicated().sum()),
                   surplus_slots=int((pv > load).sum()), surplus_days=int((pv > load).any(axis=1).sum()),
                   max_surplus_kW=float((pv-load).max()),
                   load_mean_by_weekday={str(i): float(load[dates.dayofweek == i].mean()) for i in range(7)},
                   forecast_daylight_wape_by_lead=errors, templates=templates,
                   assumptions=['样本作为其前10分钟区间的代表功率；并非已确认的区间平均值。',
                                '原始数值不插补、不平滑、不删除零值；附件3日期仅向下填充。',
                                '全年统计仅作描述，不能提前用于因果预测或调参。'])
    (OUT / 'audit_summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in summary.items() if k not in ['templates', 'source_sha256']}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
