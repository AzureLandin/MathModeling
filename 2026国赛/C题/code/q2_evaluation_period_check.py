"""Check result2 template and saved evaluation dates; no model or workbook mutation."""
from pathlib import Path
from datetime import datetime
import json
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
out = ROOT / 'results/q2_evaluation_period_check'
out.mkdir(exist_ok=True)
wb = load_workbook(ROOT / '附件/附件5/result2.xlsx', read_only=True, data_only=True)
sheets = []
for ws in wb.worksheets:
    dates = [r[0] for r in ws.iter_rows(values_only=True) if r and isinstance(r[0], datetime)]
    sheets.append(dict(sheet=ws.title, dated_rows=len(dates),
                       first=str(min(dates)) if dates else None,
                       last=str(max(dates)) if dates else None))
wb.close()
rows = []
for folder in ['q2_lightgbm_residual', 'q2_pv_support_features', 'q2_load_pv_shape']:
    for p in sorted((ROOT / 'results' / folder).glob('*/dispatch.csv')):
        f = pd.read_csv(p, usecols=['date', 'planned_cost_yuan', 'emergency_cost_yuan'])
        dates = pd.to_datetime(f.date)
        row = dict(group=p.parent.name, source=str(p.relative_to(ROOT)),
                   first=str(dates.min().date()), last=str(dates.max().date()),
                   days=int(dates.nunique()), slots=len(f),
                   january_rows=int((dates.dt.month == 1).sum()),
                   total_yuan=float(f.planned_cost_yuan.sum() + f.emergency_cost_yuan.sum()))
        assert (row['first'], row['last'], row['days'], row['slots'], row['january_rows']) == (
            '2025-02-01', '2025-12-31', 334, 48096, 0), row
        rows.append(row)
result = dict(scope='Read-only template and saved dispatch date/cost check, no forecasting or optimisation.',
              template=sheets, dispatch_checks=rows)
(out / 'check.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(result, ensure_ascii=True, indent=2))
