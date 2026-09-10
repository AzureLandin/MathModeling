"""Fill the supplied result1 workbook from the validated MILP solution."""
from pathlib import Path
import hashlib
import json
import shutil
import sys

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[1]


def main():
    assert Path(sys.prefix).name == 'math_modeling', sys.prefix
    target = ROOT/'附件/附件5/result1.xlsx'
    backup = ROOT/'results/template_backups/result1_original.xlsx'
    backup.parent.mkdir(parents=True, exist_ok=True)
    audit = json.loads((ROOT/'results/audit/audit_summary.json').read_text(encoding='utf-8'))
    original_hash = audit['source_sha256'][str(target.relative_to(ROOT))]
    log_path = ROOT/'results/q1_milp/result1_fill_validation.json'
    if log_path.exists():
        original_hash = json.loads(log_path.read_text(encoding='utf-8'))['original_sha256']
    if not backup.exists():
        assert hashlib.sha256(target.read_bytes()).hexdigest() == original_hash
        shutil.copy2(target, backup)
    assert hashlib.sha256(backup.read_bytes()).hexdigest() == original_hash
    data = pd.read_csv(ROOT/'results/q1_milp/dispatch_10min.csv')
    summary = json.loads((ROOT/'results/q1_milp/validation_summary.json').read_text(encoding='utf-8'))['summary']
    assert len(data) == 144 and summary['mip_gap'] == 0
    assert np.array_equal(data.start_minute, np.arange(0,1440,10))
    wb = openpyxl.load_workbook(backup)
    assert wb.sheetnames == ['计划购电量','充放电量']
    plan, storage = wb.worksheets
    for i,r in data.iterrows():
        cell = plan.cell(i+2,1)
        cell.comment = Comment(f'原模板标签：{cell.value}。按样本代表前10分钟区间的口径修正。','Modeling')
        cell.value = r.interval
        plan.cell(i+2,2,float(r.grid_kWh)).number_format = '0.000000'
    plan['B1'].comment = Comment(
        f"MILP结果；单位kWh。全天购电量{summary['grid_kWh']:.9f} kWh；全天购电费{summary['cost_yuan']:.9f}元。"
        '功率样本代表前10分钟区间，时段标签已修正；原模板已备份。','Modeling')
    groups = np.array([[data.charge_bus_kWh.iloc[t:t+24].sum(),
                        data.discharge_bus_kWh.iloc[t:t+24].sum()] for t in range(0,144,24)])
    for i,(charge,discharge) in enumerate(groups,2):
        storage.cell(i,2,float(charge)).number_format = '0.000000'
        storage.cell(i,3,float(discharge)).number_format = '0.000000'
    storage['E2'],storage['E3'] = float(data.state_start_kWh.iloc[0]),float(data.state_end_kWh.iloc[-1])
    for address in ['E2','E3']:
        storage[address].number_format = '0.000000'
    storage['B1'].comment = Comment('母线侧充电量，单位kWh；充电效率0.9、放电效率0.9（往返0.81）。充电可来自外网购电或光伏。','Modeling')
    storage['C1'].comment = Comment('母线侧放电量，单位kWh；同一4小时区间内可在不同10分钟时段分别充电、放电。','Modeling')
    storage['E1'].comment = Comment('电池内部储电量，单位kWh；日初和日末均为6000。','Modeling')
    staged = ROOT/'results/q1_milp/result1_filled_verified.xlsx'
    wb.save(staged)
    wb.close()
    check = openpyxl.load_workbook(staged,data_only=True)
    actual_q = np.array([r[1] for r in list(check['计划购电量'].values)[1:]])
    actual_groups = np.array([r[1:3] for r in list(check['充放电量'].values)[1:]])
    assert np.allclose(actual_q,data.grid_kWh,rtol=0,atol=1e-9)
    assert np.allclose(actual_groups,groups,rtol=0,atol=1e-9)
    assert abs(actual_q.sum()-summary['grid_kWh']) < 1e-7
    assert abs(actual_q@data.price_yuan_kWh-summary['cost_yuan']) < 1e-7
    assert check['充放电量']['E2'].value == check['充放电量']['E3'].value == 6000
    assert check['计划购电量']['A2'].value == '00:00-00:10'
    assert check['计划购电量']['A145'].value == '23:50-24:00'
    check.close()
    shutil.copyfile(staged,target)
    filled_hash = hashlib.sha256(target.read_bytes()).hexdigest()
    assert filled_hash == hashlib.sha256(staged.read_bytes()).hexdigest()
    log = dict(target=str(target),backup=str(backup),original_sha256=original_hash,
               filled_sha256=filled_hash,plan_values=144,charge_discharge_values=12,state_values=2,
               grid_kWh=float(actual_q.sum()),cost_yuan=float(actual_q@data.price_yuan_kWh),
               all_checks_passed=True,time_labels='00:00-00:10 through 23:50-24:00; original labels in comments')
    log_path.write_text(json.dumps(log,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(log,ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
