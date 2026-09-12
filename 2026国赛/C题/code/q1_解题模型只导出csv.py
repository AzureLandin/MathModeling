"""问题1：确定性日前购电与储能调度的 MILP 模型。

本脚本只负责建模、求解与结果落盘（CSV / JSON），不依赖 matplotlib。
绘图仅使用 output/q1_intervals.csv 和 output/q1_four_hour.csv。
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import numpy as np
import scipy
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix

ROOT = Path(__file__).resolve().parent
INPUT = ROOT / "附件" / "附件1.xlsx"
WORK = ROOT / "work"
OUTPUT = ROOT / "output"

RESULT_JSON = WORK / "q1_result.json"
INTERVALS_CSV = OUTPUT / "q1_intervals.csv"
FOUR_HOUR_CSV = OUTPUT / "q1_four_hour.csv"

N = 144
DT = 1.0 / 6.0
ETA_C = 0.9
ETA_D = 0.9
E_MIN = 1200.0
E_MAX = 10800.0
E_INITIAL = 6000.0
POWER_MAX = 5000.0
ENERGY_MAX = POWER_MAX * DT
FEASIBILITY_TOLERANCE = 1e-6
MIP_RELATIVE_GAP = 1e-9
MIP_TIME_LIMIT_SECONDS = 120.0

DESIGNATED_INDICES = [60, 72, 84, 96, 108, 120]

BLOCKS = ("q", "c", "d", "w", "E", "z")


def idx(block: int, t: int) -> int:
    if block < 4:
        return block * N + t
    if block == 4:
        return 4 * N + t
    if block == 5:
        return 5 * N + 1 + t
    raise ValueError(f"unknown block: {block}")


def interval_label(t: int) -> str:
    start, end = t * 10, (t + 1) * 10

    def fmt(minutes: int) -> str:
        return "24:00" if minutes == 1440 else f"{minutes // 60:02d}:{minutes % 60:02d}"

    return f"{fmt(start)}-{fmt(end)}"


def read_input() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    wb = load_workbook(INPUT, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    price = np.array([float(ws.cell(r, 2).value) for r in range(2, 146)])
    load_kw = np.array([float(ws.cell(r, 3).value) for r in range(2, 146)])
    pv_kw = np.array([float(ws.cell(r, 4).value) for r in range(2, 146)])
    wb.close()
    if not (len(price) == len(load_kw) == len(pv_kw) == N):
        raise ValueError("附件1必须包含144个10分钟采样点")
    return price, load_kw, pv_kw


def build_constraints(price, load_e, pv_e):
    nvar = 6 * N + 1
    rows = 2 * N + 2 + 2 * N
    a = lil_matrix((rows, nvar), dtype=float)
    lb = np.full(rows, -np.inf)
    ub = np.full(rows, np.inf)
    row = 0

    for t in range(N):
        a[row, idx(0, t)] = 1.0
        a[row, idx(1, t)] = -1.0
        a[row, idx(2, t)] = 1.0
        a[row, idx(3, t)] = -1.0
        rhs = load_e[t] - pv_e[t]
        lb[row] = ub[row] = rhs
        row += 1

    for t in range(N):
        a[row, idx(4, t + 1)] = 1.0
        a[row, idx(4, t)] = -1.0
        a[row, idx(1, t)] = -ETA_C
        a[row, idx(2, t)] = 1.0 / ETA_D
        lb[row] = ub[row] = 0.0
        row += 1

    a[row, idx(4, 0)] = 1.0
    lb[row] = ub[row] = E_INITIAL
    row += 1
    a[row, idx(4, N)] = 1.0
    lb[row] = ub[row] = E_INITIAL
    row += 1

    for t in range(N):
        a[row, idx(1, t)] = 1.0
        a[row, idx(5, t)] = -ENERGY_MAX
        ub[row] = 0.0
        row += 1
    for t in range(N):
        a[row, idx(2, t)] = 1.0
        a[row, idx(5, t)] = ENERGY_MAX
        ub[row] = ENERGY_MAX
        row += 1

    assert row == rows

    lower = np.zeros(nvar)
    upper = np.full(nvar, np.inf)
    lower[idx(4, 0): idx(4, N) + 1] = E_MIN
    upper[idx(4, 0): idx(4, N) + 1] = E_MAX
    lower[idx(5, 0): idx(5, N - 1) + 1] = 0.0
    upper[idx(5, 0): idx(5, N - 1) + 1] = 1.0

    integrality = np.zeros(nvar, dtype=int)
    integrality[idx(5, 0): idx(5, N - 1) + 1] = 1

    objective = np.zeros(nvar)
    objective[idx(0, 0): idx(0, N - 1) + 1] = price
    return a.tocsr(), lb, ub, Bounds(lower, upper), integrality, objective


def solve_dispatch(price, load_e, pv_e):
    a, lb, ub, bounds, integrality, objective = build_constraints(price, load_e, pv_e)
    result = milp(
        c=objective,
        integrality=integrality,
        bounds=bounds,
        constraints=LinearConstraint(a, lb, ub),
        options={
            "disp": False,
            "mip_rel_gap": MIP_RELATIVE_GAP,
            "time_limit": MIP_TIME_LIMIT_SECONDS,
        },
    )
    if not result.success:
        raise RuntimeError(f"MILP求解失败: {result.message}")

    dual_bound = float(getattr(result, "mip_dual_bound", result.fun))
    info: dict[str, float | int | str] = {
        "solver": "SciPy HiGHS MILP",
        "solver_message": str(result.message),
        "scipy_version": scipy.__version__,
        "mip_gap": float(getattr(result, "mip_gap", 0.0)),
        "mip_dual_bound_yuan": dual_bound,
        "objective_bound_gap_yuan": float(result.fun - dual_bound),
        "mip_node_count": int(getattr(result, "mip_node_count", 0)),
    }
    return result.x, info


def clean(values, tol: float = 1e-9) -> np.ndarray:
    values = np.asarray(values, dtype=float).copy()
    values[np.abs(values) < tol] = 0.0
    return values


def unpack(x):
    arrays = {
        "q": clean(x[idx(0, 0): idx(0, N - 1) + 1]),
        "c": clean(x[idx(1, 0): idx(1, N - 1) + 1]),
        "d": clean(x[idx(2, 0): idx(2, N - 1) + 1]),
        "w": clean(x[idx(3, 0): idx(3, N - 1) + 1]),
        "e": clean(x[idx(4, 0): idx(4, N) + 1]),
    }
    z_raw = x[idx(5, 0): idx(5, N - 1) + 1]
    arrays["z"] = np.rint(z_raw).astype(int)
    arrays["z_raw"] = z_raw
    return arrays


def verify(arrays, load_e, pv_e, price, solver_info):
    q, c, d, w, e = arrays["q"], arrays["c"], arrays["d"], arrays["w"], arrays["e"]
    z_raw, z = arrays["z_raw"], arrays["z"]

    balance = q + pv_e + d - load_e - c - w
    state = e[1:] - e[:-1] - ETA_C * c + d / ETA_D
    simultaneous = np.minimum(c, d)
    binary_integrality_error = float(np.max(np.abs(z_raw - z)))
    binary_gate_violation = float(
        max(0.0, np.max(c - ENERGY_MAX * z_raw), np.max(d - ENERGY_MAX * (1.0 - z_raw)))
    )
    baseline_q = np.maximum(load_e - pv_e, 0.0)
    baseline_cost = float(price @ baseline_q)
    optimized_cost = float(price @ q)

    assert np.max(np.abs(balance)) < FEASIBILITY_TOLERANCE
    assert np.max(np.abs(state)) < FEASIBILITY_TOLERANCE
    assert np.max(simultaneous) < FEASIBILITY_TOLERANCE
    assert c.max() <= ENERGY_MAX + FEASIBILITY_TOLERANCE
    assert d.max() <= ENERGY_MAX + FEASIBILITY_TOLERANCE
    assert e.min() >= E_MIN - FEASIBILITY_TOLERANCE
    assert e.max() <= E_MAX + FEASIBILITY_TOLERANCE
    assert abs(e[0] - E_INITIAL) < FEASIBILITY_TOLERANCE
    assert abs(e[-1] - E_INITIAL) < FEASIBILITY_TOLERANCE
    assert binary_integrality_error < FEASIBILITY_TOLERANCE
    assert binary_gate_violation < FEASIBILITY_TOLERANCE
    assert abs(float(solver_info["objective_bound_gap_yuan"])) < 1e-5

    checks = {
        "max_balance_residual_kwh": float(np.max(np.abs(balance))),
        "max_state_residual_kwh": float(np.max(np.abs(state))),
        "max_simultaneous_charge_discharge_kwh": float(np.max(simultaneous)),
        "max_charge_kwh_per_interval": float(c.max()),
        "max_discharge_kwh_per_interval": float(d.max()),
        "initial_storage_kwh": float(e[0]),
        "final_storage_kwh": float(e[-1]),
        "binary_charge_intervals": int(z.sum()),
        "binary_integrality_error": binary_integrality_error,
        "binary_gate_violation_kwh": binary_gate_violation,
    }
    return checks, baseline_q, baseline_cost, optimized_cost


def write_csv(path: Path, header, rows) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(header)
        writer.writerows(rows)


def main() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    OUTPUT.mkdir(parents=True, exist_ok=True)

    price, load_kw, pv_kw = read_input()
    load_e, pv_e = load_kw * DT, pv_kw * DT

    x, solver_info = solve_dispatch(price, load_e, pv_e)
    arrays = unpack(x)
    q, c, d, w, e = arrays["q"], arrays["c"], arrays["d"], arrays["w"], arrays["e"]

    checks, baseline_q, baseline_cost, optimized_cost = verify(
        arrays, load_e, pv_e, price, solver_info
    )

    four_hour = [
        {
            "period": f"{start // 6:02d}:00-{(start + 24) // 6:02d}:00",
            "charge_kwh": float(c[start:start + 24].sum()),
            "discharge_kwh": float(d[start:start + 24].sum()),
        }
        for start in range(0, N, 24)
    ]

    result = {
        "model": {
            "method": "deterministic single-stage MILP",
            "interval_minutes": 10,
            "sample_interpretation": "附件1采样时刻作为内部10分钟时段的右端点",
            "eta_charge": ETA_C,
            "eta_discharge": ETA_D,
            "battery_min_kwh": E_MIN,
            "battery_max_kwh": E_MAX,
            "battery_initial_kwh": E_INITIAL,
            "battery_final_kwh": E_INITIAL,
            "power_limit_kw": POWER_MAX,
            "feasibility_tolerance": FEASIBILITY_TOLERANCE,
            "mip_relative_gap_target": MIP_RELATIVE_GAP,
            "mip_time_limit_seconds": MIP_TIME_LIMIT_SECONDS,
        },
        "summary": {
            "optimized_purchase_kwh": float(q.sum()),
            "optimized_cost_yuan": optimized_cost,
            "baseline_purchase_kwh": float(baseline_q.sum()),
            "baseline_cost_yuan": baseline_cost,
            "cost_saving_yuan": baseline_cost - optimized_cost,
            "cost_saving_pct": (baseline_cost - optimized_cost) / baseline_cost,
            "charge_kwh": float(c.sum()),
            "discharge_kwh": float(d.sum()),
            "spilled_energy_kwh": float(w.sum()),
            "min_storage_kwh": float(e.min()),
            "max_storage_kwh": float(e.max()),
            **solver_info,
        },
        "designated_periods": [
            {"period": interval_label(t), "purchase_kwh": float(q[t])} for t in DESIGNATED_INDICES
        ],
        "four_hour_periods": four_hour,
        "checks": checks,
        "intervals": [
            {
                "period": interval_label(t),
                "price_yuan_per_kwh": float(price[t]),
                "load_kw": float(load_kw[t]),
                "pv_kw": float(pv_kw[t]),
                "purchase_kwh": float(q[t]),
                "charge_kwh": float(c[t]),
                "discharge_kwh": float(d[t]),
                "spill_kwh": float(w[t]),
                "storage_start_kwh": float(e[t]),
                "storage_end_kwh": float(e[t + 1]),
            }
            for t in range(N)
        ],
    }

    RESULT_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    write_csv(
        INTERVALS_CSV,
        [
            "period", "start_hour", "end_hour", "price_yuan_per_kwh", "load_kw", "pv_kw",
            "purchase_kwh", "charge_kwh", "discharge_kwh", "spill_kwh",
            "storage_start_kwh", "storage_end_kwh",
        ],
        [
            [
                interval_label(t), t / 6.0, (t + 1) / 6.0, float(price[t]), float(load_kw[t]),
                float(pv_kw[t]), float(q[t]), float(c[t]), float(d[t]), float(w[t]),
                float(e[t]), float(e[t + 1]),
            ]
            for t in range(N)
        ],
    )

    write_csv(
        FOUR_HOUR_CSV,
        ["period", "charge_kwh", "discharge_kwh"],
        [[r["period"], r["charge_kwh"], r["discharge_kwh"]] for r in four_hour],
    )

    print(
        json.dumps(
            {
                "summary": result["summary"],
                "checks": result["checks"],
                "files": {
                    "json": str(RESULT_JSON),
                    "intervals_csv": str(INTERVALS_CSV),
                    "four_hour_csv": str(FOUR_HOUR_CSV),
                },
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
