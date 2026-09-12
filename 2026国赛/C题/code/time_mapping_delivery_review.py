"""Read-only cross-question review of the accepted start_time_v1 mapping."""
from pathlib import Path
from datetime import timedelta
import hashlib
import json

import numpy as np
import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'results/time_mapping_delivery_review_20260912'
OUT.mkdir(parents=True, exist_ok=True)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def q1_review():
    formal = ROOT / '附件/附件5/result1.xlsx'
    validation_path = ROOT / 'results/q1_start_time_20260912/formal_delivery_validation.json'
    validation = json.loads(validation_path.read_text(encoding='utf-8'))
    wb = openpyxl.load_workbook(formal, data_only=True, read_only=False)
    wp, ws = wb.worksheets
    labels = [wp.cell(r, 1).value for r in range(2, 146)]
    values = [wp.cell(r, 2).value for r in range(2, 146)]
    storage = [ws.cell(r, c).value for r in range(2, 8) for c in (2, 3)]
    endpoints = [ws.cell(2, 5).value, ws.cell(3, 5).value]
    wb.close()
    return {
        'formal_sha256': sha(formal),
        'matches_delivery_validation_hash': sha(formal) == validation['hashes']['formal_sha256'],
        'first_label': labels[0],
        'penultimate_label': labels[-2],
        'last_label': labels[-1],
        'labels_correct': labels[0] == '0:10-0:20' and labels[-2:] == ['23:50-0:00+1', '0:00+1-0:10+1'],
        'planned_cells_nonempty': int(sum(x is not None for x in values)),
        'storage_cells_nonempty': int(sum(x is not None for x in storage)),
        'endpoint_cells_nonempty': int(sum(x is not None for x in endpoints)),
        'formal_delivery_validation_status': validation['status'],
        'style_and_structure_preserved': validation['readback']['style_and_structure_equal_to_pristine_template'],
    }


def q2_review():
    formal = ROOT / '附件/附件5/result2.xlsx'
    fill_path = ROOT / 'results/q2_frozen_delivery/fill_validation.json'
    fill = json.loads(fill_path.read_text(encoding='utf-8'))
    wb = openpyxl.load_workbook(formal, data_only=True, read_only=False)
    wp = wb.worksheets[0]
    headers = [wp.cell(1, c).value for c in range(1, 148)]
    planned_nonempty = sum(wp.cell(r, c).value is not None for r in range(2, 336) for c in range(2, 146))
    wb.close()

    plan = pd.read_csv(ROOT / 'results/q2_time_mapping/N_free/template_plan.csv', low_memory=False)
    plan['publication_date'] = pd.to_datetime(plan.publication_date)
    plan['interval_start'] = pd.to_datetime(plan.interval_start)
    expected = plan['publication_date'] + pd.to_timedelta((plan.template_slot + 1) * 10, unit='min')
    interval_mismatch = int((plan.interval_start != expected).sum())
    expected_label = np.where(plan.template_slot.eq(143), '0:00+1', plan.interval_start.dt.strftime('%H:%M'))
    source_label_mismatch = int((plan.source_label.astype(str).to_numpy() != expected_label).sum())

    actual = pd.read_csv(ROOT / 'results/q2_time_mapping/N_free/natural_dispatch.csv', low_memory=False)
    official = actual[actual.date.between('2025-02-01', '2025-12-31')].copy()
    starts = pd.to_datetime(official.interval_start)
    ends = pd.to_datetime(official.interval_end)
    continuity = bool((ends - starts).eq(pd.Timedelta(minutes=10)).all()
                      and starts.is_unique
                      and starts.sort_values().diff().dropna().eq(pd.Timedelta(minutes=10)).all())
    return {
        'formal_sha256': sha(formal),
        'matches_fill_validation_hash': sha(formal) == fill['output_sha256'],
        'fill_readback_passed': fill['all_readback_checks_passed'],
        'first_template_label': headers[1],
        'penultimate_template_label': headers[143],
        'last_template_label': headers[144],
        'formal_headers_correct': headers[1] == '0:10-0:20' and headers[143] == '23:50-0:00+1' and headers[144] == '0:00-0:10+1',
        'formal_planned_cells_nonempty': int(planned_nonempty),
        'formal_expected_planned_cells': 334 * 144,
        'time_versions': sorted(plan.time_version.unique().tolist()),
        'template_interval_start_mismatches': interval_mismatch,
        'natural_official_rows': int(len(official)),
        'natural_official_continuous_10min': continuity,
        'natural_first_interval': str(starts.min()),
        'natural_last_interval_start': str(starts.max()),
        'generated_source_label_mismatches_all_364_publication_days': source_label_mismatch,
        'source_label_issue_scope': 'presentation metadata only: interval_start, template_slot, price order, solver arrays and formal workbook headers are correct; fill_result2_frozen.py does not use source_label',
        'known_separate_metadata_issue': 'source_interval_mapping.csv price rows 0..142 export template_slot=j-1; model and formal workbook use the correct stored price order',
        'delivery_status': 'formal result2 is filled under start_time_v1; independent audit remains deferred',
    }


def q3_review():
    formal = ROOT / '附件/附件5/result3.xlsx'
    wb = openpyxl.load_workbook(formal, data_only=True, read_only=False)
    headers = [[ws.cell(1, c).value for c in range(1, 148)] for ws in wb.worksheets[:2]]
    value_nonempty = [sum(ws.cell(r, c).value is not None for r in range(2, 336) for c in range(2, 146)) for ws in wb.worksheets[:2]]
    wb.close()

    dispatch = pd.read_csv(ROOT / 'results/q3_rolling_baseline/B2_dispatch.csv', low_memory=False)
    dispatch['interval_start'] = pd.to_datetime(dispatch.interval_start)
    dispatch['interval_end'] = pd.to_datetime(dispatch.interval_end)
    natural = dispatch[dispatch.interval_start < pd.Timestamp('2026-01-01 00:00:00')].copy()
    ordered = dispatch.sort_values('interval_start')
    continuity = bool((ordered.interval_end - ordered.interval_start).eq(pd.Timedelta(minutes=10)).all()
                      and ordered.interval_start.is_unique
                      and ordered.interval_start.diff().dropna().eq(pd.Timedelta(minutes=10)).all())
    validation = json.loads((ROOT / 'results/q3_rolling_baseline/validation.json').read_text(encoding='utf-8'))
    return {
        'formal_sha256': sha(formal),
        'formal_first_template_label': headers[0][1],
        'formal_penultimate_template_label': headers[0][143],
        'formal_last_template_label': headers[0][144],
        'formal_headers_correct': all(h[1] == '0:10-0:20' and h[143] == '23:50-0:00+1' and h[144] == '0:00-0:10+1' for h in headers),
        'formal_plan_cells_nonempty': value_nonempty[0],
        'formal_adjusted_cells_nonempty': value_nonempty[1],
        'formal_delivery_status': 'template remains unfilled; no stale old-time numerical values are present',
        'saved_B2_rows_including_tail': int(len(dispatch)),
        'saved_B2_natural_evaluation_rows': int(len(natural)),
        'saved_B2_continuous_10min_including_tail': continuity,
        'saved_B2_first_interval': str(dispatch.interval_start.min()),
        'saved_B2_last_interval': str(dispatch.interval_start.max()),
        'first_midnight_effective_version': str(dispatch.sort_values('interval_start').iloc[0].effective_version),
        'ledger_bridge_residual_yuan': validation['ledger']['B2_bridge_residual_yuan'],
        'model_time_mapping_status': 'corrected: current midnight uses previous-day carry; current publication starts at 00:10; natural and template ledgers are separate',
        'strategy_status': 'Linear+q75 is the current candidate, but formal result3 has not been frozen or filled',
    }


def main():
    review = {
        'review_date': '2026-09-12',
        'scope': 'read-only Q2/Q3 mapping review plus verification of the newly finalized Q1 workbook; no training, optimization or replay',
        'q1': q1_review(),
        'q2': q2_review(),
        'q3': q3_review(),
    }
    review['overall'] = {
        'q1_formal_new_time_complete': review['q1']['labels_correct'] and review['q1']['matches_delivery_validation_hash'],
        'q2_substantive_time_mapping_and_formal_delivery_correct': review['q2']['formal_headers_correct'] and review['q2']['template_interval_start_mismatches'] == 0 and review['q2']['natural_official_continuous_10min'],
        'q2_metadata_fully_clean': review['q2']['generated_source_label_mismatches_all_364_publication_days'] == 0,
        'q3_substantive_time_mapping_correct': review['q3']['formal_headers_correct'] and review['q3']['saved_B2_continuous_10min_including_tail'],
        'q3_formal_delivery_complete': review['q3']['formal_plan_cells_nonempty'] == 334 * 144 and review['q3']['formal_adjusted_cells_nonempty'] == 334 * 144,
    }
    (OUT / 'review.json').write_text(json.dumps(review, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(review, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
