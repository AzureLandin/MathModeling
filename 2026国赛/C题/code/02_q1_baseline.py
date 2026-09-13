"""Deterministic Q1 baseline, independent physical checks and sensitivities.

Input: results/audit/q1_normalized.csv, original result1 template.
Output: results/q1_baseline, figures/q1_baseline.png, reports/问题一/问题1_Baseline结果.md.
All energy flows are AC bus-side kWh, state is battery-internal kWh.
"""
from pathlib import Path
import json
import sys
import time

import numpy as np
import pandas as pd
import scipy
from scipy.optimize import linprog, milp, Bounds, LinearConstraint
from scipy.sparse import lil_matrix, csr_matrix, hstack, vstack
import openpyxl
from openpyxl.comments import Comment
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'results' / 'q1_baseline'
CONFIG = dict(dt_hours=1/6, capacity_kWh=12000., state_min_kWh=1200.,
              state_max_kWh=10800., initial_kWh=6000., power_max_kW=5000.,
              charge_efficiency=0.9, discharge_efficiency=0.9,
              feasibility_tolerance=1e-6, milp_relative_gap=1e-9,
              milp_time_limit_seconds=120, random_seed=20260910)


def solve(load, pv, price, eta=0.9, power=5000., integer=False,
          initial_kWh=None, terminal_kWh=None, state_min_kWh=None,
          state_max_kWh=None):
    initial_kWh = CONFIG['initial_kWh'] if initial_kWh is None else float(initial_kWh)
    terminal_kWh = initial_kWh if terminal_kWh is None else float(terminal_kWh)
    state_min_kWh = CONFIG['state_min_kWh'] if state_min_kWh is None else float(state_min_kWh)
    state_max_kWh = CONFIG['state_max_kWh'] if state_max_kWh is None else float(state_max_kWh)
    if not state_min_kWh <= initial_kWh <= state_max_kWh:
        raise ValueError('initial_kWh must lie within the state bounds')
    if not state_min_kWh <= terminal_kWh <= state_max_kWh:
        raise ValueError('terminal_kWh must lie within the state bounds')
    n = len(load)
    m = 5*n + 1
    cap = power * CONFIG['dt_hours']
    # Vector: grid q, charge c, discharge d, unused w, battery states E[0:n+1].
    obj = np.zeros(m)
    obj[:n] = price
    eq = lil_matrix((2*n, m))
    rhs = np.r_[(load-pv)*CONFIG['dt_hours'], np.zeros(n)]
    for t in range(n):
        eq[t, t], eq[t, n+t], eq[t, 2*n+t], eq[t, 3*n+t] = 1, -1, 1, -1
        eq[n+t, 4*n+t+1], eq[n+t, 4*n+t] = 1, -1
        eq[n+t, n+t], eq[n+t, 2*n+t] = -eta, 1/eta
    eq = csr_matrix(eq)
    lower, upper = np.zeros(m), np.full(m, np.inf)
    upper[n:3*n] = cap
    lower[4*n:], upper[4*n:] = state_min_kWh, state_max_kWh
    lower[4*n] = upper[4*n] = initial_kWh
    lower[-1] = upper[-1] = terminal_kWh
    start = time.perf_counter()
    if integer:
        constraints = hstack([eq, csr_matrix((2*n, n))], format='csr')
        mutex = lil_matrix((2*n, m+n))
        for t in range(n):
            mutex[t, n+t], mutex[t, m+t] = 1, -cap
            mutex[n+t, 2*n+t], mutex[n+t, m+t] = 1, cap
        constraints = vstack([constraints, csr_matrix(mutex)], format='csr')
        result = milp(np.r_[obj, np.zeros(n)], integrality=np.r_[np.zeros(m), np.ones(n)],
                      bounds=Bounds(np.r_[lower, np.zeros(n)], np.r_[upper, np.ones(n)]),
                      constraints=LinearConstraint(constraints, np.r_[rhs, np.full(2*n, -np.inf)],
                                                   np.r_[rhs, np.zeros(n), np.full(n, cap)]),
                      options=dict(mip_rel_gap=CONFIG['milp_relative_gap'],
                                   time_limit=CONFIG['milp_time_limit_seconds']))
    else:
        result = linprog(obj, A_eq=eq, b_eq=rhs, bounds=list(zip(lower, upper)), method='highs')
    elapsed = time.perf_counter()-start
    if not result.success:
        raise RuntimeError(result.message)
    x = result.x[:m]
    q, c, d, w, state = x[:n], x[n:2*n], x[2*n:3*n], x[3*n:4*n], x[4*n:]
    checks = dict(bus_balance_max_abs_kWh=float(np.max(np.abs(q+pv/6+d-load/6-c-w))),
                  battery_balance_max_abs_kWh=float(np.max(np.abs(np.diff(state)-eta*c+d/eta))),
                  boundary_error_kWh=float(max(abs(state[0]-initial_kWh), abs(state[-1]-terminal_kWh))),
                   state_bound_violation_kWh=float(max(0, state_min_kWh-state.min(), state.max()-state_max_kWh)),
                  flow_bound_violation_kWh=float(max(0, -min(q.min(), c.min(), d.min(), w.min()),
                                                     c.max()-cap, d.max()-cap)),
                  simultaneous_charge_discharge_kWh=float(np.minimum(c, d).max()))
    if integer:
        z = result.x[m:]
        checks['binary_integrality_error'] = float(np.max(np.abs(z-np.rint(z))))
        checks['binary_bound_violation'] = float(max(0, -z.min(), z.max()-1))
        checks['binary_gate_violation_kWh'] = float(max(0, np.max(c-cap*z), np.max(d-cap*(1-z))))
    assert max(checks.values()) < CONFIG['feasibility_tolerance'], checks
    summary = dict(cost_yuan=float(price@q), grid_kWh=float(q.sum()), charge_kWh=float(c.sum()),
                   discharge_kWh=float(d.sum()), unused_kWh=float(w.sum()),
                   state_min_kWh=float(state.min()), state_max_kWh=float(state.max()),
                   initial_kWh=float(state[0]), final_kWh=float(state[-1]),
                   loss_kWh=float(((1-eta)*c+(1/eta-1)*d).sum()),
                   solver='HiGHS MILP' if integer else 'HiGHS LP', solver_message=result.message,
                   elapsed_seconds=elapsed, mip_gap=float(getattr(result, 'mip_gap', 0)), checks=checks)
    if integer:
        summary.update(binary_mode_raw=z.tolist(), mip_node_count=int(result.mip_node_count),
                       mip_dual_bound_yuan=float(result.mip_dual_bound),
                       objective_bound_gap_yuan=float(result.fun-result.mip_dual_bound))
    return (q, c, d, w, state), summary


def label(t):
    return f'{t//60:02d}:{t%60:02d}'


def main():
    assert Path(sys.prefix).name == 'math_modeling', sys.prefix
    np.random.seed(CONFIG['random_seed'])
    OUT.mkdir(parents=True, exist_ok=True)
    data = pd.read_csv(ROOT / 'results/audit/q1_normalized.csv')
    load, pv, price = [data[c].to_numpy() for c in ['load_kW', 'pv_kW', 'price_yuan_kWh']]
    x, baseline = solve(load, pv, price)
    _, integer_check = solve(load, pv, price, integer=True)
    assert abs(integer_check['cost_yuan']-baseline['cost_yuan']) < 1e-5
    no_grid = np.maximum(load-pv, 0)/6
    no_storage = dict(cost_yuan=float(price@no_grid), grid_kWh=float(no_grid.sum()),
                      unused_kWh=float(np.maximum(pv-load, 0).sum()/6))
    _, no_storage_solver = solve(load, pv, price, power=0)
    assert abs(no_storage_solver['cost_yuan']-no_storage['cost_yuan']) < 1e-6
    variations = {}
    for key, l, p, eff in [
        ('round_trip_efficiency_90pct', load, pv, np.sqrt(.9)),
        ('trapezoid_power_periodic_midnight', (load+np.roll(load, 1))/2, (pv+np.roll(pv, 1))/2, .9),
        ('load_plus_5pct', load*1.05, pv, .9),
        ('load_minus_5pct', load*.95, pv, .9),
        ('pv_plus_5pct', load, pv*1.05, .9),
        ('pv_minus_5pct', load, pv*.95, .9),
    ]:
        _, variations[key] = solve(l, p, price, eta=eff)
    q, c, d, w, state = x
    data.insert(0, 'interval', [f'{label(t)}-{label(t+10)}' for t in range(0,1440,10)])
    for key, value in dict(grid_kWh=q, charge_bus_kWh=c, discharge_bus_kWh=d,
                           unused_kWh=w, state_start_kWh=state[:-1], state_end_kWh=state[1:],
                           cost_yuan=q*price).items():
        data[key] = value
    data.to_csv(OUT / 'dispatch_10min.csv', index=False, encoding='utf-8-sig')
    groups = pd.DataFrame([dict(interval=f'{label(t*10)}-{label((t+24)*10)}',
                                charge_kWh=float(c[t:t+24].sum()), discharge_kWh=float(d[t:t+24].sum()))
                           for t in range(0,144,24)])
    groups.to_csv(OUT / 'charge_discharge_4hour.csv', index=False, encoding='utf-8-sig')
    selected = data[data.start_minute.isin([600,720,840,960,1080,1200])][['interval','grid_kWh']]
    selected.to_csv(OUT / 'selected_intervals.csv', index=False, encoding='utf-8-sig')
    template = ROOT / 'results/template_backups/result1_original.xlsx'
    if not template.exists():
        template = ROOT / '附件/附件5/result1.xlsx'
    wb = openpyxl.load_workbook(template)
    ws = wb['计划购电量']
    mapping = wb.create_sheet('时间映射与口径')
    mapping.append(['原模板标签', '计算时段', '原数据采样时刻（分钟）'])
    for i, row in data.iterrows():
        mapping.append([ws.cell(i+2,1).value, row.interval, int(row.end_minute)])
        ws.cell(i+2,1,row.interval)
        ws.cell(i+2,2,float(row.grid_kWh)).number_format = '0.000000'
    ws['A1'].comment = Comment('本文件为Baseline计算副本。按右端点代表功率假设修正时段；原标签见时间映射与口径。正式提交前需核对模板时间口径。', 'Modeling')
    ws = wb['充放电量']
    for i, row in groups.iterrows():
        ws.cell(i+2,2,float(row.charge_kWh))
        ws.cell(i+2,3,float(row.discharge_kWh))
    ws['E2'], ws['E3'] = float(state[0]), float(state[-1])
    note = wb.create_sheet('Baseline说明')
    for row in [
        ['状态','条件性Baseline，非最终提交文件'],
        ['时间口径','10分钟样本为前一时段的代表功率，模板标签已明确修正'],
        ['充放电量单位','母线侧kWh；储电量为电池内部kWh'],
        ['效率','充电效率0.9，放电效率0.9，往返0.81'],
        ['全天购电量_kWh',baseline['grid_kWh']], ['全天购电费_元',baseline['cost_yuan']],
        ['无储能费用_元',no_storage['cost_yuan']],
        ['LP与MILP费用差_元',abs(integer_check['cost_yuan']-baseline['cost_yuan'])],
    ]:
        note.append(row)
    wb.save(OUT / 'result1_baseline.xlsx')
    wb.close()
    check_wb = openpyxl.load_workbook(OUT / 'result1_baseline.xlsx', data_only=True)
    check_q = np.array([r[1] for r in list(check_wb['计划购电量'].values)[1:]], dtype=float)
    assert np.allclose(check_q, q, atol=1e-9, rtol=0)
    assert check_wb['计划购电量']['A2'].value == '00:00-00:10'
    assert check_wb['计划购电量']['A145'].value == '23:50-24:00'
    check_wb.close()
    (OUT/'config.json').write_text(json.dumps(CONFIG, ensure_ascii=False, indent=2), encoding='utf-8')
    result = dict(baseline=baseline, no_storage=no_storage, integer_crosscheck=integer_check,
                  sensitivities=variations, savings_yuan=no_storage['cost_yuan']-baseline['cost_yuan'],
                  savings_pct=100*(1-baseline['cost_yuan']/no_storage['cost_yuan']),
                  scipy_version=scipy.__version__, environment=sys.executable)
    (OUT/'validation_summary.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    fig, axs = plt.subplots(3,1,figsize=(11,9),sharex=True,layout='constrained')
    hours = np.arange(144)/6
    axs[0].step(hours,load,where='post',label='Load')
    axs[0].step(hours,pv,where='post',label='PV forecast')
    axs[0].set(ylabel='Power (kW)',title='Q1 baseline: deterministic dispatch (one-way efficiency 90%)')
    axs[1].step(hours,q*6,where='post',label='Grid')
    axs[1].step(hours,c*6,where='post',label='Charge')
    axs[1].step(hours,-d*6,where='post',label='Discharge (negative)')
    axs[1].set(ylabel='Bus power (kW)')
    axs[2].plot(np.arange(145)/6,state,label='Stored energy',color='#1b6c9e')
    axs[2].axhline(1200,ls=':',color='gray')
    axs[2].axhline(10800,ls=':',color='gray')
    axs[2].set(ylabel='Stored energy (kWh)',xlabel='Time of day (h)',xlim=(0,24),xticks=np.arange(0,25,2))
    price_ax = axs[2].twinx()
    price_ax.step(hours,price,where='post',color='#b85c00',alpha=.65,label='Price')
    price_ax.set_ylabel('Price (CNY/kWh)')
    price_ax.legend(loc='upper right')
    for ax in axs:
        ax.grid(alpha=.2)
        ax.legend(loc='upper left',ncol=3)
    fig.savefig(ROOT/'figures/q1_baseline.png',dpi=180)
    plt.close(fig)
    report = [
        '# 问题1 Baseline结果（条件性结果）', '',
        '本结果基于右端点代表功率、母线侧充放电、单向效率均90%、初末储电量均6000 kWh的解释。时间与效率口径仍需保留为假设；不称为无条件最终答案。', '',
        '## 问题分析', '',
        '输入附件1的144个电价、负载和光伏预测样本；输出144段购电量与充放电量及145个储电状态。目标为最小化全天购电费。负载与光伏在本问中视作给定；不含预测误差风险。', '',
        '## 数据预处理', '',
        '原始时间00:10至次日00:00映射至00:00—00:10至23:50—24:00。将kW乘以1/6小时转换为kWh；无缺失插补或异常值删除。点值作为区间代表值是离散化假设，已用周期午夜衔接的梯形功率作敏感性检验，电价保持原区间价。', '',
        '## 模型建立', '',
        r'令 $t=0,\ldots,143$，$\Delta t=1/6\ \mathrm{h}$。$L_t,V_t$为负载、光伏功率（kW）；$p_t$为电价（元/kWh）；$q_t,c_t,d_t,w_t$为购电、母线侧充电、母线侧放电和未使用电量（kWh）；$E_t$为时段起点内部储电量（kWh）。', '',
        r'$$\min \sum_{t=0}^{143}p_tq_t$$', '',
        r'$$q_t+V_t\Delta t+d_t=L_t\Delta t+c_t+w_t,$$', '',
        r'$$E_{t+1}=E_t+0.9c_t-d_t/0.9,$$', '',
        r'$$q_t,c_t,d_t,w_t\ge0,\quad c_t,d_t\le5000\Delta t,\quad1200\le E_t\le10800,\quad E_0=E_{144}=6000.$$', '',
        '不设售电收入；允许未消纳电量，不额外计费。LP先放松充放电互斥，再检查解；整数交叉模型加入二进制状态确保不同时充放电。本例LP解无同时充放电，并与MILP最优费用一致，故可直接采用LP解。', '',
        '## 模型求解', '',
        '使用SciPy的HiGHS求解LP与MILP；无需人为设定初始调度。LP包含721个连续变量、288个等式及变量界；MILP额外增加144个二进制变量及288个互斥不等式。MILP相对gap目标为1e-9，时间上限120秒，运行状态详见验证JSON。此规模无需启发式调参；不对一般MILP作多项式复杂度承诺。', '',
        '| 策略 | 全天购电量（kWh） | 全天购电费（元） |', '|---|---:|---:|',
        f"| 无储能 | {no_storage['grid_kWh']:.6f} | {no_storage['cost_yuan']:.6f} |",
        f"| LP储能调度 | {baseline['grid_kWh']:.6f} | {baseline['cost_yuan']:.6f} |", '',
        f"相对无储能节省{result['savings_yuan']:.6f}元（{result['savings_pct']:.4f}%）。这是该确定性单日场景内的储能价值，不是全年收益预测。", '',
        '| 指定时段 | 购电量（kWh） |','|---|---:|',
    ]
    report += [f'| {r.interval} | {r.grid_kWh:.6f} |' for r in selected.itertuples()]
    report += ['', '| 时间段 | 母线侧充电量（kWh） | 母线侧放电量（kWh） |','|---|---:|---:|']
    report += [f'| {r.interval} | {r.charge_kWh:.6f} | {r.discharge_kWh:.6f} |' for r in groups.itertuples()]
    report += ['',f"0:00和24:00储电量均为{state[0]:.6f} kWh。指定时段数值为保存的一组最优解，未证明调度方案唯一。",'',
               '### 验证', '',
               f"- LP与MILP费用差：{abs(integer_check['cost_yuan']-baseline['cost_yuan']):.3e}元；MILP相对gap：{integer_check['mip_gap']:.3e}。",
               f"- 最大平衡或约束违规：{max(baseline['checks'].values()):.3e} kWh；阈值1e-6 kWh。",
               '- 零充放电功率边界模型与无储能解析公式一致；输出Excel重读与CSV数据一致。',
               '- 此问是确定性调度，不拟合未知参数，因此不报告预测残差或伪造置信区间。', '',
               '| 敏感性情景 | 最优购电费（元） | 相对主模型变化（%） |', '|---|---:|---:|']
    scenario_names = dict(round_trip_efficiency_90pct='往返效率90%（单向各√0.9）',
                          trapezoid_power_periodic_midnight='梯形功率积分（午夜用周期边界）',
                          load_plus_5pct='负载增加5%',load_minus_5pct='负载减少5%',
                          pv_plus_5pct='光伏增加5%',pv_minus_5pct='光伏减少5%')
    report += [f"| {scenario_names[k]} | {v['cost_yuan']:.6f} | {100*(v['cost_yuan']/baseline['cost_yuan']-1):.4f} |" for k,v in variations.items()]
    report += ['', '每个情景重新优化，衡量输入或口径改变后的最优值；不等同于固定原策略面对误差时的鲁棒性。5%是人为设置的灵敏度幅度，不是统计误差估计。', '',
               '### 文件与复现', '',
               '- `results/q1_baseline/result1_baseline.xlsx`：保留模板主要结构，明确修正时段标签，附原标签映射及口径说明，非最终提交版。',
               '- `results/q1_baseline/dispatch_10min.csv`：逐时段完整策略及费用。',
               '- `results/q1_baseline/validation_summary.json`：校验值、对照、敏感性及求解状态。',
               '- `results/q1_baseline/config.json`：集中参数；来源为题面或本报告明确声明的假设。',
               '- `figures/q1_baseline.png`：功率、储电量和电价运行图。',
               '- 在赛题目录依次执行 `conda run -n math_modeling python code/01_audit_inputs.py` 和 `conda run -n math_modeling python code/02_q1_baseline.py`。Windows控制台可设置PYTHONUTF8=1、PYTHONIOENCODING=utf-8，并加--no-capture-output。', '']
    (ROOT/'reports/问题一/问题1_Baseline结果.md').write_text('\n'.join(report),encoding='utf-8')
    print(json.dumps(dict(baseline=baseline,no_storage=no_storage,savings_pct=result['savings_pct'],
                         integer_cost=integer_check['cost_yuan'],
                         sensitivity_costs={k:v['cost_yuan'] for k,v in variations.items()}),indent=2,ensure_ascii=False))


if __name__ == '__main__':
    main()
