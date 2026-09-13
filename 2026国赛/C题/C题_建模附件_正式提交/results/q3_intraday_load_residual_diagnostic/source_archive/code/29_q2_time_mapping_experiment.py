#!/usr/bin/env python
"""问题二 29 号实验：统一新时间映射（start_time_v1）与末态对照。

任务书 ``reports/问题二_统一新时间映射重跑实验方案.md``（单次任务书，主模型维护于问题二建模报告）。

区间起点口径（用户确认，属建模解释而非官方勘误）：

* 源标签 ``00:10`` 代表 **00:10—00:20**，源标签 ``0:00+1`` 代表 **次日 00:00—00:10**；
* 因此一个源行 k 覆盖 ``[k日00:10, (k+1)日00:10)``，是"模板日"而不是自然日；
* 自然日 d（d>=1）由 **前一行末值接当前行前 143 值** 组成，不能逐行 np.roll；
* k 日 0:00 发布覆盖 ``[k日00:10,(k+1)日00:10)`` 的 144 段模板计划，另有 1 个辅助午夜目标；
* 自然日 24:00 才是真正的 24:00：固定组约束 **S143=6000**，S144 两组都只受物理界限制。

两组唯一差别是名义末态：``N_fixed6000``（S143=6000）与 ``N_free``（仅 1200—10800）。
预测（新信息集下重建的 15 列负载 / 12 列光伏 LightGBM）、W28/q80、公共 1 月、实际反馈、
收费规则、求解器设置完全相同；旧 27 号的费用、2 月初态与 1 元豁免一律不复用。

只新增本脚本、``results/q2_time_mapping/``、``figures/q2_time_mapping/`` 与本实验报告；
旧附件、源码、签名结果只读。不填正式 ``result2.xlsx``、不锁定策略、不启动其他 Agent。
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import math
import shutil
import sys
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
from scipy.optimize import milp, Bounds, LinearConstraint
from scipy.sparse import lil_matrix, csr_matrix, hstack, vstack

ROOT = Path(__file__).resolve().parents[1]
CODE_29 = ROOT / 'code/29_q2_time_mapping_experiment.py'
OUT = ROOT / 'results/q2_time_mapping'
FIG = ROOT / 'figures/q2_time_mapping'
PLAN_MD = ROOT / 'reports/问题二_统一新时间映射重跑实验方案.md'
REPORT_MD = ROOT / 'reports/问题二_统一新时间映射实验结果报告.md'
TIME_REPORT = ROOT / 'reports/时间口径变更与跨问题影响_20260912.md'

TIME_VERSION = 'start_time_v1'
DT = 1 / 6
ETA = 0.9
STATE_MIN_KWH = 1200.0
STATE_MAX_KWH = 10800.0
POWER_MAX_KW = 5000.0
TERMINAL_FIXED_KWH = 6000.0
ALPHA = 0.80
RESIDUAL_WINDOW = 28
MIN_HISTORY_DAYS = 7
TRAIN_WINDOW = 56
MIN_TRAIN_DAYS = 14
CONSTANT_TARGET_TOL_KW = 1e-8
ENERGY_TOL_KWH = 1e-6
COST_TOL_YUAN = 1e-4
NATURAL_DAYS = 365
TEMPLATE_SLOTS = 144
TARGETS = 145
WARMUP_LAST_DAY = 30          # the shared pre-run covers natural days 0..30 (Jan 1..31)
EVAL_FIRST_DAY = 31           # 2025-02-01
EVAL_LAST_DAY = 364           # 2025-12-31
TAIL_NATURAL_DAY = 365        # 2026-01-01, only its 00:00-00:10 interval exists
EMERGENCY_MULTIPLIER = 5.0

GROUPS = [
    dict(id='N_fixed6000', mode='fixed6000',
         role='new time mapping, nominal natural-day 24:00 state pinned to 6000 kWh'),
    dict(id='N_free', mode='free',
         role='new time mapping, nominal 24:00 state only bounded by 1200-10800 kWh'),
]
GROUP_BY_ID = {g['id']: g for g in GROUPS}
CONTRAST = dict(label='N_free minus N_fixed6000 (new time mapping, ledger A)',
                treatment='N_free', baseline='N_fixed6000')

PROTECTED_TREES = ['附件', 'code', 'reports', 'figures', 'results', 'outputs']
OWN_NEW_REL = {'code/29_q2_time_mapping_experiment.py',
               'reports/问题二_统一新时间映射实验结果报告.md'}
OWN_NEW_PREFIXES = ('results/q2_time_mapping/', 'figures/q2_time_mapping/')
SHARED_APPEND_REL = {'建模上下文记忆.md', 'reports/项目进度.md'}

LOAD_FEATURES = ['l_base', 'l_b_minus_base', 'l_b_minus_b7', 'l_sameweekday_mean_minus_base',
                 'l_b_mean_minus_b7_mean', 'weekday_tue', 'weekday_wed', 'weekday_thu',
                 'weekday_fri', 'weekday_sat', 'weekday_sun', 'sin1', 'cos1', 'sin2', 'cos2']
PV_FEATURES = ['v_base', 'v_b_minus_b1', 'v_week_mean_minus_base', 'v_b_mean_minus_b1_mean',
               'sin1', 'cos1', 'sin2', 'cos2', 'support_duration_h', 'support_midpoint_h',
               'support_phase', 'support_valid']
SUPPORT_HISTORY_DAYS = 7
SUPPORT_ACTIVITY_KW = 1.0
SUPPORT_PADDING_SLOTS = 3

PARAMETERS = dict(
    time_version=TIME_VERSION,
    experiment='unified interval-start mapping, new public January, retrained 15/12-column '
               'LightGBM under the new information set, two nominal terminal modes',
    interval_rule='label 00:10 covers 00:10-00:20; label 0:00+1 covers the next day 00:00-00:10',
    natural_day_rule='natural day d = previous source row column 143 ++ current row columns 0..142',
    price_rule='natural-day price = source row 143 value then source rows 0..142',
    template_rule='k-day 0:00 publishes 144 slots covering [k 00:10, (k+1) 00:10); the current '
                  'midnight is executed from the previous day carry commitment and is never replanned',
    targets=dict(count=TARGETS, mapping='h=0 current 00:00; h=1..143 current 00:10..23:50; '
                                        'h=144 next 00:00; template slot j = h-1'),
    midnight_target_rule='both modes keep S0..S144 physically bounded; the fixed mode pins S143 '
                         '(true natural-day 24:00); nothing pins S144',
    alpha=ALPHA, residual_window_days=RESIDUAL_WINDOW, min_history_days=MIN_HISTORY_DAYS,
    train_window_days=TRAIN_WINDOW, min_train_days=MIN_TRAIN_DAYS,
    quantile_rule='per-target empirical inverse distribution, 1-based position (m*a+99)//100, index '
                  'position-1, no interpolation, negative corrections retained; h=144 has m=27 '
                  'position 22, all other targets m=28 position 23',
    load_features=LOAD_FEATURES, pv_features=PV_FEATURES,
    support_parameters=dict(history_days=SUPPORT_HISTORY_DAYS,
                            activity_threshold_kW=SUPPORT_ACTIVITY_KW,
                            padding_slots=SUPPORT_PADDING_SLOTS),
    solver='local transcription of the frozen integer kernel; relative gap 1e-9, time limit 120 s; '
           'the only difference between the two modes is the S143 bound',
    ledgers=dict(A='natural days 2025-02-01..2025-12-31, 334 days, 48096 segments',
                 B='template publications 2025-02-01..2025-12-31, 48096 segments, shifted by '
                   '10 minutes',
                 bridge='C_B - C_A = cost(tail 2026-01-01 00:00-00:10) - cost(head 2025-02-01 '
                        '00:00-00:10)'),
    reproduction='new-run internal cold-start replay of N_fixed6000; no pre-registered '
                 'new-convention cost constant exists and the old 27-module 1 CNY exemption is '
                 'not reused',
    groups=[dict(id=g['id'], mode=g['mode']) for g in GROUPS],
    comparison=CONTRAST['label'],
)

_MODULES: dict = {}


def _load(name, rel):
    spec = importlib.util.spec_from_file_location(name, ROOT / rel)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def parent():
    """Module 14: protected-tree list, figure integrity and the frozen snapshot view."""
    if 'ridge' not in _MODULES:
        _MODULES['ridge'] = _load('ridge_parent', 'code/14_q2_ridge_forecast_experiment.py')
    return _MODULES['ridge']


def scan():
    """Module 21: the registered integer quantile-position rule."""
    if 'scan' not in _MODULES:
        module = _load('quantile_scan', 'code/21_q2_quantile_level_scan.py')
        module.OUT = OUT
        module.FIG = FIG
        _MODULES['scan'] = module
    return _MODULES['scan']


def lgbm_module():
    """Module 19: the frozen LightGBM parameter dictionary."""
    if 'lgbm' not in _MODULES:
        module = _load('lightgbm_residual', 'code/19_q2_lightgbm_residual_experiment.py')
        module.OUT = OUT
        module.FIG = FIG
        _MODULES['lgbm'] = module
    return _MODULES['lgbm']


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def jsonable(obj):
    if isinstance(obj, dict):
        return {str(k): jsonable(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [jsonable(v) for v in obj]
    if isinstance(obj, np.ndarray):
        return obj.tolist()
    if isinstance(obj, np.integer):
        return int(obj)
    if isinstance(obj, np.floating):
        value = float(obj)
        return value if math.isfinite(value) else None
    if isinstance(obj, np.bool_):
        return bool(obj)
    if isinstance(obj, float):
        return obj if math.isfinite(obj) else None
    if isinstance(obj, (Path, pd.Timestamp, datetime, timedelta)):
        return str(obj)
    return obj


def save(path, obj):
    Path(path).write_text(json.dumps(jsonable(obj), ensure_ascii=False, indent=2, default=str),
                          encoding='utf-8')


def frame_to_csv(frame, path):
    frame.to_csv(path, index=False, encoding='utf-8-sig')


def protected_manifest():
    manifest = {}
    for tree in PROTECTED_TREES:
        base = ROOT / tree
        if not base.exists():
            continue
        for path in sorted(base.rglob('*')):
            if not path.is_file() or '__pycache__' in path.parts or path.name.startswith('~$'):
                continue
            relative = path.relative_to(ROOT).as_posix()
            if relative in OWN_NEW_REL or relative in SHARED_APPEND_REL:
                continue
            if relative.startswith(OWN_NEW_PREFIXES):
                continue
            manifest[relative] = digest(path)
    return manifest


# ======================================================================================
# 1. interval-start mapping: source rows -> long table -> natural days -> 145 targets
# ======================================================================================
def label_text(value):
    """Attachment time labels are datetime.time except the final '0:00+1' string."""
    if hasattr(value, 'strftime'):
        return value.strftime('%H:%M')
    text = str(value)
    return '0:00+1' if '+1' in text else text[:5]


def read_attachments():
    """Read the original attachments without changing a single stored value."""
    wb = openpyxl.load_workbook(ROOT / '附件/附件1.xlsx', read_only=True, data_only=True)
    raw1 = list(wb.worksheets[0].values)
    wb.close()
    assert len(raw1) - 1 == TEMPLATE_SLOTS, len(raw1)
    price_src = np.array([row[1] for row in raw1[1:]], dtype=float)
    price_labels = [label_text(row[0]) for row in raw1[1:]]
    wb = openpyxl.load_workbook(ROOT / '附件/附件2.xlsx', read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        sheets.append((ws.title, list(ws.values)))
    wb.close()
    assert len(sheets) == 2, [name for name, _ in sheets]
    load_rows, pv_rows = sheets[0][1], sheets[1][1]
    assert len(load_rows) - 1 == NATURAL_DAYS and len(pv_rows) - 1 == NATURAL_DAYS
    source_dates = [row[0] for row in load_rows[1:]]
    assert [row[0] for row in pv_rows[1:]] == source_dates
    load_src = np.asarray([row[1:] for row in load_rows[1:]], dtype=float)
    pv_src = np.asarray([row[1:] for row in pv_rows[1:]], dtype=float)
    assert load_src.shape == pv_src.shape == (NATURAL_DAYS, TEMPLATE_SLOTS)
    assert price_src.shape == (TEMPLATE_SLOTS,)
    assert np.isfinite(load_src).all() and np.isfinite(pv_src).all() and np.isfinite(price_src).all()
    assert (load_src >= 0).all() and (pv_src >= 0).all() and (price_src > 0).all()
    labels = [label_text(cell) for cell in load_rows[0][1:]]
    assert len(labels) == TEMPLATE_SLOTS and labels[-1] == '0:00+1', labels[-1]
    assert [label_text(cell) for cell in pv_rows[0][1:]] == labels, 'load and pv headers differ'
    return dict(price_src=price_src, price_labels=price_labels, load_src=load_src, pv_src=pv_src,
                source_dates=source_dates, labels=labels,
                sheets=[name for name, _ in sheets])


def natural_arrays(load_src, pv_src):
    """natural day d slot 0 = previous row's column 143; slots 1..143 = row d columns 0..142."""
    nat_load = np.full((NATURAL_DAYS + 1, TEMPLATE_SLOTS), np.nan)
    nat_pv = np.full((NATURAL_DAYS + 1, TEMPLATE_SLOTS), np.nan)
    for d in range(NATURAL_DAYS):
        nat_load[d, 1:] = load_src[d, 0:TEMPLATE_SLOTS - 1]
        nat_pv[d, 1:] = pv_src[d, 0:TEMPLATE_SLOTS - 1]
        nat_load[d + 1, 0] = load_src[d, TEMPLATE_SLOTS - 1]
        nat_pv[d + 1, 0] = pv_src[d, TEMPLATE_SLOTS - 1]
    # natural day 0 has no observed 00:00-00:10: slot 0 stays NaN (the gap), never faked as zero
    nat_load[0, 0] = np.nan
    nat_pv[0, 0] = np.nan
    return nat_load, nat_pv


def natural_price(price_src):
    """Natural-day price = source last-row value then source rows 0..142 (daily clock pattern)."""
    return np.r_[price_src[TEMPLATE_SLOTS - 1], price_src[:TEMPLATE_SLOTS - 1]]


def truth_targets(load_src, pv_src, k):
    """The 145 realised targets of publication day k: previous row's column 143 ++ row k."""
    assert k >= 1, 'day 0 has no previous source row'
    return np.r_[load_src[k - 1, TEMPLATE_SLOTS - 1], load_src[k, :]], \
        np.r_[pv_src[k - 1, TEMPLATE_SLOTS - 1], pv_src[k, :]]


def build_long_table(attachments, dates):
    """One row per stored power value and per price value, with both coordinates."""
    load_src, pv_src, price_src = (attachments['load_src'], attachments['pv_src'],
                                   attachments['price_src'])
    rows = []
    for sheet, array in (('小区负载', load_src), ('光伏发电实际功率', pv_src)):
        for r in range(NATURAL_DAYS):
            source_date = dates[r]
            for j in range(TEMPLATE_SLOTS):
                if j < TEMPLATE_SLOTS - 1:
                    start = source_date + timedelta(minutes=10 + 10 * j)
                    natural_date = source_date
                    natural_slot = j + 1
                else:
                    start = source_date + timedelta(days=1)
                    natural_date = start
                    natural_slot = 0
                end = start + timedelta(minutes=10)
                rows.append(dict(source_file='附件/附件2.xlsx', source_sheet=sheet,
                                 source_row=r + 2, source_column=j + 2,
                                 source_date=str(source_date.date()),
                                 source_label=attachments['labels'][j],
                                 interval_start=start.isoformat(sep=' '),
                                 interval_end=end.isoformat(sep=' '),
                                 available_at=end.isoformat(sep=' '),
                                 natural_date=str(natural_date.date()),
                                 natural_slot=natural_slot,
                                 template_date=str(source_date.date()), template_slot=j,
                                 value_kW=float(array[r, j])))
    # price: 144 stored values, one daily clock pattern; they are not copied into 365 days
    for j in range(TEMPLATE_SLOTS):
        rows.append(dict(source_file='附件/附件1.xlsx', source_sheet='sheet1', source_row=j + 2,
                         source_column=2, source_date='daily-clock-pattern',
                         source_label=attachments['price_labels'][j], interval_start='',
                         interval_end='', available_at='', natural_date='daily-clock-pattern',
                         natural_slot=(0 if j == TEMPLATE_SLOTS - 1 else j + 1),
                         template_date='daily-clock-pattern',
                         template_slot=(TEMPLATE_SLOTS - 1 if j == TEMPLATE_SLOTS - 1 else j - 1),
                         value_kW=float(price_src[j])))
    return pd.DataFrame(rows)


def mapping_checks(attachments, dates, nat_load, nat_pv, price_src):
    """Reversibility, no-overlap/no-gap coverage, and the two-coordinate bookkeeping."""
    load_src, pv_src = attachments['load_src'], attachments['pv_src']
    rounds = []
    worst = 0.0
    for r in (0, 1, 100, 200, 364):
        for j in (0, 1, 71, 142):
            worst = max(worst, abs(load_src[r, j] - nat_load[r, j + 1]),
                        abs(pv_src[r, j] - nat_pv[r, j + 1]))
        worst = max(worst, abs(load_src[r, TEMPLATE_SLOTS - 1] - nat_load[r + 1, 0]),
                    abs(pv_src[r, TEMPLATE_SLOTS - 1] - nat_pv[r + 1, 0]))
    rounds.append(dict(check='row-to-natural bijection', worst=worst, tolerance=0.0,
                       passed=bool(worst == 0.0)))
    # natural day d must contain exactly the 144 stored values of the matching physical intervals
    total = 0
    for d in range(0, NATURAL_DAYS + 1):
        total += int(np.isfinite(nat_load[d]).sum() + np.isfinite(nat_pv[d]).sum())
    rounds.append(dict(check='natural-day value count',
                       worst=float(abs(total - 2 * NATURAL_DAYS * TEMPLATE_SLOTS)),
                       tolerance=0.0,
                       passed=bool(total == 2 * NATURAL_DAYS * TEMPLATE_SLOTS)))
    # 145-target identity against the natural arrays
    gap = 0.0
    for k in (1, 100, 250, 364):
        t_load, t_pv = truth_targets(load_src, pv_src, k)
        gap = max(gap, abs(t_load[0] - nat_load[k, 0]), abs(t_pv[0] - nat_pv[k, 0]),
                  float(np.max(np.abs(t_load[1:TEMPLATE_SLOTS] - nat_load[k, 1:TEMPLATE_SLOTS]))),
                  float(np.max(np.abs(t_pv[1:TEMPLATE_SLOTS] - nat_pv[k, 1:TEMPLATE_SLOTS]))),
                  abs(t_load[TEMPLATE_SLOTS] - nat_load[k + 1, 0]),
                  abs(t_pv[TEMPLATE_SLOTS] - nat_pv[k + 1, 0]))
    rounds.append(dict(check='145-target identity', worst=gap, tolerance=0.0, passed=bool(gap == 0.0)))
    # price order and conservation
    price_nat = natural_price(price_src)
    order_error = max(abs(price_nat[0] - price_src[TEMPLATE_SLOTS - 1]),
                      float(np.max(np.abs(price_nat[1:] - price_src[:TEMPLATE_SLOTS - 1]))))
    rounds.append(dict(check='natural price = last row then rows 0..142',
                       worst=max(order_error, abs(price_nat.sum() - price_src.sum())),
                       tolerance=0.0,
                       passed=bool(order_error == 0.0
                                   and abs(price_nat.sum() - price_src.sum()) < 1e-9)))
    # the template objective price equals the stored source order
    obj_error = float(np.max(np.abs(price_nat[1:TEMPLATE_SLOTS] - price_src[:TEMPLATE_SLOTS - 1])))
    rounds.append(dict(check='template objective price order', worst=obj_error, tolerance=0.0,
                       passed=bool(obj_error == 0.0)))
    # interval coverage: 52560 power values, no overlap, no gap, monotone 10-minute steps
    covered = np.array([(dates[0] + timedelta(days=r, minutes=10 + 10 * j))
                        for r in range(NATURAL_DAYS) for j in range(TEMPLATE_SLOTS)])
    diffs = np.diff(covered.astype('datetime64[m]').astype(np.int64))
    rounds.append(dict(check='source intervals are contiguous 10-minute steps',
                       worst=float(np.max(np.abs(diffs - 10))), tolerance=0.0,
                       passed=bool(int(np.max(np.abs(diffs - 10))) == 0
                                   and len(covered) == 52560)))
    out = dict(rounds=rounds, source_interval_count=int(len(covered)),
               first_interval=str(covered[0]), last_interval=str(covered[-1]),
               gap_slot_is_nan=bool(np.isnan(nat_load[0, 0]) and np.isnan(nat_pv[0, 0])),
               values_preserved=int(load_src.size + pv_src.size + price_src.size),
               all_passed=bool(all(r['passed'] for r in rounds)))
    return out


# ======================================================================================
# 2. features, LightGBM training and the published forecast archive
# ======================================================================================
def weekday_one_hot(a_dates, a):
    weekday = a_dates[a].weekday()
    return np.array([1.0 if weekday == d else 0.0 for d in range(1, 7)])


def harmonics():
    t = np.arange(TEMPLATE_SLOTS, dtype=float)
    return np.column_stack([np.sin(2 * np.pi * t / 144), np.cos(2 * np.pi * t / 144),
                            np.sin(4 * np.pi * t / 144), np.cos(4 * np.pi * t / 144)])


HARMONICS = harmonics()

_FEATURE_CACHE: dict = {}


def naive_load_day(nat_load, dates, a):
    """b^L: most recent complete same-weekday day before a; fall back to b = a-1 semantics."""
    offset = 7
    day = a - offset
    while day >= 1:
        if dates[day].weekday() == dates[a].weekday():
            return nat_load[day].copy()
        day -= offset
    return nat_load[a - 1].copy()


def naive_pv_day(nat_pv, k):
    return nat_pv[k - 1].copy()


def naive_rows_145(nat_load, nat_pv, dates, k, kind):
    """The naive baseline as 145 targets, without touching any feature column."""
    if kind == 'load':
        base_cur = naive_load_day(nat_load, dates, k)
        base_next = (naive_load_day(nat_load, dates, min(k + 1, NATURAL_DAYS))
                     if k + 1 <= NATURAL_DAYS else nat_load[k - 1].copy())
        return np.r_[base_cur[0], base_cur[1:TEMPLATE_SLOTS], base_next[0]]
    base = naive_pv_day(nat_pv, k)
    return np.r_[base[0], base[1:TEMPLATE_SLOTS], base[0]]


def support_window(nat_pv, k):
    """s, e, valid from the seven complete days b-6..b; empty or too little history -> full day."""
    b = k - 1
    history = [d for d in range(max(1, b - SUPPORT_HISTORY_DAYS + 1), b + 1)]
    if len(history) < SUPPORT_HISTORY_DAYS:
        return 0, TEMPLATE_SLOTS - 1, 0
    block = nat_pv[history]
    active = np.flatnonzero(block.max(axis=0) > SUPPORT_ACTIVITY_KW)
    if active.size == 0:
        return 0, TEMPLATE_SLOTS - 1, 0
    start = max(0, int(active.min()) - SUPPORT_PADDING_SLOTS)
    stop = min(TEMPLATE_SLOTS - 1, int(active.max()) + SUPPORT_PADDING_SLOTS)
    return start, stop, 1


def load_features_day(nat_load, dates, k, a):
    """144 x 15 feature rows indexed by natural clock slot t, using only history known at k 00:00."""
    b = k - 1
    base = naive_load_day(nat_load, dates, a)
    out = np.empty((TEMPLATE_SLOTS, len(LOAD_FEATURES)))
    out[:, 0] = base
    out[:, 1] = nat_load[b] - base
    out[:, 2] = nat_load[b] - nat_load[b - 7]
    ids = [r for r in (1, 2, 3, 4) if a - 7 * r >= 1]
    out[:, 3] = nat_load[[a - 7 * r for r in ids]].mean(axis=0) - base
    out[:, 4] = nat_load[b].mean() - nat_load[b - 7].mean()
    out[:, 5:11] = weekday_one_hot(dates, a)
    out[:, 11:15] = HARMONICS
    return out, base


def pv_features_day(nat_pv, k):
    """144 x 12 feature rows indexed by natural clock slot t."""
    b = k - 1
    base = naive_pv_day(nat_pv, k)
    out = np.empty((TEMPLATE_SLOTS, len(PV_FEATURES)))
    out[:, 0] = base
    out[:, 1] = nat_pv[b] - nat_pv[b - 1]
    out[:, 2] = nat_pv[max(1, b - 6):b + 1].mean(axis=0) - base
    out[:, 3] = nat_pv[b].mean() - nat_pv[b - 1].mean()
    out[:, 4:8] = HARMONICS
    start, stop, valid = support_window(nat_pv, k)
    duration = (stop - start + 1) / 6.0
    midpoint = (start + stop + 1) / 12.0
    phase = (np.arange(TEMPLATE_SLOTS) - start + 0.5) / (stop - start + 1)
    out[:, 8] = duration
    out[:, 9] = midpoint
    out[:, 10] = phase
    out[:, 11] = float(valid)
    gate = (np.arange(TEMPLATE_SLOTS) >= start) & (np.arange(TEMPLATE_SLOTS) <= stop)
    return out, base, gate.astype(float), (start, stop, valid)


def first_feature_day(kind):
    """Load needs b-7 >= 1 (k >= 9); PV needs b-6 >= 1 (k >= 8)."""
    return 9 if kind == 'load' else 8


def feature_rows_145(nat_load, nat_pv, dates, k, kind):
    """145 x F design rows for one publication day; row h uses its own target date and clock.

    Rows: h=0 uses target date k at clock 0; h=1..143 use target date k at clock h;
    h=144 uses target date k+1 at clock 0 (the next midnight). Everything else uses b = k-1.
    """
    key = (kind, k)
    if key in _FEATURE_CACHE:
        return _FEATURE_CACHE[key]
    if kind == 'load':
        feat_cur, base_cur = load_features_day(nat_load, dates, k, k)
        feat_next, base_next = load_features_day(nat_load, dates, k, min(k + 1, NATURAL_DAYS))
        rows = np.vstack([feat_cur[0:1], feat_cur[1:TEMPLATE_SLOTS], feat_next[0:1]])
        base = np.r_[base_cur[0], base_cur[1:TEMPLATE_SLOTS], base_next[0]]
        result = (rows, base, None, None)
    else:
        feat_cur, base_cur, gate_cur, window = pv_features_day(nat_pv, k)
        rows = np.vstack([feat_cur[0:1], feat_cur[1:TEMPLATE_SLOTS], feat_cur[0:1]])
        base = np.r_[base_cur[0], base_cur[1:TEMPLATE_SLOTS], base_cur[0]]
        gate = np.r_[gate_cur[0], gate_cur[1:TEMPLATE_SLOTS], gate_cur[0]]
        result = (rows, base, gate, window)
    _FEATURE_CACHE[key] = result
    return result


def label_rows_144(nat_load, nat_pv, dates, j, kind, load_src, pv_src):
    """Training rows for publication day j: its 144 template targets (h=1..144)."""
    features, base, gate, window = feature_rows_145(nat_load, nat_pv, dates, j, kind)
    truth_load, truth_pv = truth_targets(load_src, pv_src, j)
    truth = truth_load if kind == 'load' else truth_pv
    return features[1:TARGETS], truth[1:TARGETS] - base[1:TARGETS]


def build_forecast_archive(nat_load, nat_pv, dates, load_src, pv_src, save_models=False,
                           verbose=False, day_from=1, model_days=None, day_to=NATURAL_DAYS):
    """One LightGBM per target per publication day under the new information set.

    Day k trains on the feature-valid, fully-labelled publication days j <= k-2 inside the trailing
    56-day window, needs at least 14 of them, and falls back to the naive baseline otherwise.
    """
    from lightgbm import LGBMRegressor

    params = dict(lgbm_module().LGBM_PARAMS)
    issued_load = np.full((NATURAL_DAYS, TARGETS), np.nan)
    issued_pv = np.full((NATURAL_DAYS, TARGETS), np.nan)
    naive_load = np.full((NATURAL_DAYS, TARGETS), np.nan)
    naive_pv = np.full((NATURAL_DAYS, TARGETS), np.nan)
    gate_archive = np.ones((NATURAL_DAYS, TARGETS))
    support_archive = np.full((NATURAL_DAYS, 3), np.nan)
    fit_log = []
    models_dir = OUT / 'models'
    if save_models:
        models_dir.mkdir(parents=True, exist_ok=True)

    for k in range(max(1, int(day_from)), min(int(day_to), NATURAL_DAYS)):
        for kind in ('load', 'pv'):
            threshold = first_feature_day(kind)
            entry = dict(target=kind, day_index=k, date=str(dates[k].date()), trained=False,
                         branch='', n_train_days=0, n_rows=0, model_file='', elapsed_seconds=0.0)
            if k < threshold:
                # before the first feature-valid day the plan still publishes the causal naive
                # baseline; no feature column is read, so no window can reach before day 1
                base = naive_rows_145(nat_load, nat_pv, dates, k, kind)
                gate = np.ones(TARGETS)
                window = (0, TEMPLATE_SLOTS - 1, 0)
                issued = base
                entry['branch'] = 'feature_history_fallback'
                entry['reason'] = f'day_index<{threshold}'
            else:
                features, base, gate, window = feature_rows_145(nat_load, nat_pv, dates, k, kind)
                issued = base
                train_days = [j for j in range(max(1, k - TRAIN_WINDOW), k - 1) if j >= threshold]
                entry['n_train_days'] = len(train_days)
                if len(train_days) < MIN_TRAIN_DAYS:
                    entry['branch'] = 'insufficient_history'
                    entry['reason'] = f'train_days={len(train_days)}<{MIN_TRAIN_DAYS}'
                else:
                    rows_sets = [label_rows_144(nat_load, nat_pv, dates, j, kind,
                                                load_src, pv_src) for j in train_days]
                    X = np.concatenate([rows for rows, _ in rows_sets], axis=0)
                    y = np.concatenate([labels for _, labels in rows_sets], axis=0)
                    if not (np.isfinite(X).all() and np.isfinite(y).all()):
                        raise RuntimeError(f'{kind} day {k}: non-finite feature or label')
                    entry['n_rows'] = int(X.shape[0])
                    entry['train_first'] = str(dates[train_days[0]].date())
                    entry['train_last'] = str(dates[train_days[-1]].date())
                    started = time.perf_counter()
                    if float(np.max(y) - np.min(y)) <= CONSTANT_TARGET_TOL_KW:
                        correction = np.full(TARGETS, float(np.mean(y)))
                        entry['branch'] = 'constant_target'
                        entry['constant_value_kW'] = float(np.mean(y))
                    else:
                        model = LGBMRegressor(**params)
                        model.fit(X, y)
                        correction = np.asarray(model.predict(features), dtype=float)
                        entry['branch'] = 'fitted'
                        entry['n_trees'] = int(model.booster_.num_trees())
                        if save_models and (model_days is None or k in model_days):
                            name = f'{kind}_{dates[k].date()}.txt'
                            text = model.booster_.model_to_string()
                            (models_dir / name).write_text(text, encoding='utf-8')
                            entry['model_file'] = f'models/{name}'
                    entry['elapsed_seconds'] = time.perf_counter() - started
                    issued = base + correction
                    entry['trained'] = True
            published = np.maximum(0.0, issued)
            if kind == 'load':
                naive_load[k] = base
                issued_load[k] = published
            else:
                naive_pv[k] = base
                gate_archive[k] = gate
                support_archive[k] = window
                issued_pv[k] = published * gate_archive[k]
            fit_log.append(entry)
        if verbose and k % 90 == 0:
            print(f'    forecast archive: day {k} ({dates[k].date()})', flush=True)



    formal = slice(EVAL_FIRST_DAY, NATURAL_DAYS)
    summary = dict(
        load_first_feature_day=first_feature_day('load'),
        pv_first_feature_day=first_feature_day('pv'),
        load_first_trainable_day=int(next((k for k in range(1, NATURAL_DAYS)
                                           if any(e['target'] == 'load' and e['day_index'] == k
                                                  and e['trained'] for e in fit_log)), -1)),
        pv_first_trainable_day=int(next((k for k in range(1, NATURAL_DAYS)
                                         if any(e['target'] == 'pv' and e['day_index'] == k
                                                and e['trained'] for e in fit_log)), -1)),
        load_trained_days=int(sum(1 for e in fit_log if e['target'] == 'load' and e['trained'])),
        pv_trained_days=int(sum(1 for e in fit_log if e['target'] == 'pv' and e['trained'])),
        formal_finite=bool(np.isfinite(issued_load[formal]).all()
                           and np.isfinite(issued_pv[formal]).all()),
        formal_nonnegative=bool((issued_load[formal] >= 0).all()
                                and (issued_pv[formal] >= 0).all()),
        gated_off_slots=int((gate_archive[formal] == 0).sum()),
        support_valid_days=int((support_archive[formal][:, 2] == 1).sum()),
        support_invalid_days=int((support_archive[formal][:, 2] == 0).sum()),
    )
    return dict(issued_load=issued_load, issued_pv=issued_pv, naive_load=naive_load,
                naive_pv=naive_pv, gate=gate_archive, support=support_archive, fit_log=fit_log,
                summary=summary)


def build_net_forecast(issued_load, issued_pv):
    return (issued_load - issued_pv) * DT


def errors_145(load_src, pv_src, nat_load, nat_pv, net_forecast):
    """Per (publication day, target) realised net-demand error, NaN where not issued."""
    errors = np.full((NATURAL_DAYS, TARGETS), np.nan)
    for k in range(1, NATURAL_DAYS):
        truth_load, truth_pv = truth_targets(load_src, pv_src, k)
        errors[k] = (truth_load - truth_pv) * DT - net_forecast[k]
    return errors


def target_available(j, h, k):
    """interval_end(publication j, target h) <= publication k 00:00."""
    if h <= TEMPLATE_SLOTS - 1:
        return True                       # ends by j 23:50 for h<=143
    return j + 1 < k                      # h=144 ends at (j+1) 00:10


def protection_for(errors, k, h):
    """Registered integer-rule quantile protection for one (publication day, target)."""
    low = max(1, k - RESIDUAL_WINDOW)
    candidates = [j for j in range(low, k)
                  if np.isfinite(errors[j, h]) and target_available(j, h, k)]
    m = len(candidates)
    if m < MIN_HISTORY_DAYS:
        return np.zeros(1)[0], m, f'insufficient_history(m={m}<{MIN_HISTORY_DAYS})'
    position = scan().quantile_index(m, ALPHA)
    ordered = np.sort(np.array([errors[j, h] for j in candidates]))
    return float(ordered[position - 1]), m, ''


def protection_archive(errors):
    """Protection amount, sample size and reason for every (publication day, target)."""
    protection = np.zeros((NATURAL_DAYS, TARGETS))
    m_archive = np.zeros((NATURAL_DAYS, TARGETS), dtype=int)
    reasons = {}
    for k in range(WARMUP_LAST_DAY + 1, NATURAL_DAYS):
        for h in range(TARGETS):
            value, m, reason = protection_for(errors, k, h)
            protection[k, h] = value
            m_archive[k, h] = m
            if reason:
                reasons[f'{k}:{h}'] = reason
    return protection, m_archive, reasons


# ======================================================================================
# 3. nominal MILP with the true 24:00 terminal, and the actual feedback
# ======================================================================================
def solver_options():
    return dict(mip_rel_gap=1e-9, time_limit=120)


def build_model(protected_kwh, price_day, initial, mode):
    """144 template slots; S0 = the estimated state at 00:10, terminal index 143 = true 24:00.

    Variable order: q[144], c[144], d[144], w[144], S[0..144], z[144] binary. The two modes differ
    only in the S143 lower/upper bound; S144 stays physically bounded in both.
    """
    assert mode in ('fixed6000', 'free'), mode
    n = TEMPLATE_SLOTS
    n_vars = 5 * n + 1
    cap = POWER_MAX_KW * DT
    objective = np.zeros(n_vars)
    objective[:n] = price_day
    balance = lil_matrix((2 * n, n_vars))
    rhs = np.r_[np.asarray(protected_kwh, dtype=float), np.zeros(n)]
    for t in range(n):
        balance[t, t], balance[t, n + t] = 1, -1
        balance[t, 2 * n + t], balance[t, 3 * n + t] = 1, -1
        balance[n + t, 4 * n + t + 1], balance[n + t, 4 * n + t] = 1, -1
        balance[n + t, n + t], balance[n + t, 2 * n + t] = -ETA, 1 / ETA
    balance = csr_matrix(balance)
    lower, upper = np.zeros(n_vars), np.full(n_vars, np.inf)
    upper[n:3 * n] = cap
    lower[4 * n:], upper[4 * n:] = STATE_MIN_KWH, STATE_MAX_KWH
    lower[4 * n] = upper[4 * n] = float(initial)
    terminal_index = 4 * n + (TEMPLATE_SLOTS - 1)
    if mode == 'fixed6000':
        lower[terminal_index] = upper[terminal_index] = TERMINAL_FIXED_KWH
        target = TERMINAL_FIXED_KWH
    else:
        target = None
    constraints = hstack([balance, csr_matrix((2 * n, n))], format='csr')
    mutex = lil_matrix((2 * n, n_vars + n))
    for t in range(n):
        mutex[t, n + t], mutex[t, n_vars + t] = 1, -cap
        mutex[n + t, 2 * n + t], mutex[n + t, n_vars + t] = 1, cap
    constraints = vstack([constraints, csr_matrix(mutex)], format='csr')
    integrality = np.r_[np.zeros(n_vars), np.ones(n)]
    bounds = Bounds(np.r_[lower, np.zeros(n)], np.r_[upper, np.ones(n)])
    linear = LinearConstraint(constraints, np.r_[rhs, np.full(2 * n, -np.inf)],
                              np.r_[rhs, np.zeros(n), np.full(n, cap)])
    return dict(n=n, n_vars=n_vars, cap=float(cap), objective=np.r_[objective, np.zeros(n)],
                integrality=integrality, bounds=bounds, constraints=linear, lower=lower,
                upper=upper, mode=mode, terminal_index=int(terminal_index),
                terminal_target=target,
                terminal_lower_kWh=float(lower[terminal_index]),
                terminal_upper_kWh=float(upper[terminal_index]),
                variable_order='q[0..143], c[0..143], d[0..143], w[0..143], S[0..144], z[0..143]')


def solve_day(protected_kwh, price_day, initial, mode):
    model = build_model(protected_kwh, price_day, initial, mode)
    started = time.perf_counter()
    result = milp(model['objective'], integrality=model['integrality'], bounds=model['bounds'],
                  constraints=model['constraints'], options=solver_options())
    elapsed = time.perf_counter() - started
    if not result.success:
        raise RuntimeError(f'HiGHS: {result.message}')
    n = model['n']
    x = result.x[:model['n_vars']]
    q, c, d, w = x[:n], x[n:2 * n], x[2 * n:3 * n], x[3 * n:4 * n]
    state = x[4 * n:5 * n + 1]
    z = result.x[model['n_vars']:]
    protected = np.asarray(protected_kwh, dtype=float)
    checks = dict(
        bus_balance_max_abs_kWh=float(np.max(np.abs(q + d - protected - c - w))),
        battery_balance_max_abs_kWh=float(np.max(np.abs(np.diff(state) - ETA * c + d / ETA))),
        first_state_error_kWh=float(abs(state[0] - float(initial))),
        terminal_target_error_kWh=float(abs(state[TEMPLATE_SLOTS - 1] - TERMINAL_FIXED_KWH)
                                        if mode == 'fixed6000' else 0.0),
        state_bound_violation_kWh=float(max(0.0, STATE_MIN_KWH - state.min(),
                                            state.max() - STATE_MAX_KWH)),
        flow_bound_violation_kWh=float(max(0.0, -min(q.min(), c.min(), d.min(), w.min()),
                                           c.max() - model['cap'], d.max() - model['cap'])),
        simultaneous_charge_discharge_kWh=float(np.minimum(c, d).max()),
        binary_integrality_error=float(np.max(np.abs(z - np.rint(z)))),
        binary_gate_violation_kWh=float(max(0.0, np.max(c - model['cap'] * z),
                                            np.max(d - model['cap'] * (1 - z)))))
    if max(checks.values()) >= ENERGY_TOL_KWH:
        raise AssertionError(dict(mode=mode, checks=checks))
    summary = dict(cost_yuan=float(price_day @ q), elapsed_seconds=elapsed, checks=checks,
                   mip_gap=float(getattr(result, 'mip_gap', 0.0)),
                   mip_node_count=int(result.mip_node_count),
                   mip_dual_bound_yuan=float(result.mip_dual_bound),
                   objective_bound_gap_yuan=float(result.fun - result.mip_dual_bound),
                   mode=mode, terminal_target=model['terminal_target'],
                   terminal_lower_kWh=model['terminal_lower_kWh'],
                   terminal_upper_kWh=model['terminal_upper_kWh'],
                   s143_kWh=float(state[TEMPLATE_SLOTS - 1]), s144_kWh=float(state[-1]))
    nominal = dict(charge=np.array(c), discharge=np.array(d), unused=np.array(w),
                   binary=np.asarray(z, dtype=float), state=np.array(state))
    return q, c, d, w, state, summary, nominal


def feedback_step(u, state, battery=True):
    """One step of the frozen feedback in kWh: surge charges, deficit discharges, rest emergency."""
    cap = POWER_MAX_KW * DT if battery else 0.0
    if u >= 0:
        charge = min(u, cap, max(0.0, (STATE_MAX_KWH - state) / ETA))
        unused = u - charge
        return charge, 0.0, 0.0, unused, state + ETA * charge
    discharge = min(-u, cap, max(0.0, (state - STATE_MIN_KWH) * ETA))
    emergency = -u - discharge
    return 0.0, discharge, emergency, 0.0, state - discharge / ETA


def feedback(plan_kwh, net_kwh, initial, battery=True):
    plan = np.asarray(plan_kwh, dtype=float)
    net = np.asarray(net_kwh, dtype=float)
    c, d, e, w = [np.zeros(len(plan)) for _ in range(4)]
    states = np.zeros(len(plan) + 1)
    states[0] = float(initial)
    for t in range(len(plan)):
        c[t], d[t], e[t], w[t], states[t + 1] = feedback_step(
            plan[t] - net[t], states[t], battery=battery)
    return c, d, e, w, states


# ======================================================================================
# 4. day loop: carry commitment, plans, both ledgers
# ======================================================================================
def plan_for_day(k, mode, carry, initial, net_forecast, protection, price_src, dates_k=''):
    """One publication day: zeros on Jan 1-2, naive in the January pre-run, issued+q80 after."""
    if k <= 1:
        plan = np.zeros(TEMPLATE_SLOTS)
        nominal_state = np.full(TEMPLATE_SLOTS + 1, float(initial))
        solver = dict(cost_yuan=0.0, elapsed_seconds=0.0, mip_gap=0.0, checks={}, mode=mode,
                      terminal_target=None, terminal_lower_kWh=float('nan'),
                      terminal_upper_kWh=float('nan'), s143_kWh=float(initial),
                      s144_kWh=float(initial), objective_bound_gap_yuan=float('nan'))
        return plan, nominal_state, solver, 'zeros', float(initial), 0.0
    if k <= WARMUP_LAST_DAY:
        protected = net_forecast['naive'][k, 1:TARGETS].copy()
        aux = float(net_forecast['naive'][k, 0])
        terminal_mode = 'fixed6000'
        kind = 'naive'
    else:
        template_r, aux_r = protection[k, 1:TARGETS], float(protection[k, 0])
        protected = net_forecast['issued'][k, 1:TARGETS] + template_r
        aux = float(net_forecast['issued'][k, 0]) + aux_r
        terminal_mode = mode
        kind = 'issued'
    _, _, nominal_emergency, _, estimated = feedback_step(carry - aux, float(initial))
    try:
        plan, _, _, _, nominal_state, solver, _ = solve_day(protected, price_src,
                                                            float(estimated), terminal_mode)
    except Exception as exc:                                        # noqa: BLE001 - re-raise w/ ctx
        raise RuntimeError(f'day {k} ({dates_k}) source={kind} estimated={estimated} '
                           f'finite={bool(np.isfinite(protected).all())}: {exc}') from exc
    return plan, nominal_state, solver, kind, float(estimated), float(nominal_emergency)


def run_group(group, nat_load, nat_pv, dates, price_src, net_forecast, protection,
              day_stop=NATURAL_DAYS):
    """Sequential natural-day execution with the midnight carry; builds ledger A and B."""
    mode = group['mode']
    strategy_id = group['id']
    nat_net = (nat_load - nat_pv) * DT
    price_nat = natural_price(price_src)
    state = TERMINAL_FIXED_KWH
    carry = 0.0
    natural_rows, solver_rows = [], []
    plan_store, nominal_store, e_store, states_store = {}, {}, {}, {}
    nominal_emergency_log = {}
    physical = {}
    for k in range(min(day_stop, NATURAL_DAYS)):
        initial = state
        plan, nominal_state, solver, kind, estimated, nominal_emergency = plan_for_day(
            k, mode, carry, initial, net_forecast, protection, price_src, str(dates[k].date()))
        # 2025-01-01 keeps the battery idle; greedy feedback resumes Jan 2
        battery = k > 0
        if k == 0:
            # 2025-01-01 00:00-00:10 has no source observation: the battery idles at 6000 kWh and
            # the segment is neither billed nor used as a training/label value
            c0 = d0 = e0 = w0 = 0.0
            after_first = float(initial)
        else:
            c0, d0, e0, w0, after_first = feedback_step(carry - nat_net[k, 0], initial,
                                                        battery=battery)
        # template slots 0..142 are executed on natural slots 1..143; slot 143 is the carry
        c, d, e, w, trail = feedback(plan[:TEMPLATE_SLOTS - 1], nat_net[k, 1:TEMPLATE_SLOTS],
                                     after_first, battery=battery)
        states = np.r_[initial, after_first, trail[1:]]
        c_full = np.r_[c0, c]
        d_full = np.r_[d0, d]
        e_full = np.r_[e0, e]
        w_full = np.r_[w0, w]
        plan_full = np.r_[carry, plan[:TEMPLATE_SLOTS - 1]]
        balance_vector = plan_full + d_full + e_full - nat_net[k] - c_full - w_full
        if k == 0:
            balance_vector[0] = 0.0                 # the unknown opening segment is excluded
        for name, value in (('balance', float(np.max(np.abs(balance_vector)))),
                            ('state_recursion',
                             float(np.max(np.abs(np.diff(states) - ETA * c_full
                                                  + d_full / ETA)))),
                            ('capacity', float(max(0.0, STATE_MIN_KWH - states.min(),
                                                   states.max() - STATE_MAX_KWH))),
                            ('power', float(max(0.0, c_full.max() - POWER_MAX_KW * DT,
                                                d_full.max() - POWER_MAX_KW * DT))),
                            ('nonnegative',
                             float(max(0.0, -min(c_full.min(), d_full.min(), e_full.min(),
                                                 w_full.min())))),
                            ('mutex', float(np.minimum(c_full, d_full).max()))):
            physical[name] = max(physical.get(name, 0.0), value)
        natural_rows.append(pd.DataFrame({
            'strategy_id': strategy_id, 'time_version': TIME_VERSION,
            'date': str(dates[k].date()),
            'natural_slot': np.arange(TEMPLATE_SLOTS),
            'interval_start': [(dates[k] + timedelta(minutes=10 * t)).isoformat(sep=' ')
                               for t in range(TEMPLATE_SLOTS)],
            'interval_end': [(dates[k] + timedelta(minutes=10 * (t + 1))).isoformat(sep=' ')
                             for t in range(TEMPLATE_SLOTS)],
            'plan_source': ['carry from previous template slot 143' if t == 0
                            else f'{kind} template slot {t - 1}' for t in range(TEMPLATE_SLOTS)],
            'price_yuan_kWh': price_nat, 'load_kW': nat_load[k], 'pv_kW': nat_pv[k],
            'net_kWh': nat_net[k], 'plan_kWh': plan_full, 'charge_kWh': c_full,
            'discharge_kWh': d_full, 'emergency_kWh': e_full, 'unused_kWh': w_full,
            'state_start_kWh': states[:-1], 'state_end_kWh': states[1:],
            'planned_cost_yuan': plan_full * price_nat,
            'emergency_cost_yuan': EMERGENCY_MULTIPLIER * e_full * price_nat}))
        if kind != 'zeros':
            solver_rows.append(dict(strategy_id=strategy_id, time_version=TIME_VERSION, mode=mode,
                                    date=str(dates[k].date()), day_index=k, plan_source=kind,
                                    terminal_target=(np.nan if solver['terminal_target'] is None
                                                     else solver['terminal_target']),
                                    terminal_lower_kWh=solver['terminal_lower_kWh'],
                                    terminal_upper_kWh=solver['terminal_upper_kWh'],
                                    s143_kWh=solver['s143_kWh'], s144_kWh=solver['s144_kWh'],
                                    estimated_00_10_kWh=estimated,
                                    nominal_first_emergency_kWh=nominal_emergency,
                                    elapsed_seconds=solver['elapsed_seconds'],
                                    mip_gap=solver['mip_gap'],
                                    objective_gap_yuan=solver.get('objective_bound_gap_yuan',
                                                                  np.nan),
                                    **{f'check_{key}': float(value)
                                       for key, value in solver['checks'].items()}))
        plan_store[k], nominal_store[k], e_store[k], states_store[k] = (plan, nominal_state,
                                                                        e_full, states)
        nominal_emergency_log[k] = nominal_emergency
        carry = float(plan[TEMPLATE_SLOTS - 1])
        state = float(states[-1])
    natural_frame = pd.concat(natural_rows, ignore_index=True)
    gap = ((natural_frame.date == '2025-01-01') & (natural_frame.natural_slot == 0)).to_numpy()
    natural_frame['gap_unknown'] = gap
    natural_frame.loc[gap, 'net_kWh'] = np.nan
    natural_frame.loc[gap, 'planned_cost_yuan'] = np.nan
    natural_frame.loc[gap, 'emergency_cost_yuan'] = np.nan
    if day_stop < NATURAL_DAYS:
        return dict(natural=natural_frame, template=pd.DataFrame(), tail=None,
                    solver=pd.DataFrame(solver_rows), physical=physical, final_state=state,
                    nominal_emergency_log=nominal_emergency_log, group=group)
    # 2026-01-01 00:00-00:10 executes the 2025-12-31 template slot 143 exactly once
    tail_c, tail_d, tail_e, tail_w, tail_end = feedback_step(
        carry - nat_net[TAIL_NATURAL_DAY, 0], state)
    tail = dict(date='2026-01-01', interval='00:00-00:10',
                price=float(price_src[TEMPLATE_SLOTS - 1]), planned_kWh=float(carry),
                charge_kWh=float(tail_c), discharge_kWh=float(tail_d),
                emergency_kWh=float(tail_e), unused_kWh=float(tail_w),
                state_start_kWh=float(state), state_end_kWh=float(tail_end),
                planned_cost_yuan=float(carry * price_src[TEMPLATE_SLOTS - 1]),
                emergency_cost_yuan=float(EMERGENCY_MULTIPLIER * tail_e
                                          * price_src[TEMPLATE_SLOTS - 1]))
    template_rows = []
    for k in range(1, NATURAL_DAYS):
        plan = plan_store[k]
        next_emergency = float(e_store[k + 1][0]) if k + 1 < NATURAL_DAYS else float(tail_e)
        next_cost = (float(EMERGENCY_MULTIPLIER * next_emergency
                           * price_src[TEMPLATE_SLOTS - 1])
                     if k + 1 < NATURAL_DAYS else float(tail['emergency_cost_yuan']))
        template_rows.append(pd.DataFrame({
            'strategy_id': strategy_id, 'time_version': TIME_VERSION,
            'publication_date': str(dates[k].date()), 'template_slot': np.arange(TEMPLATE_SLOTS),
            'source_label': [(f'{10 * ((t + 1) // 6) // 60:02d}:{10 * (t + 1) % 60:02d}'
                              if t < TEMPLATE_SLOTS - 1 else '0:00+1')
                             for t in range(TEMPLATE_SLOTS)],
            'interval_start': [(dates[k] + timedelta(minutes=10 * (t + 1))).isoformat(sep=' ')
                               for t in range(TEMPLATE_SLOTS)],
            'price_yuan_kWh': price_src, 'planned_kWh': plan,
            'emergency_kWh': np.r_[e_store[k][1:TEMPLATE_SLOTS], next_emergency],
            'planned_cost_yuan': plan * price_src,
            'emergency_cost_yuan': np.r_[EMERGENCY_MULTIPLIER * e_store[k][1:TEMPLATE_SLOTS]
                                         * price_src[:TEMPLATE_SLOTS - 1], next_cost],
            'nominal_state_start_kWh': np.asarray(nominal_store[k])[:-1],
            'nominal_state_end_kWh': np.asarray(nominal_store[k])[1:],
            'actual_state_start_kWh': states_store[k][1:],
            'actual_state_end_kWh': np.r_[states_store[k][2:],
                                          (states_store[k + 1][1] if k + 1 < NATURAL_DAYS
                                           else tail_end)],
            'terminal_target_kWh': (TERMINAL_FIXED_KWH if mode == 'fixed6000' else np.nan)
                                   if k > WARMUP_LAST_DAY else np.nan}))
    return dict(natural=natural_frame,
                template=pd.concat(template_rows, ignore_index=True), tail=tail,
                solver=pd.DataFrame(solver_rows), physical=physical, final_state=state,
                nominal_emergency_log=nominal_emergency_log, group=group)


def ledger_totals(result, price_src=None):
    """Ledger A (natural days Feb 1..Dec 31) and ledger B (template publications), bridge."""
    natural, template = result['natural'], result['template']
    a = natural[(pd.to_datetime(natural.date) >= pd.Timestamp('2025-02-01'))
                & (pd.to_datetime(natural.date) <= pd.Timestamp('2025-12-31'))]
    b = template[(pd.to_datetime(template.publication_date) >= pd.Timestamp('2025-02-01'))
                 & (pd.to_datetime(template.publication_date) <= pd.Timestamp('2025-12-31'))]
    assert len(a) == 334 * TEMPLATE_SLOTS and len(b) == 334 * TEMPLATE_SLOTS, (len(a), len(b))
    head = a[a.natural_slot == 0].iloc[0]
    tail = result['tail']
    ledger_a = dict(
        planned_kWh=float(a.plan_kWh.sum()), emergency_kWh=float(a.emergency_kWh.sum()),
        unused_kWh=float(a.unused_kWh.sum()), charge_kWh=float(a.charge_kWh.sum()),
        discharge_kWh=float(a.discharge_kWh.sum()),
        loss_kWh=float((0.1 * a.charge_kWh + (1 / ETA - 1) * a.discharge_kWh).sum()),
        planned_cost_yuan=float(a.planned_cost_yuan.sum()),
        emergency_cost_yuan=float(a.emergency_cost_yuan.sum()),
        total_cost_yuan=float(a.planned_cost_yuan.sum() + a.emergency_cost_yuan.sum()),
        emergency_slots=int((a.emergency_kWh > ENERGY_TOL_KWH).sum()),
        emergency_days=int(a[a.emergency_kWh > ENERGY_TOL_KWH].date.nunique()),
        initial_kWh=float(head.state_start_kWh),
        final_kWh=float(a[a.natural_slot == TEMPLATE_SLOTS - 1].iloc[-1].state_end_kWh),
        segments=int(len(a)))
    ledger_b = dict(
        planned_kWh=float(b.planned_kWh.sum()), emergency_kWh=float(b.emergency_kWh.sum()),
        planned_cost_yuan=float(b.planned_cost_yuan.sum()),
        emergency_cost_yuan=float(b.emergency_cost_yuan.sum()),
        total_cost_yuan=float(b.planned_cost_yuan.sum() + b.emergency_cost_yuan.sum()),
        segments=int(len(b)))
    head_cost = float(head.planned_cost_yuan + head.emergency_cost_yuan)
    tail_cost = float(tail['planned_cost_yuan'] + tail['emergency_cost_yuan'])
    return dict(ledger_A=ledger_a, ledger_B=ledger_b, head_cost_yuan=head_cost,
                tail_cost_yuan=tail_cost,
                bridge_difference_yuan=float(ledger_b['total_cost_yuan']
                                             - ledger_a['total_cost_yuan']),
                bridge_expected_yuan=float(tail_cost - head_cost),
                bridge_residual_yuan=float(abs(ledger_b['total_cost_yuan']
                                               - ledger_a['total_cost_yuan']
                                               - (tail_cost - head_cost))))


# ======================================================================================
# 5. metric frames
# ======================================================================================
def evaluation_solver(result):
    """The solver log restricted to the formal evaluation window (Feb 1..Dec 31)."""
    solver = result['solver']
    mask = ((pd.to_datetime(solver.date) >= pd.Timestamp('2025-02-01'))
            & (pd.to_datetime(solver.date) <= pd.Timestamp('2025-12-31')))
    return solver[mask]


def summary_frame(results):
    rows = []
    for result in results:
        a = ledger_totals(result)
        solver = evaluation_solver(result)
        row = dict(strategy_id=result['group']['id'], mode=result['group']['mode'],
                   time_version=TIME_VERSION)
        row.update(a['ledger_A'])
        row.update({f'B_{key}': value for key, value in a['ledger_B'].items()})
        row['bridge_residual_yuan'] = a['bridge_residual_yuan']
        row['nominal_S143_mean_kWh'] = float(solver.s143_kWh.mean())
        row['nominal_S143_max_kWh'] = float(solver.s143_kWh.max())
        row['nominal_S144_mean_kWh'] = float(solver.s144_kWh.mean())
        row['solver_seconds'] = float(solver.elapsed_seconds.sum())
        row['max_mip_gap'] = float(solver.mip_gap.max())
        row['solver_failures'] = 0
        row['formal_solves'] = int(len(solver))
        rows.append(row)
    return pd.DataFrame(rows)


def contrast_row(results):
    by_id = {r['group']['id']: r for r in results}
    treat = ledger_totals(by_id['N_free'])
    base = ledger_totals(by_id['N_fixed6000'])
    a, b = treat['ledger_A'], base['ledger_A']
    return dict(comparison=CONTRAST['label'], treatment='N_free', baseline='N_fixed6000',
                A_delta_planned_cost_yuan=a['planned_cost_yuan'] - b['planned_cost_yuan'],
                A_delta_emergency_cost_yuan=a['emergency_cost_yuan'] - b['emergency_cost_yuan'],
                A_delta_total_cost_yuan=a['total_cost_yuan'] - b['total_cost_yuan'],
                A_delta_total_percent=100 * (a['total_cost_yuan'] - b['total_cost_yuan'])
                / b['total_cost_yuan'],
                A_delta_emergency_kWh=a['emergency_kWh'] - b['emergency_kWh'],
                A_delta_planned_kWh=a['planned_kWh'] - b['planned_kWh'],
                A_delta_loss_kWh=a['loss_kWh'] - b['loss_kWh'],
                A_delta_final_kWh=a['final_kWh'] - b['final_kWh'],
                A_delta_emergency_days=a['emergency_days'] - b['emergency_days'],
                B_delta_total_cost_yuan=treat['ledger_B']['total_cost_yuan']
                - base['ledger_B']['total_cost_yuan'])


def daily_contrast_frame(results):
    by_id = {r['group']['id']: r for r in results}
    base, treat = by_id['N_fixed6000']['natural'], by_id['N_free']['natural']
    mask = ((pd.to_datetime(base.date) >= pd.Timestamp('2025-02-01'))
            & (pd.to_datetime(base.date) <= pd.Timestamp('2025-12-31')))
    base, treat = base[mask], treat[mask]
    cost_base = (base.planned_cost_yuan + base.emergency_cost_yuan).groupby(base.date).sum()
    cost_treat = (treat.planned_cost_yuan + treat.emergency_cost_yuan).groupby(treat.date).sum()
    em_base = base.emergency_kWh.groupby(base.date).sum()
    em_treat = treat.emergency_kWh.groupby(treat.date).sum()
    return pd.DataFrame({'date': cost_base.index, 'fixed_total_cost_yuan': cost_base.to_numpy(),
                         'free_total_cost_yuan': cost_treat.to_numpy(),
                         'free_minus_fixed_yuan': cost_treat.to_numpy() - cost_base.to_numpy(),
                         'fixed_emergency_kWh': em_base.to_numpy(),
                         'free_emergency_kWh': em_treat.to_numpy(),
                         'delta_emergency_kWh': em_treat.to_numpy() - em_base.to_numpy()})


def monthly_contrast_frame(results):
    daily = daily_contrast_frame(results).assign(month=lambda frame: frame.date.str[:7])
    return daily.groupby('month').agg(
        fixed_total_cost_yuan=('fixed_total_cost_yuan', 'sum'),
        free_total_cost_yuan=('free_total_cost_yuan', 'sum'),
        free_minus_fixed_yuan=('free_minus_fixed_yuan', 'sum'),
        fixed_emergency_kWh=('fixed_emergency_kWh', 'sum'),
        free_emergency_kWh=('free_emergency_kWh', 'sum'),
        delta_emergency_kWh=('delta_emergency_kWh', 'sum')).reset_index()


def window_frame(results, price_src):
    price_nat = natural_price(price_src)
    windows = [('00-06', 0, 36), ('06-10', 36, 60), ('10-24', 60, 144)]
    diagnostic = [('19-21', 114, 126), ('23-24', 138, 144)]
    rows = []
    for result in results:
        natural = result['natural']
        mask = ((pd.to_datetime(natural.date) >= pd.Timestamp('2025-02-01'))
                & (pd.to_datetime(natural.date) <= pd.Timestamp('2025-12-31')))
        a = natural[mask]
        emergency = a.emergency_kWh.to_numpy().reshape(334, TEMPLATE_SLOTS)
        planned = a.plan_kWh.to_numpy().reshape(334, TEMPLATE_SLOTS)
        grid = np.broadcast_to(price_nat, (334, TEMPLATE_SLOTS))
        for scope, group in (('main', windows), ('diagnostic', diagnostic)):
            for name, start, stop in group:
                rows.append(dict(strategy_id=result['group']['id'], scope=scope, window=name,
                                 emergency_kWh=float(emergency[:, start:stop].sum()),
                                 emergency_cost_yuan=float(
                                     EMERGENCY_MULTIPLIER * (emergency[:, start:stop]
                                                             * grid[:, start:stop]).sum()),
                                 planned_cost_yuan=float((planned[:, start:stop]
                                                          * grid[:, start:stop]).sum())))
    return pd.DataFrame(rows)


def stability_frame(results):
    daily = daily_contrast_frame(results)
    monthly = monthly_contrast_frame(results)
    return pd.DataFrame([dict(
        comparison=CONTRAST['label'],
        days_improved=int((daily.free_minus_fixed_yuan < -COST_TOL_YUAN).sum()),
        days_worse=int((daily.free_minus_fixed_yuan > COST_TOL_YUAN).sum()),
        days_total=int(len(daily)),
        months_improved=int((monthly.free_minus_fixed_yuan < -COST_TOL_YUAN).sum()),
        months_worse=int((monthly.free_minus_fixed_yuan > COST_TOL_YUAN).sum()),
        months_total=int(len(monthly)),
        worst_day=str(daily.loc[daily.free_minus_fixed_yuan.idxmax(), 'date']),
        worst_day_yuan=float(daily.free_minus_fixed_yuan.max()),
        best_day=str(daily.loc[daily.free_minus_fixed_yuan.idxmin(), 'date']),
        best_day_yuan=float(daily.free_minus_fixed_yuan.min()))])


def inventory_frame(results, price_src):
    nu = float(np.median(natural_price(price_src)) / ETA)
    rows = []
    for result in results:
        a = ledger_totals(result)['ledger_A']
        solver = evaluation_solver(result)
        rows.append(dict(strategy_id=result['group']['id'], mode=result['group']['mode'],
                         nominal_S143_mean_kWh=float(solver.s143_kWh.mean()),
                         nominal_S143_min_kWh=float(solver.s143_kWh.min()),
                         nominal_S143_max_kWh=float(solver.s143_kWh.max()),
                         nominal_S143_days_at_6000=int((np.abs(solver.s143_kWh
                                                               - TERMINAL_FIXED_KWH) < 1e-6).sum()),
                         nominal_S143_days_at_floor=int((np.abs(solver.s143_kWh
                                                                - STATE_MIN_KWH) < 1e-6).sum()),
                         nominal_S144_mean_kWh=float(solver.s144_kWh.mean()),
                         actual_initial_kWh=a['initial_kWh'], actual_final_kWh=a['final_kWh'],
                         change_kWh=a['final_kWh'] - a['initial_kWh'],
                         nu_yuan_per_internal_kWh=nu,
                         inventory_adjustment_yuan=float(-nu * (a['final_kWh']
                                                                - a['initial_kWh'])),
                         adjusted_cost_yuan=float(a['total_cost_yuan']
                                                  - nu * (a['final_kWh'] - a['initial_kWh']))))
    return pd.DataFrame(rows)


def energy_frame(results):
    by_id = {r['group']['id']: ledger_totals(r)['ledger_A'] for r in results}
    a, b = by_id['N_free'], by_id['N_fixed6000']
    row = dict(comparison=CONTRAST['label'],
               delta_planned_kWh=a['planned_kWh'] - b['planned_kWh'],
               delta_emergency_kWh=a['emergency_kWh'] - b['emergency_kWh'],
               delta_charge_kWh=a['charge_kWh'] - b['charge_kWh'],
               delta_discharge_kWh=a['discharge_kWh'] - b['discharge_kWh'],
               delta_unused_kWh=a['unused_kWh'] - b['unused_kWh'],
               delta_loss_kWh=a['loss_kWh'] - b['loss_kWh'],
               delta_end_kWh=a['final_kWh'] - b['final_kWh'])
    row['identity_residual_kWh'] = float(row['delta_planned_kWh'] + row['delta_emergency_kWh']
                                         - row['delta_unused_kWh'] - row['delta_loss_kWh']
                                         - row['delta_end_kWh'])
    return pd.DataFrame([row])


def soc_frame(results):
    frames = []
    for result in results:
        solver = evaluation_solver(result)
        natural = result['natural']
        day_end = natural[natural.natural_slot == TEMPLATE_SLOTS - 1][['date', 'state_end_kWh']]
        mask = ((pd.to_datetime(day_end.date) >= pd.Timestamp('2025-02-01'))
                & (pd.to_datetime(day_end.date) <= pd.Timestamp('2025-12-31')))
        day_end = day_end[mask]
        assert len(day_end) == len(solver) == 334, (len(day_end), len(solver))
        frames.append(pd.DataFrame({'strategy_id': result['group']['id'],
                                    'date': pd.to_datetime(solver.date).to_numpy(),
                                    'nominal_s143': solver.s143_kWh.to_numpy(),
                                    'actual_end': day_end.state_end_kWh.to_numpy()}))
    return pd.concat(frames, ignore_index=True)


# ======================================================================================
# 6. checks
# ======================================================================================
def artificial_mapping_check(attachments, dates, nat_load, nat_pv):
    """Three days recomputed independently from the raw rows with explicit clock arithmetic."""
    load_src = attachments['load_src']
    rows = []
    worst = 0.0
    for k in (1, 100, 364):
        # natural day k slot 0 <- row k-1 column 143 ; slots 1..143 <- row k columns 0..142
        expect_nat = np.empty(TEMPLATE_SLOTS)
        expect_nat[0] = load_src[k - 1, TEMPLATE_SLOTS - 1]
        expect_nat[1:] = load_src[k, :TEMPLATE_SLOTS - 1]
        gap = float(np.max(np.abs(expect_nat - nat_load[k])))
        # the truth targets of publication k: previous midnight in front, then row k
        truth = np.empty(TARGETS)
        truth[0] = load_src[k - 1, TEMPLATE_SLOTS - 1]
        truth[1:] = load_src[k, :]
        gap = max(gap, float(np.max(np.abs(truth - np.r_[load_src[k - 1, TEMPLATE_SLOTS - 1],
                                                         load_src[k, :]]))))
        # clock arithmetic: h=0 -> 00:00-00:10, h=1 -> 00:10-00:20, h=143 -> 23:50-24:00,
        # h=144 -> next day 00:00-00:10
        stamps = {0: (dates[k] + timedelta(0), dates[k] + timedelta(minutes=10)),
                  1: (dates[k] + timedelta(minutes=10), dates[k] + timedelta(minutes=20)),
                  143: (dates[k] + timedelta(hours=23, minutes=50), dates[k] + timedelta(days=1)),
                  144: (dates[k] + timedelta(days=1), dates[k] + timedelta(days=1, minutes=10))}
        for h, (start, end) in stamps.items():
            assert (end - start) == timedelta(minutes=10), (k, h)
        worst = max(worst, gap)
        rows.append(dict(day_index=int(k), date=str(dates[k].date()), max_difference=gap,
                         h0_interval=f'{stamps[1][0].isoformat(sep=" ")}',
                         h144_interval=f'{stamps[144][0].isoformat(sep=" ")}'))
    return dict(cases=rows, worst=worst, all_passed=bool(worst == 0.0))


def forecast_checks(archive, nat_pv, nat_load, dates, price_src):
    """First feature/trainable days, gate defaults, midnight naive fallback, unfaked series."""
    out = dict(summary=archive['summary'], rounds=[])
    out['rounds'].append(dict(check='load first feature day is 9 (Jan 10 formed as 0-based 9)',
                              value=first_feature_day('load'), expected=9,
                              passed=bool(first_feature_day('load') == 9)))
    out['rounds'].append(dict(check='pv first feature day is 8 (Jan 9)',
                              value=first_feature_day('pv'), expected=8,
                              passed=bool(first_feature_day('pv') == 8)))
    # the midnight target must never read the current day's still-unknown midnight value:
    # perturb it and require the published feature row to be unchanged
    k = 200
    probe_pv = nat_pv.copy()
    probe_pv[k, 0] = nat_pv[k, 0] + 12345.0
    _FEATURE_CACHE.clear()
    _, base_before, _, _ = feature_rows_145(nat_load, nat_pv, dates, k, 'pv')
    _, base_after, _, _ = feature_rows_145(nat_load, probe_pv, dates, k, 'pv')
    _FEATURE_CACHE.clear()
    unchanged = bool(np.array_equal(base_before, base_after))
    uses_previous = bool(base_before[TARGETS - 1] == nat_pv[k - 1, 0])
    out['rounds'].append(dict(check='h=144 pv base ignores the unknown current midnight',
                              value=[float(base_before[TARGETS - 1]),
                                     float(base_after[TARGETS - 1])],
                              expected=[float(nat_pv[k - 1, 0]), float(nat_pv[k - 1, 0])],
                              passed=bool(unchanged and uses_previous)))
    # an empty support window must not zero the whole day
    empty = np.zeros_like(nat_pv)
    start, stop, valid = support_window(empty, k)
    gate_all_one = bool(start == 0 and stop == TEMPLATE_SLOTS - 1 and valid == 0)
    out['rounds'].append(dict(check='empty support window -> s=0,e=143,v=0 and no zeroing',
                              value=[start, stop, valid], expected=[0, 143, 0],
                              passed=gate_all_one))
    # the gap is never faked as a value
    out['rounds'].append(dict(check='2025-01-01 00:00-00:10 stays unknown',
                              value=float(np.isnan(nat_load[0, 0])),
                              expected=1.0, passed=bool(np.isnan(nat_load[0, 0]))))
    out['rounds'].append(dict(check='every stored power value is carried into the natural arrays',
                              value=float(np.isfinite(nat_load).sum()
                                         + np.isfinite(nat_pv).sum()),
                              expected=float(2 * NATURAL_DAYS * TEMPLATE_SLOTS),
                              passed=bool(int(np.isfinite(nat_load).sum()
                                              + np.isfinite(nat_pv).sum())
                                          == 2 * NATURAL_DAYS * TEMPLATE_SLOTS)))
    out['all_passed'] = bool(all(r['passed'] for r in out['rounds']))
    return out


def mode_boundary_checks(price_src):
    """Mode switch, committed-carry shortfall, empty/full battery and negative protection."""
    rounds = []
    protected = np.full(TEMPLATE_SLOTS, 300.0)
    fixed = build_model(protected, price_src, 6000.0, 'fixed6000')
    free = build_model(protected, price_src, 6000.0, 'free')
    index = fixed['terminal_index']
    rounds.append(dict(check='mode switch changes only the S143 bound',
                       value=[int(v) for v in np.flatnonzero(fixed['bounds'].lb
                                                             != free['bounds'].lb)],
                       expected=[index],
                       passed=bool(np.array_equal(np.flatnonzero(fixed['bounds'].lb
                                                                 != free['bounds'].lb), [index])
                                   and np.array_equal(np.flatnonzero(fixed['bounds'].ub
                                                                     != free['bounds'].ub),
                                                      [index])
                                   and np.array_equal(fixed['objective'], free['objective'])
                                   and np.array_equal(fixed['integrality'], free['integrality'])
                                   and int((fixed['constraints'].A
                                            != free['constraints'].A).nnz) == 0)))
    rounds.append(dict(check='terminal index is S143 (true natural-day 24:00)',
                       value=int(index), expected=int(4 * TEMPLATE_SLOTS + 143),
                       passed=bool(index == 4 * TEMPLATE_SLOTS + 143)))
    # committed carry below the auxiliary protected need still yields a feasible plan, and the
    # nominal first-interval emergency is recorded rather than forced to zero
    _, _, emergency, _, estimated = feedback_step(-2000.0, STATE_MIN_KWH)
    plan, _, _, _, state, _, _ = solve_day(np.full(TEMPLATE_SLOTS, 100.0), price_src,
                                           float(estimated), 'fixed6000')
    rounds.append(dict(check='nominal first-interval emergency recorded; plan stays feasible',
                       value=[float(emergency), float(state[TEMPLATE_SLOTS - 1])],
                       expected=[2000.0, TERMINAL_FIXED_KWH],
                       passed=bool(abs(emergency - 2000.0) < 1e-9
                                   and abs(state[TEMPLATE_SLOTS - 1]
                                           - TERMINAL_FIXED_KWH) < 1e-6)))
    # empty battery cannot discharge below the floor
    _, discharge, emergency_empty, _, after = feedback_step(-2000.0, STATE_MIN_KWH)
    rounds.append(dict(check='empty battery cannot discharge below 1200 kWh',
                       value=[float(discharge), float(after), float(emergency_empty)],
                       expected=[0.0, STATE_MIN_KWH, 2000.0],
                       passed=bool(discharge == 0.0 and abs(after - STATE_MIN_KWH) < 1e-12)))
    # full battery absorbs surplus up to the power cap only
    charge, _, _, unused, after_full = feedback_step(3000.0, STATE_MAX_KWH)
    rounds.append(dict(check='full battery absorbs only up to the power cap',
                       value=[float(charge), float(unused), float(after_full)],
                       expected=[0.0, 3000.0, STATE_MAX_KWH],
                       passed=bool(charge == 0.0 and unused == 3000.0)))
    # negative protection is retained: the plan never buys, and the absorbed energy is split
    # between charging and unused - with a free terminal those two are tied optima
    negative = -np.full(TEMPLATE_SLOTS, 5.0)
    q, c, d, w, _, summary_neg, _ = solve_day(negative, price_src, 6000.0, 'free')
    rounds.append(dict(check='negative protection retained (no purchase, balance closes)',
                       value=[float(q.sum()), float(c.sum() + w.sum() - d.sum()),
                              float(summary_neg['cost_yuan'])],
                       expected=[0.0, 5.0 * TEMPLATE_SLOTS, 0.0],
                       passed=bool(q.sum() == 0.0
                                   and abs(c.sum() + w.sum() - d.sum() - 5.0 * TEMPLATE_SLOTS)
                                   < 1e-6
                                   and summary_neg['cost_yuan'] == 0.0)))
    # the feedback matches the frozen module-05 controller on a real day
    bm = _load('bm_check', 'code/05_q2_baseline.py')
    plan = np.full(TEMPLATE_SLOTS, 500.0)
    load_day = np.full(TEMPLATE_SLOTS, 600.0)
    pv_day = np.full(TEMPLATE_SLOTS, 100.0)
    mine = feedback(plan, (load_day - pv_day) * DT, 6000.0)
    theirs = bm.control(plan, load_day, pv_day, 6000.0)
    gap = float(max(np.max(np.abs(mine[0] - theirs[0])), np.max(np.abs(mine[1] - theirs[1])),
                    np.max(np.abs(mine[2] - theirs[2])), np.max(np.abs(mine[3] - theirs[3])),
                    np.max(np.abs(mine[4] - theirs[4]))))
    rounds.append(dict(check='feedback equals the frozen controller', value=gap, expected=0.0,
                       passed=bool(gap == 0.0)))
    return dict(rounds=rounds, all_passed=bool(all(r['passed'] for r in rounds)))


def perturbed_forecasts(k, factor_load, factor_pv, nat_load, nat_pv, dates, load_src, pv_src,
                        clean_net, clean_protection):
    """Rebuild only the publication days at or after k and merge with the unchanged prefix."""
    load2 = load_src.copy()
    pv2 = pv_src.copy()
    load2[k:] *= factor_load
    pv2[k:] *= factor_pv
    nat2_load, nat2_pv = natural_arrays(load2, pv2)
    _FEATURE_CACHE.clear()
    arch2 = build_forecast_archive(nat2_load, nat2_pv, dates, load2, pv2, day_from=k)
    days = np.arange(NATURAL_DAYS)
    mask = (days >= k)[:, None]
    net2 = dict(
        issued=np.where(mask, (arch2['issued_load'] - arch2['issued_pv']) * DT,
                        clean_net['issued']),
        naive=np.where(mask, (arch2['naive_load'] - arch2['naive_pv']) * DT, clean_net['naive']))
    errors2 = errors_145(load2, pv2, nat2_load, nat2_pv, net2['issued'])
    protection2 = np.zeros_like(clean_protection)
    protection2[:k] = clean_protection[:k]
    rebuilt, m_arch, _ = protection_archive(errors2)
    protection2[k:] = rebuilt[k:]
    return dict(nat_load=nat2_load, nat2_pv=nat2_pv, net=net2, protection=protection2,
                errors=errors2, archive=arch2)


def perturbation_checks(results, nat_load, nat_pv, dates, price_src, load_src, pv_src,
                        clean_net, clean_protection):
    """Six future-truth perturbations: the whole prefix and the perturbation-day plan must not
    move."""
    cases = []
    worst_prefix = 0.0
    worst_plan = 0.0
    for k in (EVAL_FIRST_DAY, 171, 354):
        for name, factor_load, factor_pv in (('load_x1.2', 1.2, 1.0), ('pv_x0.7', 1.0, 0.7)):
            rebuild = perturbed_forecasts(k, factor_load, factor_pv, nat_load, nat_pv, dates,
                                          load_src, pv_src, clean_net, clean_protection)
            entry = dict(day_index=int(k), date=str(dates[k].date()), case=name, groups={})
            for group in GROUPS:
                clean_result = {r['group']['id']: r for r in results}[group['id']]
                probe = run_group(group, rebuild['nat_load'], rebuild['nat2_pv'], dates, price_src,
                                  rebuild['net'], rebuild['protection'], day_stop=k + 1)
                clean_natural = clean_result['natural']
                today = str(dates[k].date())
                # everything strictly before the perturbation day must be bit-identical
                strict = clean_natural[clean_natural.date < today]
                probe_strict = probe['natural'][probe['natural'].date < today]
                assert len(strict) == len(probe_strict) == k * TEMPLATE_SLOTS
                plan_gap = float(np.max(np.abs(strict.plan_kWh.to_numpy()
                                               - probe_strict.plan_kWh.to_numpy())))
                state_gap = float(np.max(np.abs(strict.state_end_kWh.to_numpy()
                                                - probe_strict.state_end_kWh.to_numpy())))
                clean_day = clean_natural[clean_natural.date == today]
                probe_day = probe['natural'][probe['natural'].date == today]
                day_plan_gap = float(np.max(np.abs(clean_day.plan_kWh.to_numpy()
                                                   - probe_day.plan_kWh.to_numpy())))
                entering_gap = float(abs(float(clean_day.state_start_kWh.iloc[0])
                                         - float(probe_day.state_start_kWh.iloc[0])))
                protection_gap = float(np.max(np.abs(clean_protection[:k + 1]
                                                     - rebuild['protection'][:k + 1])))
                entry['groups'][group['id']] = dict(
                    prefix_plan_difference_kWh=plan_gap, prefix_state_difference_kWh=state_gap,
                    perturbation_day_plan_difference_kWh=day_plan_gap,
                    perturbation_day_entering_state_difference_kWh=entering_gap,
                    prefix_protection_difference_kWh=protection_gap,
                    perturbation_day_actual_emergency_kWh=float(probe_day.emergency_kWh.sum()),
                    clean_day_actual_emergency_kWh=float(clean_day.emergency_kWh.sum()),
                    perturbation_day_state_end_kWh=float(probe_day.state_end_kWh.iloc[-1]),
                    clean_day_state_end_kWh=float(clean_day.state_end_kWh.iloc[-1]),
                    passed=bool(plan_gap == 0.0 and state_gap == 0.0 and day_plan_gap == 0.0
                                and entering_gap == 0.0 and protection_gap == 0.0))
                worst_prefix = max(worst_prefix, plan_gap, state_gap, entering_gap)
                worst_plan = max(worst_plan, day_plan_gap)
            entry['passed'] = bool(all(v['passed'] for v in entry['groups'].values()))
            cases.append(entry)
    return dict(cases=cases, max_prefix_difference=worst_prefix,
                max_perturbation_day_plan_difference=worst_plan,
                note='a perturbation from day k on may change the actual execution of day k, but '
                     'every earlier plan, state, feature, label and protection amount must stay '
                     'bit-identical, and so must the plan published on day k itself',
                all_passed=bool(all(c['passed'] for c in cases)))


def halfday_check(results, nat_load, nat_pv, dates, price_src):
    """Natural-day half-day prefix: only slot 72 and later change, so slots 0..71 stay."""
    cases = []
    worst = 0.0
    for result in results:
        natural = result['natural']
        day = natural[natural.date == '2025-06-21']
        plan = day.plan_kWh.to_numpy()
        initial = float(day.state_start_kWh.iloc[0])
        nat_net = (nat_load - nat_pv) * DT
        clean = feedback(plan, nat_net[171], initial)
        probe_load = nat_load[171].copy()
        probe_pv = nat_pv[171].copy()
        probe_load[72:] *= 1.2
        probe_pv[72:] *= 0.7
        probe = feedback(plan, (probe_load - probe_pv) * DT, initial)
        gap = float(max(np.max(np.abs(probe[0][:72] - clean[0][:72])),
                        np.max(np.abs(probe[1][:72] - clean[1][:72])),
                        np.max(np.abs(probe[2][:72] - clean[2][:72])),
                        np.max(np.abs(probe[3][:72] - clean[3][:72])),
                        np.max(np.abs(probe[4][:73] - clean[4][:73]))))
        cases.append(dict(strategy_id=result['group']['id'], date='2025-06-21',
                          prefix_slots=72, max_prefix_difference=gap,
                          state_entering_slot_72_kWh=float(clean[4][72]),
                          passed=bool(gap == 0.0)))
        worst = max(worst, gap)
    return dict(cases=cases, max_prefix_difference=worst, all_passed=bool(worst == 0.0))


def carry_check(results, nat_load, nat_pv):
    """The carry equals the previous template's slot 143, and every interval is billed once."""
    out = {}
    worst = 0.0
    for result in results:
        natural = result['natural']
        template = result['template']
        dates_used = natural.date.unique()
        for k in range(1, NATURAL_DAYS):
            first = natural[(natural.date == dates_used[k]) & (natural.natural_slot == 0)]
            prev = template[(template.publication_date == dates_used[k - 1])
                            & (template.template_slot == TEMPLATE_SLOTS - 1)]
            if len(first) and len(prev):
                worst = max(worst, float(abs(first.plan_kWh.iloc[0]
                                             - prev.planned_kWh.iloc[0])))
        a = ledger_totals(result)['ledger_A']
        out[result['group']['id']] = a['segments']
    return dict(max_carry_difference_kWh=worst, ledger_A_segments=out,
                all_passed=bool(worst == 0.0 and set(out.values()) == {334 * TEMPLATE_SLOTS}))


def reproduction_check(group, nat_load, nat_pv, dates, price_src, net_forecast, protection,
                       clean):
    """Cold-start re-run of the fixed group: every segment and the annual cash cost must match."""
    replay = run_group(group, nat_load, nat_pv, dates, price_src, net_forecast, protection)
    left, right = clean['natural'], replay['natural']
    gaps = {}
    for column in ('plan_kWh', 'emergency_kWh', 'charge_kWh', 'discharge_kWh', 'unused_kWh',
                   'state_end_kWh'):
        gaps[column] = float(np.max(np.abs(left[column].to_numpy() - right[column].to_numpy())))
    totals_left = ledger_totals(clean)['ledger_A']['total_cost_yuan']
    totals_right = ledger_totals(replay)['ledger_A']['total_cost_yuan']
    out = dict(segment_max_differences=gaps,
               max_segment_difference_kWh=float(max(gaps.values())),
               cost_difference_yuan=float(abs(totals_left - totals_right)),
               ledger_A_left=totals_left, ledger_A_right=totals_right)
    out['all_passed'] = bool(out['max_segment_difference_kWh'] < ENERGY_TOL_KWH
                             and out['cost_difference_yuan'] <= COST_TOL_YUAN)
    return out


# ======================================================================================
# 7. figures
# ======================================================================================
def figures(summary, monthly, soc):
    produced = []
    fig, axis = plt.subplots(figsize=(11, 3.6), layout='constrained')
    axis.set_xlim(0, 24)
    axis.set_ylim(0, 3)
    axis.axis('off')
    for row, (label, start, stop, colour) in enumerate([
            ('source row k-1: label 0:00+1', 0, 0.167, '#1b6c9e'),
            ('source row k: labels 00:10 .. 23:50', 0.167, 24, '#b8501b')]):
        axis.add_patch(plt.Rectangle((start, 2.1 - 0.6 * row), stop - start, 0.45,
                                     color=colour, alpha=0.75))
        axis.text(start + (stop - start) / 2, 2.32 - 0.6 * row, label, ha='center', fontsize=9)
    axis.annotate('', xy=(0.167, 1.05), xytext=(0, 1.05),
                  arrowprops=dict(arrowstyle='->', color='black'))
    axis.text(0.09, 1.12, 'natural day k 00:00-00:10 (from the previous commitment)', fontsize=8,
              ha='center')
    axis.annotate('', xy=(24, 0.75), xytext=(0.167, 0.75),
                  arrowprops=dict(arrowstyle='->', color='black'))
    axis.text(12, 0.82, 'natural day k 00:10-24:00 (this day plan, 143 slots)', fontsize=8,
              ha='center')
    axis.text(12, 0.35, 'template day k = [k 00:10, (k+1) 00:10) = 144 slots; '
                        'the fixed mode pins S143 (true 24:00), S144 stays free',
              fontsize=8, ha='center')
    axis.set_title('start_time_v1 interval-start mapping', fontsize=10)
    path = FIG / 'q2_time_mapping_midnight.png'
    fig.savefig(path, dpi=180)
    plt.close(fig)
    produced.append(path)

    fig, axes = plt.subplots(1, 2, figsize=(11, 4.5), layout='constrained')
    index = np.arange(len(summary))
    axes[0].bar(index - 0.18, summary.B_planned_cost_yuan / 1e6, 0.36,
                label='template ledger B planned (million CNY)', color='#1b6c9e')
    axes[0].bar(index + 0.18, summary.B_emergency_cost_yuan / 1e6, 0.36,
                label='template ledger B emergency (million CNY)', color='#b8501b')
    axes[0].set_xticks(index)
    axes[0].set_xticklabels(summary.strategy_id, fontsize=8)
    axes[0].set_ylabel('Million CNY')
    axes[0].set_title('Two groups, template ledger B')
    months = np.arange(len(monthly))
    axes[1].bar(months, monthly.free_minus_fixed_yuan / 1e3, 0.6, color='#7a4fa3',
                label='free - fixed (thousand CNY)')
    axes[1].axhline(0, color='black', lw=0.8)
    axes[1].set_xticks(months)
    axes[1].set_xticklabels(monthly.month, rotation=45, ha='right', fontsize=8)
    axes[1].set_ylabel('Thousand CNY (free - fixed)')
    axes[1].set_title('Monthly cash difference, natural-day ledger A')
    for axis in axes:
        axis.grid(alpha=0.2, axis='y')
        axis.legend(loc='best', fontsize=8)
    path = FIG / 'q2_time_mapping_costs.png'
    fig.savefig(path, dpi=180)
    plt.close(fig)
    produced.append(path)

    fig, axes = plt.subplots(2, 1, figsize=(11, 7), sharex=True, layout='constrained')
    for group_id, colour in ((GROUPS[0]['id'], '#1b6c9e'), (GROUPS[1]['id'], '#b8501b')):
        block = soc[soc.strategy_id == group_id]
        axes[0].plot(block.date, block.nominal_s143, lw=0.8, color=colour, label=group_id)
        axes[1].plot(block.date, block.actual_end, lw=0.8, color=colour, label=group_id)
    for axis, title in ((axes[0], 'Nominal S143 = true natural-day 24:00'),
                        (axes[1], 'Actual natural-day 24:00 stored energy')):
        for level in (STATE_MIN_KWH, TERMINAL_FIXED_KWH, STATE_MAX_KWH):
            axis.axhline(level, color='gray', ls=':', lw=0.8)
        axis.set_ylabel('kWh')
        axis.set_title(title)
        axis.grid(alpha=0.2)
        axis.legend(loc='best', fontsize=8)
    axes[1].set_xlabel('Date (2025-02-01..2025-12-31)')
    path = FIG / 'q2_time_mapping_soc.png'
    fig.savefig(path, dpi=180)
    plt.close(fig)
    produced.append(path)
    return produced


# ======================================================================================
# 8. validation workbooks
# ======================================================================================
def write_validation_workbooks(results, mapping_frame):
    produced = []
    for result in results:
        path = OUT / f"{result['group']['id']}_validation.xlsx"
        a = ledger_totals(result)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '口径说明'
        for row in [['时间版本', TIME_VERSION],
                    ['源标签含义', '00:10 对应 00:10-00:20；0:00+1 对应次日 00:00-00:10'],
                    ['自然日构造', '自然日 d = 前一源行第 143 列 ++ 当前源行第 0..142 列'],
                    ['模板日', 'k日0:00发布 [k日00:10,(k+1)日00:10) 的 144 段计划'],
                    ['午夜承诺', '自然日 00:00-00:10 执行前一日模板末段，当日不得重解'],
                    ['固定组末态', 'S143 = 6000 kWh（真正的自然日 24:00）；S144 只受物理界限制'],
                    ['账本A', '自然日 2025-02-01..2025-12-31，334天48096段'],
                    ['账本B', '模板发布 2025-02-01..2025-12-31，48096段，整体后移10分钟'],
                    ['桥接', 'C_B - C_A = 2026-01-01 00:00-00:10 成本 - 2025-02-01 00:00-00:10 成本'],
                    ['本副本说明', '校验用副本，不替换正式 附件/附件5/result2.xlsx'],
                    ['组ID', result['group']['id']], ['末态模式', result['group']['mode']],
                    ['账本A现金总费_元', a['ledger_A']['total_cost_yuan']],
                    ['账本B现金总费_元', a['ledger_B']['total_cost_yuan']],
                    ['桥接残差_元', a['bridge_residual_yuan']]]:
            ws.append(row)
        ws = wb.create_sheet('自然日汇总')
        natural = result['natural']
        daily = natural.groupby('date').agg(
            plan_kWh=('plan_kWh', 'sum'), emergency_kWh=('emergency_kWh', 'sum'),
            unused_kWh=('unused_kWh', 'sum'), planned_cost_yuan=('planned_cost_yuan', 'sum'),
            emergency_cost_yuan=('emergency_cost_yuan', 'sum'),
            state_start_kWh=('state_start_kWh', 'first'), state_end_kWh=('state_end_kWh', 'last'))
        ws.append(['date', 'plan_kWh', 'emergency_kWh', 'unused_kWh', 'planned_cost_yuan',
                   'emergency_cost_yuan', 'state_start_kWh', 'state_end_kWh'])
        for row in daily.reset_index().itertuples(index=False):
            ws.append(list(row))
        ws = wb.create_sheet('模板计划样例')
        ws.append(['publication_date', 'source_label', 'interval_start', 'price_yuan_kWh',
                   'planned_kWh', 'emergency_kWh', 'nominal_state_end_kWh',
                   'actual_state_end_kWh'])
        sample = result['template'][result['template'].publication_date.isin(
            ['2025-02-01', '2025-12-31'])]
        for row in sample.itertuples(index=False):
            ws.append([row.publication_date, row.source_label, row.interval_start,
                       row.price_yuan_kWh, row.planned_kWh, row.emergency_kWh,
                       row.nominal_state_end_kWh, row.actual_state_end_kWh])
        ws = wb.create_sheet('映射样例')
        ws.append(['source_date', 'source_label', 'template_slot', 'interval_start',
                   'interval_end', 'available_at', 'natural_date', 'natural_slot', 'value_kW'])
        for row in mapping_frame.head(288).itertuples(index=False):
            ws.append([row.source_date, row.source_label, row.template_slot, row.interval_start,
                       row.interval_end, row.available_at, row.natural_date, row.natural_slot,
                       getattr(row, 'value_kW', '')])
        wb.save(path)
        wb.close()
        produced.append(path)
    return produced


# ======================================================================================
# 9. main
# ======================================================================================
def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--mode', choices=['register', 'repro', 'full'], default='full')
    parser.add_argument('--amend-reason', default=None)
    args = parser.parse_args()
    assert Path(sys.prefix).name == 'math_modeling', sys.prefix
    OUT.mkdir(parents=True, exist_ok=True)
    FIG.mkdir(parents=True, exist_ok=True)
    started = time.perf_counter()
    started_utc = datetime.now(timezone.utc).isoformat()

    attachments = read_attachments()
    # 366 dates: day 364 is 2025-12-31 and day 365 is 2026-01-01 (h=144 target date)
    dates = pd.date_range('2025-01-01', periods=NATURAL_DAYS + 1)
    assert len(dates) == NATURAL_DAYS + 1 and dates[NATURAL_DAYS] == pd.Timestamp('2026-01-01')
    load_src, pv_src, price_src = (attachments['load_src'], attachments['pv_src'],
                                   attachments['price_src'])
    nat_load, nat_pv = natural_arrays(load_src, pv_src)
    mapping_frame = build_long_table(attachments, dates)
    frame_to_csv(mapping_frame, OUT / 'source_interval_mapping.csv')

    inputs = ['code/14_q2_ridge_forecast_experiment.py',
              'code/19_q2_lightgbm_residual_experiment.py',
              'code/21_q2_quantile_level_scan.py', 'code/29_q2_time_mapping_experiment.py',
              'code/q1_start_time_recompute.py',
              '附件/附件1.xlsx', '附件/附件2.xlsx',
              'reports/问题二_统一新时间映射重跑实验方案.md',
              'reports/时间口径变更与跨问题影响_20260912.md']
    hashes = {rel: digest(ROOT / rel) for rel in inputs}
    dependency = dict(first_load_feature_day=first_feature_day('load'),
                      first_pv_feature_day=first_feature_day('pv'),
                      quantile_position_28=int(scan().quantile_index(28, ALPHA)),
                      quantile_position_27=int(scan().quantile_index(27, ALPHA)),
                      lgbm_params_sha256=hashlib.sha256(json.dumps(
                          lgbm_module().LGBM_PARAMS, sort_keys=True).encode()).hexdigest(),
                      lgbm_param_keys=sorted(lgbm_module().LGBM_PARAMS.keys()))
    signature = hashlib.sha256(json.dumps(
        dict(parameters=PARAMETERS, dependency=dependency, input_sha256=hashes,
             code=digest(CODE_29)), sort_keys=True).encode()).hexdigest()
    registration_path = OUT / 'registration.json'
    if args.mode == 'register' or not registration_path.exists():
        record = dict(registered_utc=datetime.now(timezone.utc).isoformat(),
                      specification='reports/问题二_统一新时间映射重跑实验方案.md',
                      time_version=TIME_VERSION, parameters=PARAMETERS, dependency=dependency,
                      input_sha256=hashes, code_sha256=digest(CODE_29), executable=sys.executable,
                      python=sys.version, numpy=np.__version__, pandas=pd.__version__)
        if registration_path.exists():
            existing = json.loads(registration_path.read_text(encoding='utf-8'))
            if existing['signature'] != signature:
                assert args.amend_reason, 'a different registration exists; pass --amend-reason'
                for key in ('time_version', 'interval_rule', 'natural_day_rule', 'price_rule',
                            'template_rule', 'targets', 'midnight_target_rule', 'alpha',
                            'residual_window_days', 'train_window_days', 'min_train_days',
                            'quantile_rule', 'load_features', 'pv_features', 'groups',
                            'comparison', 'solver', 'ledgers'):
                    assert existing['parameters'][key] == PARAMETERS[key], key
                amendments = existing.get('amendments', [])
                amendments.append(dict(amended_utc=datetime.now(timezone.utc).isoformat(),
                                       previous_signature=existing['signature'],
                                       previous_code_sha256=existing.get('code_sha256'),
                                       new_signature=signature,
                                       new_code_sha256=digest(CODE_29),
                                       reason=args.amend_reason, frozen_spec_unchanged=True))
                record['amendments'] = amendments
        record['signature'] = signature
        save(registration_path, record)
        snapshot = OUT / 'source_archive'
        (snapshot / 'code').mkdir(parents=True, exist_ok=True)
        for rel in inputs:
            if rel.startswith('code/'):
                shutil.copy2(ROOT / rel, snapshot / 'code' / Path(rel).name)
        shutil.copy2(PLAN_MD, snapshot / 'experiment_plan.md')
        shutil.copy2(TIME_REPORT, snapshot / 'time_change_report.md')
        print(f'registered: signature={signature[:16]}', flush=True)
        if args.mode == 'register':
            return
    else:
        existing = json.loads(registration_path.read_text(encoding='utf-8'))
        assert existing['signature'] == signature, (
            'inputs or code changed since registration; rerun --mode register --amend-reason')

    protected = protected_manifest()
    save(OUT / 'protected_before.json', protected)
    timings = {}

    mapping = mapping_checks(attachments, dates, nat_load, nat_pv, price_src)
    save(OUT / 'mapping_checks.json', mapping)
    assert mapping['all_passed'], mapping
    artificial = artificial_mapping_check(attachments, dates, nat_load, nat_pv)
    save(OUT / 'artificial_mapping.json', artificial)
    assert artificial['all_passed'], artificial
    print(f"mapping: {mapping['source_interval_count']} source intervals, bijection and "
          f"145-target identity exact; first {mapping['first_interval']}, "
          f"last {mapping['last_interval']}", flush=True)

    tick = time.perf_counter()
    archive = build_forecast_archive(nat_load, nat_pv, dates, load_src, pv_src, save_models=True,
                                     verbose=True,
                                     model_days=set(range(0, NATURAL_DAYS, 30))
                                     | {23, 24, 31, 171, 354})
    timings['forecast_seconds'] = time.perf_counter() - tick
    frame_to_csv(pd.DataFrame(archive['fit_log']), OUT / 'forecast_fit_log.csv')
    save(OUT / 'forecast_summary.json', archive['summary'])
    formal_ok = archive['summary']['formal_finite'] and archive['summary']['formal_nonnegative']
    assert formal_ok, archive['summary']
    print(f"forecast archive: load first trainable day "
          f"{archive['summary']['load_first_trainable_day']} "
          f"({dates[archive['summary']['load_first_trainable_day']].date()}), pv "
          f"{archive['summary']['pv_first_trainable_day']} "
          f"({dates[archive['summary']['pv_first_trainable_day']].date()})", flush=True)

    net_forecast = dict(issued=build_net_forecast(archive['issued_load'], archive['issued_pv']),
                        naive=build_net_forecast(archive['naive_load'], archive['naive_pv']))
    errors = errors_145(load_src, pv_src, nat_load, nat_pv, net_forecast['issued'])
    protection, m_archive, reasons = protection_archive(errors)
    np.savez_compressed(OUT / 'archive_float.npz', issued_load=archive['issued_load'],
                        issued_pv=archive['issued_pv'], errors=errors, protection=protection,
                        m_archive=m_archive)
    for name, array in (('issued_load', archive['issued_load']),
                        ('issued_pv', archive['issued_pv']),
                        ('naive_load', archive['naive_load']), ('naive_pv', archive['naive_pv']),
                        ('gate', archive['gate'])):
        np.save(OUT / f'{name}.npy', array)
    quantile = dict(position_28=int(scan().quantile_index(28, ALPHA)),
                    position_27=int(scan().quantile_index(27, ALPHA)),
                    h144_m_values=sorted(set(int(v)
                                             for v in m_archive[EVAL_FIRST_DAY:, TARGETS - 1])),
                    other_m_values=sorted(set(int(v)
                                              for v in m_archive[EVAL_FIRST_DAY:,
                                                                 :TARGETS - 1].ravel())),
                    insufficient_reason_count=len(reasons))
    quantile['all_passed'] = bool(quantile['h144_m_values'] == [27]
                                  and quantile['other_m_values'] == [28]
                                  and quantile['position_27'] == 22
                                  and quantile['position_28'] == 23)
    save(OUT / 'quantile_validation.json', quantile)
    assert quantile['all_passed'], quantile
    print(f"protection: h=144 uses m={quantile['h144_m_values']} position "
          f"{quantile['position_27']}, all other targets m={quantile['other_m_values']} "
          f"position {quantile['position_28']}", flush=True)

    forecasts = forecast_checks(archive, nat_pv, nat_load, dates, price_src)
    save(OUT / 'forecast_checks.json', forecasts)
    assert forecasts['all_passed'], forecasts
    boundaries = mode_boundary_checks(price_src)
    save(OUT / 'boundary_checks.json', boundaries)
    assert boundaries['all_passed'], boundaries
    print(f"boundary and forecast checks: {len(boundaries['rounds'])} + "
          f"{len(forecasts['rounds'])} rounds passed", flush=True)

    results = []
    reference = {}
    tick = time.perf_counter()
    fixed = run_group(GROUP_BY_ID['N_fixed6000'], nat_load, nat_pv, dates, price_src,
                      net_forecast, protection)
    timings['scenario_N_fixed6000_seconds'] = time.perf_counter() - tick
    results.append(fixed)
    reference['N_fixed6000'] = ledger_totals(fixed)
    print(f"  N_fixed6000: ledger A "
          f"{reference['N_fixed6000']['ledger_A']['total_cost_yuan']:,.6f} CNY, "
          f"ledger B {reference['N_fixed6000']['ledger_B']['total_cost_yuan']:,.6f} CNY",
          flush=True)
    if args.mode == 'repro':
        replay = reproduction_check(GROUP_BY_ID['N_fixed6000'], nat_load, nat_pv, dates,
                                    price_src, net_forecast, protection, fixed)
        save(OUT / 'internal_reproduction.json', replay)
        assert replay['all_passed'], replay
        print(f'  internal reproduction passed: max segment '
              f'{replay["max_segment_difference_kWh"]:.3e} kWh, cost '
              f'{replay["cost_difference_yuan"]:.3e} CNY', flush=True)
        print(f'wall seconds = {time.perf_counter() - started:.1f}', flush=True)
        return

    tick = time.perf_counter()
    free = run_group(GROUP_BY_ID['N_free'], nat_load, nat_pv, dates, price_src, net_forecast,
                     protection)
    timings['scenario_N_free_seconds'] = time.perf_counter() - tick
    results.append(free)
    reference['N_free'] = ledger_totals(free)
    print(f"  N_free: ledger A {reference['N_free']['ledger_A']['total_cost_yuan']:,.6f} CNY, "
          f"ledger B {reference['N_free']['ledger_B']['total_cost_yuan']:,.6f} CNY", flush=True)

    tick = time.perf_counter()
    replay = reproduction_check(GROUP_BY_ID['N_fixed6000'], nat_load, nat_pv, dates, price_src,
                                net_forecast, protection, fixed)
    timings['reproduction_seconds'] = time.perf_counter() - tick
    save(OUT / 'internal_reproduction.json', replay)
    assert replay['all_passed'], replay
    print(f"  internal reproduction: max segment "
          f"{replay['max_segment_difference_kWh']:.3e} kWh, cost "
          f"{replay['cost_difference_yuan']:.3e} CNY", flush=True)

    carry = carry_check(results, nat_load, nat_pv)
    save(OUT / 'carry_check.json', carry)
    assert carry['all_passed'], carry
    halfday = halfday_check(results, nat_load, nat_pv, dates, price_src)
    save(OUT / 'halfday_check.json', halfday)
    assert halfday['all_passed'], halfday
    tick = time.perf_counter()
    perturbation = perturbation_checks(results, nat_load, nat_pv, dates, price_src, load_src,
                                       pv_src, net_forecast, protection)
    timings['perturbation_seconds'] = time.perf_counter() - tick
    save(OUT / 'perturbation_checks.json', perturbation)
    assert perturbation['all_passed'], perturbation
    print(f"  carry, half-day and {len(perturbation['cases'])} perturbation cases passed",
          flush=True)

    summary = summary_frame(results)
    frame_to_csv(summary, OUT / 'summary.csv')
    contrast = pd.DataFrame([contrast_row(results)])
    frame_to_csv(contrast, OUT / 'contrast.csv')
    daily = daily_contrast_frame(results)
    frame_to_csv(daily, OUT / 'daily_contrasts.csv')
    monthly = monthly_contrast_frame(results)
    frame_to_csv(monthly, OUT / 'monthly_contrasts.csv')
    window = window_frame(results, price_src)
    frame_to_csv(window, OUT / 'window_emergency.csv')
    stability = stability_frame(results)
    frame_to_csv(stability, OUT / 'stability.csv')
    inventory = inventory_frame(results, price_src)
    frame_to_csv(inventory, OUT / 'inventory.csv')
    energy = energy_frame(results)
    frame_to_csv(energy, OUT / 'energy_contrasts.csv')
    for result in results:
        folder = OUT / result['group']['id']
        folder.mkdir(parents=True, exist_ok=True)
        frame_to_csv(result['natural'], folder / 'natural_dispatch.csv')
        frame_to_csv(result['template'], folder / 'template_plan.csv')
        frame_to_csv(result['solver'], folder / 'solver_log.csv')
        save(folder / 'ledger.json', ledger_totals(result))
        save(folder / 'validation.json', dict(physical=result['physical'],
                                              final_state_kWh=result['final_state'],
                                              group=result['group']))
    save(OUT / 'all_ledgers.json', reference)
    frame_to_csv(pd.DataFrame([dict(strategy_id=r['group']['id'], **r['tail']) for r in results]),
                 OUT / 'tail_interval.csv')

    tick = time.perf_counter()
    soc = soc_frame(results)
    produced = figures(summary, monthly, soc)
    timings['figures_seconds'] = time.perf_counter() - tick
    integrity = parent().figure_integrity(produced)
    save(OUT / 'figure_integrity.json', integrity)
    workbooks = write_validation_workbooks(results, mapping_frame)

    protected_after = protected_manifest()
    changed = sorted(name for name in protected if protected_after.get(name) != protected[name])
    missing = sorted(name for name in protected if name not in protected_after)
    save(OUT / 'protected_after.json', protected_after)

    physical_max = max(max(v.values()) for v in (r['physical'] for r in results))
    checks = dict(mapping=mapping, artificial_mapping=artificial, forecast=forecasts,
                  boundary=boundaries, quantile=quantile, reproduction=replay, carry=carry,
                  halfday=halfday, perturbation=perturbation, ledger=reference,
                  energy=energy.to_dict('records'), physical_max=float(physical_max),
                  protected_unchanged=bool(not changed and not missing),
                  protected_changed=changed, protected_missing=missing,
                  protected_count=len(protected),
                  thresholds=dict(energy_kWh=ENERGY_TOL_KWH, cost_yuan=COST_TOL_YUAN),
                  timings=timings,
                  workbooks=[str(p.relative_to(ROOT)) for p in workbooks])
    save(OUT / 'checks.json', checks)

    # the experiment writes data only; the report is rendered by code/30 (report layer),
    # so editing prose never invalidates the registration signature or forces a re-run
    artifact_hashes = {path.relative_to(OUT).as_posix(): digest(path)
                       for path in sorted(OUT.rglob('*'))
                       if path.is_file() and path.name not in ('artifact_hashes.json',
                                                               'run_manifest.json')}
    save(OUT / 'artifact_hashes.json', artifact_hashes)
    manifest = dict(started_utc=started_utc,
                    finished_utc=datetime.now(timezone.utc).isoformat(),
                    wall_seconds=time.perf_counter() - started, mode=args.mode,
                    time_version=TIME_VERSION, executable=sys.executable, python=sys.version,
                    numpy=np.__version__, pandas=pd.__version__, signature=signature,
                    input_sha256=hashes, parameters=PARAMETERS, dependency=dependency,
                    protected_before_count=len(protected),
                    protected_unchanged=bool(not changed and not missing),
                    ledgers=reference, contrast=contrast.iloc[0].to_dict(),
                    forecast_summary=archive['summary'], quantile=quantile,
                    self_checks_passed=dict(mapping=mapping['all_passed'],
                                            artificial=artificial['all_passed'],
                                            forecast=forecasts['all_passed'],
                                            boundary=boundaries['all_passed'],
                                            quantile=quantile['all_passed'],
                                            reproduction=replay['all_passed'],
                                            carry=carry['all_passed'],
                                            halfday=halfday['all_passed'],
                                            perturbation=perturbation['all_passed']),
                    timings=timings, artifact_count=len(artifact_hashes),
                    report=str(REPORT_MD.relative_to(ROOT)))
    save(OUT / 'run_manifest.json', manifest)
    assert manifest['protected_unchanged'], (changed, missing)
    print(summary[['strategy_id', 'total_cost_yuan', 'B_total_cost_yuan', 'emergency_kWh',
                   'final_kWh']].to_string(index=False), flush=True)
    print(json.dumps(jsonable(contrast.iloc[0].to_dict()), ensure_ascii=False, indent=2),
          flush=True)
    print(f'self-checks {manifest["self_checks_passed"]}', flush=True)
    print(f'wall seconds = {time.perf_counter() - started:.1f}', flush=True)
    return manifest


if __name__ == '__main__':
    main()
