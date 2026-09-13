"""Finalize Question 1 under the accepted interval-start time convention.

Delivery-only operation: no optimization is run. The script reads the verified
``start_time_v1`` payload and patches numeric cells directly into the pristine
OOXML package, preserving every non-worksheet package part (including printer
settings, relationships and shared strings). It backs up the legacy formal file,
atomically replaces the formal workbook, and writes a read-back validation record.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
FORMAL = ROOT / "附件/附件5/result1.xlsx"
PRISTINE = ROOT / "results/template_backups/result1_original.xlsx"
OLD_BACKUP = ROOT / "results/template_backups/result1_old_time_before_start_time_v1.xlsx"
PAYLOAD = ROOT / "results/q1_start_time_20260912/workbook_payload.json"
NUMERIC_VALIDATION = ROOT / "results/q1_start_time_20260912/validation.json"
OUT_DIR = ROOT / "results/q1_start_time_20260912"
VERIFIED_COPY = OUT_DIR / "result1_formal_verified.xlsx"
DELIVERY_VALIDATION = OUT_DIR / "formal_delivery_validation.json"
TMP = FORMAL.with_name("result1.tmp.xlsx")
TOL = 1e-9
SHEET_PARTS = {"xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def num(value: float) -> str:
    return format(float(value), ".17g")


def column_number(coord: str) -> int:
    letters = re.match(r"[A-Z]+", coord).group(0)
    result = 0
    for ch in letters:
        result = result * 26 + ord(ch) - 64
    return result


def set_numeric_cell(xml: str, coord: str, value: float) -> str:
    """Set an existing or missing numeric cell without reserializing the sheet."""
    self_re = re.compile(rf'<c\b([^>]*\br="{re.escape(coord)}"[^>]*)/>')
    match = self_re.search(xml)
    if match:
        attrs = match.group(1).rstrip('/')
        attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
        replacement = f'<c{attrs}><v>{num(value)}</v></c>'
        return xml[:match.start()] + replacement + xml[match.end():]

    full_re = re.compile(rf'<c\b([^>]*\br="{re.escape(coord)}"[^>]*)>(.*?)</c>')
    match = full_re.search(xml)
    if match:
        attrs = re.sub(r'\s+t="[^"]*"', "", match.group(1))
        replacement = f'<c{attrs}><v>{num(value)}</v></c>'
        return xml[:match.start()] + replacement + xml[match.end():]

    row_number = int(re.search(r"\d+", coord).group(0))
    row_re = re.compile(rf'(<row\b[^>]*\br="{row_number}"[^>]*>)(.*?)(</row>)')
    row_match = row_re.search(xml)
    if not row_match:
        raise AssertionError(f"row {row_number} not found for {coord}")
    body = row_match.group(2)
    new_cell = f'<c r="{coord}"><v>{num(value)}</v></c>'
    target_col = column_number(coord)
    insertion = len(body)
    for existing in re.finditer(r'<c\b[^>]*\br="([A-Z]+\d+)"[^>]*(?:/>|>.*?</c>)', body):
        if column_number(existing.group(1)) > target_col:
            insertion = existing.start()
            break
    new_body = body[:insertion] + new_cell + body[insertion:]
    return xml[:row_match.start()] + row_match.group(1) + new_body + row_match.group(3) + xml[row_match.end():]


def build_exact_candidate(payload: dict) -> None:
    sheet_values = {
        "xl/worksheets/sheet1.xml": {
            **{f"B{r}": payload["purchases"][r - 2][0] for r in range(2, 146)},
        },
        "xl/worksheets/sheet2.xml": {
            **{f"B{r}": payload["storage"][r - 2][0] for r in range(2, 8)},
            **{f"C{r}": payload["storage"][r - 2][1] for r in range(2, 8)},
            "E2": payload["endpoints"][0][0],
            "E3": payload["endpoints"][1][0],
        },
    }
    if TMP.exists():
        TMP.unlink()
    with ZipFile(PRISTINE, "r") as source, ZipFile(TMP, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename in sheet_values:
                text = data.decode("utf-8")
                for coord, value in sheet_values[info.filename].items():
                    text = set_numeric_cell(text, coord, value)
                data = text.encode("utf-8")
            target.writestr(info, data)


def package_preservation(candidate: Path) -> dict:
    with ZipFile(PRISTINE, "r") as original, ZipFile(candidate, "r") as filled:
        original_names = original.namelist()
        filled_names = filled.namelist()
        same_members = original_names == filled_names
        changed = []
        for name in original_names:
            if original.read(name) != filled.read(name):
                changed.append(name)
        non_sheet_parts_identical = all(name in SHEET_PARTS for name in changed)
        required = {
            "xl/sharedStrings.xml",
            "xl/printerSettings/printerSettings1.bin",
            "xl/worksheets/_rels/sheet2.xml.rels",
        }
        return {
            "member_list_equal_to_pristine": same_members,
            "member_count": len(filled_names),
            "changed_members": changed,
            "only_worksheet_xml_changed": non_sheet_parts_identical,
            "required_printer_and_shared_string_parts_present": required.issubset(filled_names),
        }


def workbook_structure(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=False)
    result = {"sheetnames": wb.sheetnames, "sheets": []}
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for cell in row:
                cells.append(
                    (cell.coordinate, cell.style_id, cell.number_format, str(cell.font),
                     str(cell.fill), str(cell.border), str(cell.alignment), str(cell.protection))
                )
        result["sheets"].append(
            {
                "title": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "merged": sorted(str(x) for x in ws.merged_cells.ranges),
                "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
                "row_heights": {str(i): d.height for i, d in ws.row_dimensions.items() if d.height is not None},
                "column_widths": {str(i): d.width for i, d in ws.column_dimensions.items() if d.width is not None},
                "cells": cells,
            }
        )
    wb.close()
    return result


def max_abs_error(actual, expected) -> float:
    return max(abs(float(a) - float(e)) for a, e in zip(actual, expected))


def read_values(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    pws, sws = wb.worksheets
    data = {
        "labels": [pws.cell(r, 1).value for r in range(2, 146)],
        "purchases": [pws.cell(r, 2).value for r in range(2, 146)],
        "storage": [sws.cell(r, c).value for r in range(2, 8) for c in (2, 3)],
        "endpoints": [sws.cell(2, 5).value, sws.cell(3, 5).value],
    }
    wb.close()
    return data


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    OLD_BACKUP.parent.mkdir(parents=True, exist_ok=True)
    payload = json.loads(PAYLOAD.read_text(encoding="utf-8"))
    numeric = json.loads(NUMERIC_VALIDATION.read_text(encoding="utf-8"))
    assert numeric["periodic_midnight_assumption"] is True
    assert numeric["physical_day"] == "00:00 to 24:00"
    assert abs(float(numeric["initial_kWh"]) - 6000.0) <= TOL
    assert abs(float(numeric["final_kWh"]) - 6000.0) <= TOL
    assert len(payload["purchases"]) == 144 and len(payload["storage"]) == 6 and len(payload["endpoints"]) == 2

    formal_before_run_hash = sha256(FORMAL)
    if not OLD_BACKUP.exists():
        shutil.copy2(FORMAL, OLD_BACKUP)
    legacy_formal_hash = sha256(OLD_BACKUP)

    pristine_values = read_values(PRISTINE)
    assert pristine_values["labels"] == payload["labels"]
    assert pristine_values["labels"][0] == "0:10-0:20"
    assert pristine_values["labels"][-2:] == ["23:50-0:00+1", "0:00+1-0:10+1"]

    build_exact_candidate(payload)
    candidate = read_values(TMP)
    purchases_expected = [row[0] for row in payload["purchases"]]
    storage_expected = [x for pair in payload["storage"] for x in pair]
    endpoints_expected = [row[0] for row in payload["endpoints"]]
    assert candidate["labels"] == pristine_values["labels"]
    purchase_error = max_abs_error(candidate["purchases"], purchases_expected)
    storage_error = max_abs_error(candidate["storage"], storage_expected)
    endpoint_error = max_abs_error(candidate["endpoints"], endpoints_expected)
    assert max(purchase_error, storage_error, endpoint_error) <= TOL

    structure_equal = workbook_structure(TMP) == workbook_structure(PRISTINE)
    package = package_preservation(TMP)
    assert structure_equal
    assert package["member_list_equal_to_pristine"]
    assert package["only_worksheet_xml_changed"]
    assert package["required_printer_and_shared_string_parts_present"]

    # Replace when the exact-package candidate differs; subsequent reruns are byte-idempotent.
    candidate_hash = sha256(TMP)
    formal_replaced = candidate_hash != formal_before_run_hash
    if formal_replaced:
        os.replace(TMP, FORMAL)
    else:
        TMP.unlink()
    shutil.copy2(FORMAL, VERIFIED_COPY)
    formal_hash = sha256(FORMAL)
    assert formal_hash == sha256(VERIFIED_COPY) == candidate_hash

    report = {
        "status": "complete",
        "operation": "formal Question 1 delivery switched to start_time_v1 without re-solving",
        "time_version": "start_time_v1",
        "interval_rule": "source label 00:10 covers 00:10-00:20; final source label 0:00+1 covers next-day 00:00-00:10",
        "physical_boundary_rule": "battery state is 6000 kWh at true 00:00 and true 24:00",
        "template_rule": "purchase rows retain the official 00:10-to-next-day-00:10 labels and order",
        "formal_workbook": str(FORMAL),
        "pristine_template": str(PRISTINE),
        "old_formal_backup": str(OLD_BACKUP),
        "verified_copy": str(VERIFIED_COPY),
        "formal_replaced_this_run": formal_replaced,
        "hashes": {
            "pristine_template_sha256": sha256(PRISTINE),
            "formal_before_run_sha256": formal_before_run_hash,
            "old_formal_sha256": legacy_formal_hash,
            "old_backup_sha256": sha256(OLD_BACKUP),
            "formal_sha256": formal_hash,
            "verified_copy_sha256": sha256(VERIFIED_COPY),
            "payload_sha256": sha256(PAYLOAD),
            "numeric_validation_sha256": sha256(NUMERIC_VALIDATION),
            "script_sha256": sha256(Path(__file__)),
        },
        "readback": {
            "planned_cells": 144,
            "storage_cells": 12,
            "endpoint_cells": 2,
            "total_filled_result_cells": 158,
            "first_label": candidate["labels"][0],
            "penultimate_label": candidate["labels"][-2],
            "last_label": candidate["labels"][-1],
            "purchase_max_abs_error": purchase_error,
            "storage_max_abs_error": storage_error,
            "endpoint_max_abs_error": endpoint_error,
            "style_and_structure_equal_to_pristine_template": structure_equal,
            "ooxml_package_preservation": package,
        },
        "numeric_result": {
            "grid_kWh": numeric["grid_kWh"],
            "cost_yuan": numeric["cost_yuan"],
            "initial_kWh": numeric["initial_kWh"],
            "final_kWh": numeric["final_kWh"],
            "state_at_00_10_kWh": numeric["template_period_start_kWh"],
            "mip_gap": numeric["mip_gap"],
            "lp_gap_yuan": numeric["lp_gap_yuan"],
        },
    }
    DELIVERY_VALIDATION.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
