"""Causal Q2 backtest; run with conda run -n math_modeling python code/05_q2_baseline.py.

Inputs: attachment 1 fixed prices and attachment 2 historical load/PV only.
Outputs: results/q2_baseline, reports/Q2_Baseline.md, figures/q2_baseline.png.
Policies and windows are fixed before evaluation; no full-year parameter fitting.
"""
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import importlib.util
import json
import sys

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.comments import Comment
import scipy

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'results/q2_baseline'
spec = importlib.util.spec_from_file_location('q1_core', ROOT/'code/02_q1_baseline.py')
core = importlib.util.module_from_spec(spec)
spec.loader.exec_module(core)
CONFIG = dict(dt_hours=1/6, eta_c=.9, eta_d=.9, state_min_kWh=1200., state_max_kWh=10800.,
              initial_kWh=6000., power_max_kW=5000., emergency_multiplier=5.,
              pv_mean_days=7, same_weekday_days=4, tolerance=1e-6,
              evaluation_start='2025-02-01', seed=20260910,
              cold_start='First day: zero plan, battery idle, PV serves load, emergency fills deficit.',
              terminal='Nominal end-of-day state equals actual beginning state; actual state carries forward.',
              actual_control='Contemporaneous representative power and ideal fast feedback; no future samples.',
              scenarios=['lag','mean','mean_no_storage','mean_terminal6000'])
NAMES = {'lag':'上周负载+昨日光伏', 'mean':'同星期4日均值+光伏7日均值',
         'mean_no_storage':'均值预测/无储能', 'mean_terminal6000':'均值预测/预测末态6000'}
DATES = pd.date_range('2025-01-01','2025-12-31')


def read_sources():
    audit = json.loads((ROOT/'results/audit/audit_summary.json').read_text(encoding='utf-8'))
    hashes = {}
    for name in ['附件1.xlsx','附件2.xlsx']:
        path = ROOT/'附件'/name
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        relative = str(path.relative_to(ROOT))
        audit_hashes = audit['source_sha256']
        expected = audit_hashes.get(relative)
        if expected is None:  # audit JSON may store Windows separators
            expected = audit_hashes.get(relative.replace('/', chr(92)))
        assert digest == expected, relative
        hashes[name] = digest
    wb = openpyxl.load_workbook(ROOT/'附件/附件1.xlsx',read_only=True,data_only=True)
    raw = list(wb.worksheets[0].values)
    prices = np.array([r[1] for r in raw[1:]],dtype=float)
    wb.close()
    wb = openpyxl.load_workbook(ROOT/'附件/附件2.xlsx',read_only=True,data_only=True)
    arrays = []
    for ws in wb.worksheets:
        rows = list(ws.values)
        assert pd.DatetimeIndex([r[0] for r in rows[1:]]).equals(DATES)
        def minute(value):
            text = str(value)
            hour,minute = text.replace('+1','').split(':')[:2]
            return int(hour)*60+int(minute)+(1440 if '+1' in text else 0)
        assert [minute(v) for v in rows[0][1:]] == [minute(r[0]) for r in raw[1:]] == list(range(10,1441,10))
        arrays.append(np.asarray([r[1:] for r in rows[1:]],dtype=float))
    wb.close()
    assert prices.shape == (144,) and (prices>0).all()
    assert all(a.shape==(365,144) and np.isfinite(a).all() and (a>=0).all() for a in arrays)
    return arrays[0],arrays[1],prices,hashes


def predict(history_load, history_pv, family):
    count = len(history_load)
    assert count == len(history_pv) and count>0
    if family == 'lag':
        ids = np.array([count-7 if count>=7 else count-1])
        pv_ids = np.array([count-1])
    else:
        ids = np.arange(count-7,-1,-7)[:CONFIG['same_weekday_days']]
        if not len(ids):
            ids = np.arange(count)
        pv_ids = np.arange(max(0,count-CONFIG['pv_mean_days']),count)
    assert max(ids.max(),pv_ids.max()) < count
    return history_load[ids].mean(axis=0),history_pv[pv_ids].mean(axis=0),int(max(ids.max(),pv_ids.max()))


def control(plan, load, pv, initial, battery=True):
    q = np.asarray(plan)
    c,d,e,w = [np.zeros(len(q)) for _ in range(4)]
    states = np.r_[initial,np.zeros(len(q))]
    cap = CONFIG['power_max_kW']*CONFIG['dt_hours'] if battery else 0.
    for t in range(len(q)):
        surplus = q[t]+(pv[t]-load[t])*CONFIG['dt_hours']
        if surplus>=0:
            c[t] = min(surplus,cap,max(0,(CONFIG['state_max_kWh']-states[t])/CONFIG['eta_c']))
            w[t] = surplus-c[t]
        else:
            d[t] = min(-surplus,cap,max(0,(states[t]-CONFIG['state_min_kWh'])*CONFIG['eta_d']))
            e[t] = -surplus-d[t]
        states[t+1] = states[t]+CONFIG['eta_c']*c[t]-d[t]/CONFIG['eta_d']
    balance = q+pv*CONFIG['dt_hours']+d+e-load*CONFIG['dt_hours']-c-w
    checks = dict(balance_kWh=float(np.abs(balance).max()),
                  state_recursion_kWh=float(np.abs(np.diff(states)-CONFIG['eta_c']*c+d/CONFIG['eta_d']).max()),
                  state_bounds_kWh=float(max(0,CONFIG['state_min_kWh']-states.min(),states.max()-CONFIG['state_max_kWh'])),
                  power_bounds_kWh=float(max(0,c.max()-cap,d.max()-cap)),
                  nonnegative_kWh=float(max(0,-min(q.min(),c.min(),d.min(),e.min(),w.min()))),
                  simultaneous_kWh=float(np.minimum(c,d).max()))
    assert max(checks.values())<CONFIG['tolerance'],checks
    return c,d,e,w,states,checks


def validate_core(load,pv,price):
    # Analytic boundary checks cover exhausted/full batteries and power-limited delivery.
    x = control(np.array([0.]),np.array([12000.]),np.array([0.]),1200.)
    assert x[1][0] == 0 and abs(x[2][0]-2000)<1e-9
    x = control(np.array([0.]),np.array([12000.]),np.array([0.]),10800.)
    assert abs(x[1][0]-5000/6)<1e-9 and abs(x[2][0]-(2000-5000/6))<1e-9
    x = control(np.array([2000.]),np.array([0.]),np.array([0.]),10800.)
    assert x[0][0] == 0 and x[3][0] == 2000
    a = predict(load[:31],pv[:31],'mean')
    corrupted_load,corrupted_pv = load.copy(),pv.copy()
    corrupted_load[31:]=1e9
    corrupted_pv[31:]=0
    b = predict(corrupted_load[:31],corrupted_pv[:31],'mean')
    assert np.array_equal(a[0],b[0]) and np.array_equal(a[1],b[1])
    sample = pd.read_csv(ROOT/'results/audit/q1_normalized.csv')
    _,q1 = core.solve(sample.load_kW.to_numpy(),sample.pv_kW.to_numpy(),price,integer=True)
    locked = json.loads((ROOT/'results/q1_milp/validation_summary.json').read_text(encoding='utf-8'))['summary']
    assert abs(q1['cost_yuan']-locked['cost_yuan'])<1e-6
    q,nominal = core.solve(load[0],pv[0],price,integer=True,initial_kWh=1200.)
    assert nominal['initial_kWh']==nominal['final_kWh']==1200
    c,d,e,w,state,_ = control(q[0],load[0],pv[0],1200.)
    assert e.max()<1e-6
    # A controller's prefix cannot depend on a later load observation.
    modified = load[0].copy()
    modified[72:]*=2
    other = control(q[0],modified,pv[0],1200.)
    assert np.array_equal(state[:73],other[4][:73])
    return dict(boundary_cases=True,forecast_history_only=True,controller_prefix_invariance=True,
                q1_regression=True,arbitrary_initial_state=True,perfect_forecast_no_emergency=True)


def event_rows(date,emergency):
    active = emergency>1e-6
    starts = np.flatnonzero(active & ~np.r_[False,active[:-1]])
    stops = np.flatnonzero(active & ~np.r_[active[1:],False])+1
    return [dict(date=str(date.date()),interval=f'{core.label(int(s)*10)}-{core.label(int(t)*10)}',
                 start_slot=int(s),end_slot=int(t),emergency_kWh=float(emergency[s:t].sum()))
            for s,t in zip(starts,stops)]


def run_policy(key,load,pv,price):
    state=CONFIG['initial_kWh']
    records,daily,events=[],[],[]
    checks={}
    family='lag' if key=='lag' else 'mean'
    for day,date in enumerate(DATES):
        initial=state
        if day==0:
            pred_l,pred_v=np.full(144,np.nan),np.full(144,np.nan)
            q=np.zeros(144)
            reference=np.full(145,initial)
            solver=dict(elapsed_seconds=0.,mip_gap=0.,checks={})
            history_end=pd.NaT
        else:
            pred_l,pred_v,last_id=predict(load[:day],pv[:day],family)
            history_end=DATES[last_id]+pd.Timedelta(days=1)
            assert history_end<=date
            if key=='mean_no_storage':
                q=np.maximum(pred_l-pred_v,0)*CONFIG['dt_hours']
                reference=np.full(145,initial)
                solver=dict(elapsed_seconds=0.,mip_gap=0.,checks={})
            else:
                terminal=6000. if key=='mean_terminal6000' else initial
                x,solver=core.solve(pred_l,pred_v,price,integer=True,initial_kWh=initial,terminal_kWh=terminal)
                q=x[0].copy()
                reference=x[4]
        fixed_q=q.copy()
        c,d,e,w,states,physical=control(q,load[day],pv[day],initial,
                                       battery=day>0 and key!='mean_no_storage')
        assert np.array_equal(q,fixed_q)
        assert abs(states[0]-state)<1e-9
        state=float(states[-1])
        for name,value in physical.items():
            checks[name]=max(checks.get(name,0),value)
        for name,value in solver['checks'].items():
            checks['nominal_'+name]=max(checks.get('nominal_'+name,0),value)
        day_events=event_rows(date,e)
        events.extend(day_events)
        part=pd.DataFrame(dict(date=str(date.date()),issue_time=date,history_latest_end=history_end,
                    slot=np.arange(144),start_time=pd.date_range(date,periods=144,freq='10min'),
                    interval=[f'{core.label(t)}-{core.label(t+10)}' for t in range(0,1440,10)],
                    load_kW=load[day],pv_kW=pv[day],load_forecast_kW=pred_l,pv_forecast_kW=pred_v,
                    price_yuan_kWh=price,planned_kWh=q,charge_kWh=c,discharge_kWh=d,
                    emergency_kWh=e,unused_kWh=w,state_start_kWh=states[:-1],state_end_kWh=states[1:],
                    nominal_state_end_kWh=reference[1:],planned_cost_yuan=q*price,emergency_cost_yuan=e*5*price))
        records.append(part)
        daily.append(dict(date=str(date.date()),initial_kWh=initial,final_kWh=state,
                          planned_kWh=float(q.sum()),emergency_kWh=float(e.sum()),unused_kWh=float(w.sum()),
                          charge_kWh=float(c.sum()),discharge_kWh=float(d.sum()),
                          planned_cost_yuan=float(q@price),emergency_cost_yuan=float(5*e@price),
                          total_cost_yuan=float((q+5*e)@price),emergency_events=len(day_events),
                          solver_seconds=solver['elapsed_seconds'],mip_gap=solver['mip_gap']))
        if day in [30,120,240,364]:
            print(f'{key}: completed {date.date()}',flush=True)
    whole=pd.concat(records,ignore_index=True)
    daily=pd.DataFrame(daily)
    assert np.allclose(daily.initial_kWh.iloc[1:],daily.final_kWh.iloc[:-1],atol=1e-9,rtol=0)
    evaluated=whole[whole.date>=CONFIG['evaluation_start']].copy()
    assert len(evaluated)==334*144
    ds=daily[daily.date>=CONFIG['evaluation_start']]
    assert abs(evaluated.planned_cost_yuan.sum()+evaluated.emergency_cost_yuan.sum()-ds.total_cost_yuan.sum())<1e-6
    forecast_error=(evaluated.load_kW-evaluated.pv_kW)-(evaluated.load_forecast_kW-evaluated.pv_forecast_kW)
    totals={k:float(ds[k].sum()) for k in ['planned_kWh','emergency_kWh','unused_kWh','charge_kWh','discharge_kWh',
                                          'planned_cost_yuan','emergency_cost_yuan','total_cost_yuan']}
    totals.update(policy=key,emergency_events=int(ds.emergency_events.sum()),
                  emergency_days=int((ds.emergency_kWh>1e-6).sum()),
                  emergency_cost_share=totals['emergency_cost_yuan']/totals['total_cost_yuan'],
                  load_mae_kW=float((evaluated.load_kW-evaluated.load_forecast_kW).abs().mean()),
                  pv_mae_kW=float((evaluated.pv_kW-evaluated.pv_forecast_kW).abs().mean()),
                  net_mae_kW=float(forecast_error.abs().mean()),net_rmse_kW=float(np.sqrt((forecast_error**2).mean())),
                  net_bias_actual_minus_forecast_kW=float(forecast_error.mean()),
                  evaluation_initial_kWh=float(ds.initial_kWh.iloc[0]),final_kWh=float(ds.final_kWh.iloc[-1]),
                  max_mip_gap=float(ds.mip_gap.max()),solver_total_seconds=float(daily.solver_seconds.sum()),
                  worst_day=str(ds.loc[ds.total_cost_yuan.idxmax(),'date']))
    path=OUT/key
    path.mkdir(parents=True,exist_ok=True)
    whole.to_csv(path/'dispatch_all_2025.csv',index=False,encoding='utf-8-sig')
    daily.to_csv(path/'daily_summary.csv',index=False,encoding='utf-8-sig')
    ev=pd.DataFrame(events,columns=['date','interval','start_slot','end_slot','emergency_kWh'])
    ev=ev[ev.date>=CONFIG['evaluation_start']]
    ev.to_csv(path/'emergency_events.csv',index=False,encoding='utf-8-sig')
    assert abs(ev.emergency_kWh.sum()-totals['emergency_kWh'])<1e-5
    monthly=ds.assign(month=ds.date.str[:7]).groupby('month')[[
        'planned_cost_yuan','emergency_cost_yuan','total_cost_yuan','emergency_kWh','unused_kWh','emergency_events']].sum()
    monthly.to_csv(path/'monthly_summary.csv',encoding='utf-8-sig')
    errors=evaluated.assign(month=evaluated.date.str[:7],net_error_kW=forecast_error)
    errors.groupby('month').net_error_kW.agg(['mean','std','min','max']).to_csv(path/'monthly_net_errors.csv',encoding='utf-8-sig')
    peak=evaluated[evaluated.emergency_kWh>1e-6].copy()
    peak['hour']=peak.slot//6
    peak.groupby('hour')[['emergency_kWh','emergency_cost_yuan']].sum().to_csv(path/'emergency_by_hour.csv',encoding='utf-8-sig')
    (path/'validation.json').write_text(json.dumps(dict(checks=checks,totals=totals),ensure_ascii=False,indent=2),encoding='utf-8')
    return evaluated,ds,ev,monthly,totals


def export_excel(data,daily,events):
    wb=openpyxl.Workbook()
    plan=wb.active
    plan.title='计划购电量'
    labels=data.iloc[:144].interval.tolist()
    plan.append(['日期\\时间']+labels+['全天购电量（kWh）','全天计划购电费（元）'])
    storage=wb.create_sheet('充放电量')
    storage.append(['日期','时间段','充电量','放电量','时刻','储电量'])
    emergency=wb.create_sheet('紧急购电量')
    emergency.append(['日期','购电时间段','购电量'])
    for date,group in data.groupby('date',sort=True):
        stamp=datetime.fromisoformat(date)
        plan.append([stamp]+group.planned_kWh.tolist()+[float(group.planned_kWh.sum()),float(group.planned_cost_yuan.sum())])
        for block in range(6):
            segment=group.iloc[block*24:(block+1)*24]
            storage.append([stamp,f'{block*4:02d}:00-{(block+1)*4:02d}:00',
                            float(segment.charge_kWh.sum()),float(segment.discharge_kWh.sum()),
                            '0:00' if block==0 else ('24:00' if block==1 else None),
                            float(group.state_start_kWh.iloc[0]) if block==0 else (float(group.state_end_kWh.iloc[-1]) if block==1 else None)])
    for row in events.itertuples():
        emergency.append([datetime.fromisoformat(row.date),row.interval,row.emergency_kWh])
    note=wb.create_sheet('Baseline说明')
    note.append(['状态','首轮因果Baseline，非最终方案，未替换附件5原模板'])
    note.append(['预测','负载上周同段，光伏昨日同段；历史不足用昨日；首日零计划且储能静置'])
    note.append(['电价','仅附件1，每日重复；应急5倍，不用附件3或附件4'])
    note.append(['时间','右端点代表前10分钟区间；已修正原模板偏移'])
    note.append(['末态','预测末态=当日初态，实际状态跨日连续；初始6000 kWh'])
    note.append(['效率','单向各0.9；充放电为母线侧，储电量为电池内部'])
    note.append(['执行','当期供需快速反馈；富余充电、缺口放电、剩余应急；非全局最优控制'])
    note.append(['全天总费','计划表末列为计划费；含应急总费见每日费用账本'])
    ledger=wb.create_sheet('每日费用账本')
    ledger.append(daily.columns.tolist())
    for row in daily.itertuples(index=False,name=None):
        ledger.append(row)
    for ws in [plan,storage,emergency]:
        ws.freeze_panes='B2'
        ws.column_dimensions['A'].width=15
        for row in ws.iter_rows(min_row=2):
            row[0].number_format='yyyy/mm/dd'
            for cell in row[1:]:
                if isinstance(cell.value,(int,float)):
                    cell.number_format='0.000000'
    plan['A1'].comment=Comment('时间标签按前10分钟区间修正，计算口径见Baseline说明。','Modeling')
    target=OUT/'result2_baseline.xlsx'
    wb.save(target)
    wb.close()
    check=openpyxl.load_workbook(target,read_only=True,data_only=True)
    rows=list(check['计划购电量'].values)[1:]
    assert len(rows)==334 and np.allclose(np.array([r[1:145] for r in rows]),data.planned_kWh.to_numpy().reshape(334,144),atol=1e-9,rtol=0)
    stored=list(check['充放电量'].values)[1:]
    assert len(stored)==334*6
    assert abs(sum(r[2] for r in stored)-data.charge_kWh.sum())<1e-6
    assert abs(sum(r[3] for r in stored)-data.discharge_kWh.sum())<1e-6
    assert np.allclose([stored[i*6][5] for i in range(334)],daily.initial_kWh,atol=1e-9,rtol=0)
    assert np.allclose([stored[i*6+1][5] for i in range(334)],daily.final_kWh,atol=1e-9,rtol=0)
    assert abs(sum(r[2] for r in list(check['紧急购电量'].values)[1:])-data.emergency_kWh.sum())<1e-5
    check.close()


def make_report(runs):
    import matplotlib.pyplot as plt
    fig,axes=plt.subplots(2,1,figsize=(11,8),layout='constrained')
    for key in ['lag','mean','mean_terminal6000']:
        monthly=runs[key][3]
        axes[0].plot(np.arange(2,13),monthly.total_cost_yuan/1e6,marker='o',label=key)
    axes[0].set(title='Q2 causal backtest: monthly realized cost',xlabel='Month in 2025',ylabel='Cost (million CNY)',xticks=np.arange(2,13))
    daily=runs['lag'][1]
    axes[1].plot(pd.to_datetime(daily.date),daily.initial_kWh,label='Actual daily initial energy')
    axes[1].set(title='Lag baseline: continuous battery state',xlabel='Date',ylabel='Stored energy (kWh)')
    for ax in axes:
        ax.grid(alpha=.2)
        ax.legend()
    fig.savefig(ROOT/'figures/q2_baseline.png',dpi=160)
    plt.close(fig)
    lines=['# 第二问首轮Baseline与诊断','','状态：已完成因果Baseline全年运行及2—12月评价，尚未收敛为最终模型。', '',
           '## 1. 问题分析','',
           '日前不知道当天真实供需，计划多买仍付款，不足需5倍应急。比较的目标是实际计划费与应急费之和，不是只优化预测误差。', '',
           '## 2. 数据预处理','',
           '直接读取附件1价格与附件2两表，核对来源哈希、365天×144点、非负与时间一致性。仅使用这些输入，未读取附件3/4用于决策。样本代表前10分钟区间，功率乘1/6小时。', '',
           '主Baseline负载用上周同段、光伏用昨日同段；不足7天的负载回退昨日。候选均值方案负载取最近4个同星期日、光伏取最近7天；不足时仅使用已有历史。这些窗口为预先设定候选，不是全年拟合最优参数。每天0:00最多使用刚结束前一日24:00的观测，默认数据无额外发布延迟。', '',
           "### 2.1 具体预测什么",
           "",
           "每天0:00一次性预测当天144个10分钟时段的小区负载功率与光伏发电功率，共两条曲线、288个预测值，单位均为kW。令日期索引为$k$（2025-01-01为$k=0$），日内时段为$t=0,\\ldots,143$；$L_{k,t},V_{k,t}$表示实际功率，$\\widehat L_{k,t},\\widehat V_{k,t}$表示当日0:00发出的预测。$t=0$对应00:00—00:10，沿用右端点代表功率解释。",
           "",
           "只预测负载与光伏；电价使用附件1的已知重复日曲线，0:00储电量使用继承的实际状态。购电量和充放电量由后续MILP优化决定，应急量由实际执行时的供需缺口决定，不是本轮预测器的输出。预测模块尚未估计概率分布或预测区间。",
           "",
           "### 2.2 主Baseline：季节性朴素预测",
           "",
           "主Baseline（代码标识lag）负载采用以7天为周期的季节性朴素模型（seasonal naive），光伏采用以1天为周期的季节性朴素模型，也称日前日曲线持续性预测（daily-profile persistence）：",
           "",
           "$$\\widehat L_{k,t}=L_{k-7,t},\\qquad\\widehat V_{k,t}=V_{k-1,t}.$$",
           "",
           "负载公式表达同星期的日内曲线具有重复性；光伏公式假设昨日曲线可作为今天的初步估计，不能据此保证天气不变。光伏取昨日每个对应时段值，不是把昨日最后一个时点复制到整天。该模型不拟合回归系数，不使用ARIMA、神经网络或天气外生变量。",
           "",
           "### 2.3 对照：同星期均值与同一时段移动平均",
           "",
           "均值预测（代码标识mean）对负载使用最近4个同星期日的逐时段平均，对光伏使用最近7个已结束日期的逐时段移动平均。历史充足时为：",
           "",
           "$$\\widehat L_{k,t}=\\frac14\\sum_{j=1}^{4}L_{k-7j,t},\\qquad\\widehat V_{k,t}=\\frac17\\sum_{j=1}^{7}V_{k-j,t}.$$",
           "",
           "光伏的7日窗口是每天相同时段的7个样本，不是相邻的7个10分钟样本；负载均值跨4周，不是连续4天。两个平均模型都是简单的等权历史估计，窗口在本轮评价前指定，不是通过全年数据选出的最优窗口。mean_no_storage和mean_terminal6000复用相同mean预测，改变的是储能使用或规划终端条件。",
           "",
           "### 2.4 冷启动与后续衔接",
           "",
           "lag在已有历史不足7天时，负载回退为昨日同段值；mean只取已经存在的同星期日，不足4个时按实际样本数平均；完全没有此前同星期日时，负载用全部已有日期同段均值。光伏不足7天时用已有日期同段均值。首日完全无历史，不调用预测器，不编造预测值，沿用已声明的零普通购电计划与储能静置初始化。",
           "",
           "预测曲线组合成净负荷功率$\\widehat N_{k,t}=\\widehat L_{k,t}-\\widehat V_{k,t}$（kW），再乘$\\Delta t=1/6$小时得到净需求电量（kWh），作为日前MILP的输入。净负荷可为负，表示预测光伏富余；不提前截断，以保留储能充电机会。预测器解决“当天供需可能是多少”，MILP解决“据此买多少电、如何安排电池”。",
           "",
           "每个评价日均在0:00一次性给出整天预测；本轮不在白天更新这些日前预测，也不将当天实际值回填预测。历史来自附件2，未使用附件3预报或附件4电价。",
           "",
           '## 3. 模型建立','',
           '主Baseline与均值方案都用点预测MILP确定并冻结144段计划；预测末态=当日实际初态，抑制规划末端放空；该约定不是题目要求，实际末态不强制相等。另比较预测末态6000 kWh的终端敏感性。', '',
           r'实际费用$C=\sum_t p_tq_t+5\sum_t p_te_t$；供需平衡$q_t+V_t\Delta t+d_t+e_t=L_t\Delta t+c_t+w_t$，$\Delta t=1/6$小时。$q,e,c,d,w$均为母线侧kWh，$L,V$为kW，$p$为元/kWh，$E$为电池内部kWh。', '',
           r'储能状态$E_{t+1}=E_t+0.9c_t-d_t/0.9$，$1200\le E_t\le10800$；日前MILP以二进制变量保证充放电互斥，母线侧功率上限5000 kW。', '',
           '执行采用当期测量下富余充电、缺额放电、剩余应急的贪心控制；快速反馈和代表功率是离散化假设。它不预知后续真值，但也不保证跨时段应急费用最优。末态传给次日，从未每日重置。1月1日无历史：计划零、电池静置6000、当期光伏供负载并应急补缺。1月2日起滚动预测和运行；1月费用不计入题目2—12月评价。', '',
           '无储能对照使用同一均值预测，电池保持6000不参与；不同策略经1月运行后2月初态可能不同，下表列出初末电量，不能忽视储能价值差异。', '',
           '## 4. 模型求解与结果','',
           'HiGHS（scipy.optimize.milp）分支割框架；每日日前相对gap目标1e-9，时限120秒。实时反馈每段常数运算，总仿真线性于时间段数；MILP不作一般多项式复杂度承诺。', '',
           '| 方案 | 计划费（元） | 应急费（元） | 实际总费（元） | 应急量（kWh） |', '|---|---:|---:|---:|---:|']
    for key,run in runs.items():
        t=run[4]
        lines.append(f"| {NAMES[key]} | {t['planned_cost_yuan']:.2f} | {t['emergency_cost_yuan']:.2f} | {t['total_cost_yuan']:.2f} | {t['emergency_kWh']:.2f} |")
    lines+=['','| 方案 | 净负荷MAE（kW） | 应急天数 | 2月1日初态（kWh） | 12月31日末态（kWh） |','|---|---:|---:|---:|---:|']
    for key,run in runs.items():
        t=run[4]
        lines.append(f"| {NAMES[key]} | {t['net_mae_kW']:.2f} | {t['emergency_days']} | {t['evaluation_initial_kWh']:.3f} | {t['final_kWh']:.3f} |")
    lines+=['','以上为固定候选策略的因果回测表现，不代表选取全年表现最好方案后能对同一全年宣称无偏样本外优势。所有MILP的最优性只针对当天点预测模型，不是全年不确定性控制的最优性。', '',
            '### 指定日期：主Baseline','',
            '| 日期 | 计划电量（kWh） | 应急电量（kWh） | 计划费（元） | 应急费（元） | 实际总费（元） |','|---|---:|---:|---:|---:|---:|']
    primary=runs['lag']
    selected_dates=['2025-03-20','2025-06-21','2025-09-23','2025-12-21']
    for r in primary[1][primary[1].date.isin(selected_dates)].itertuples():
        lines.append(f'| {r.date} | {r.planned_kWh:.3f} | {r.emergency_kWh:.3f} | {r.planned_cost_yuan:.2f} | {r.emergency_cost_yuan:.2f} | {r.total_cost_yuan:.2f} |')
    selected=primary[0][primary[0].date.isin(selected_dates)]
    selected.to_csv(OUT/'selected_dates_dispatch.csv',index=False,encoding='utf-8-sig')
    primary[2][primary[2].date.isin(selected_dates)].to_csv(OUT/'selected_dates_emergency_events.csv',index=False,encoding='utf-8-sig')
    lines+=['','### 验证与限制','',
            '- 每段验证能量平衡、状态递推、非负、容量、功率和互斥，逐日验证状态连续；预测只接收已结束日期的历史数组。',
            '- 改动未来输入不改变预测和控制前缀；耗尽、充满、功率受限边界测试通过；参数化初态后第一问最优值回归通过。',
            '- 按日和逐段费用一致；连续应急事件电量总和与逐段总和一致；Excel重读验证计划、储能汇总及初末状态。',
            '- 残差和应急小时分布保存在各方案CSV；不把零停电作为优势，因为无限制的应急购电保证供需补足。',
            '- 终端目标与贪心控制尚需诊断，冷启动和2月初态影响未穷尽；10分钟代表功率无法评估更快的控制延迟或瞬时峰值。',
            '- 本轮没有概率风险优化，没有把未知预测误差当作零风险，不宣称本轮结果已最终收敛。','',
            '### 文件与复现','',
            '- code/05_q2_baseline.py：四组固定策略全年运行、验证和报告生成。',
            '- results/q2_baseline/comparison.csv、run_summary.json：总表、参数、来源、验证状态。',
            '- 各策略子目录：dispatch_all_2025.csv、daily_summary.csv、monthly_summary.csv、monthly_net_errors.csv、emergency_events.csv、emergency_by_hour.csv、validation.json。',
            '- results/q2_baseline/result2_baseline.xlsx：主Baseline计算副本，不覆盖附件5；每日费用账本含应急费。',
            '- figures/q2_baseline.png：各月费用及主Baseline储电状态。',
            '- 运行：conda run -n math_modeling python code/05_q2_baseline.py。','']
    (ROOT/'reports/Q2_Baseline.md').write_text('\n'.join(lines),encoding='utf-8')


def main():
    assert Path(sys.prefix).name=='math_modeling',sys.prefix
    np.random.seed(CONFIG['seed'])
    OUT.mkdir(parents=True,exist_ok=True)
    load,pv,prices,hashes=read_sources()
    tests=validate_core(load,pv,prices)
    print('Core and causality checks passed.',flush=True)
    runs={key:run_policy(key,load,pv,prices) for key in CONFIG['scenarios']}
    table=pd.DataFrame([r[4] for r in runs.values()])
    table.to_csv(OUT/'comparison.csv',index=False,encoding='utf-8-sig')
    export_excel(*runs['lag'][:3])
    make_report(runs)
    metadata=dict(run_time_utc=datetime.now(timezone.utc).isoformat(),config=CONFIG,
                  python=sys.version,executable=sys.executable,numpy=np.__version__,pandas=pd.__version__,scipy=scipy.__version__,
                  source_sha256=hashes,tests=tests,excel_readback_passed=True,comparisons=table.to_dict(orient='records'))
    (OUT/'run_summary.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8')
    print(table[['policy','total_cost_yuan','emergency_cost_yuan','net_mae_kW','evaluation_initial_kWh','final_kWh']].to_string(index=False),flush=True)


if __name__=='__main__':
    main()
