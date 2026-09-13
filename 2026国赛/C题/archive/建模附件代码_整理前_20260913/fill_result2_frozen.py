"""Fill the original result2 template from frozen N_free outputs; no model execution."""
from pathlib import Path
from copy import copy
from datetime import datetime, time, timedelta
import hashlib
import json
import shutil

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / '附件/附件5/result2.xlsx'
SOURCE = ROOT / 'results/q2_time_mapping/N_free'
OUT = ROOT / 'results/q2_frozen_delivery'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    paths = [SOURCE / 'template_plan.csv', SOURCE / 'natural_dispatch.csv']
    source_hashes = {str(p.relative_to(ROOT)): sha(p) for p in paths}
    plan = pd.read_csv(paths[0], low_memory=False)
    actual = pd.read_csv(paths[1], low_memory=False)
    plan = plan[plan.publication_date.between('2025-02-01', '2025-12-31')].copy()
    actual = actual[actual.date.between('2025-02-01', '2025-12-31')].copy()
    plan.sort_values(['publication_date', 'template_slot'], inplace=True)
    actual.sort_values(['date', 'natural_slot'], inplace=True)
    assert len(plan) == len(actual) == 334 * 144
    assert set(plan.strategy_id) == set(actual.strategy_id) == {'N_free'}
    assert set(plan.time_version) == set(actual.time_version) == {'start_time_v1'}
    dates = pd.date_range('2025-02-01', '2025-12-31')
    assert np.isfinite(plan[['planned_kWh', 'price_yuan_kWh']]).all().all()
    assert np.isfinite(actual[['charge_kWh', 'discharge_kWh', 'emergency_kWh',
                                     'state_start_kWh', 'state_end_kWh']]).all().all()
    backup = OUT / ('result2_before_fill_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f') + '.xlsx')
    shutil.copy2(TARGET, backup)
    wb = load_workbook(TARGET)
    assert wb.sheetnames == ['计划购电量', '充放电量', '紧急购电量']
    wp, ws, we = wb.worksheets
    original_headers = [c.value for c in wp[1]]
    original_dates = [wp.cell(r, 1).value for r in range(2, 336)]
    expected = {}

    def put(sheet, row, col, value):
        sheet.cell(row, col, value)
        expected[(sheet.title, row, col)] = value
        if isinstance(value, (float, np.floating)):
            sheet.cell(row, col).number_format = '0.000000'

    # Preserve the original 144 interval labels, including its next-midnight column.
    for row, date in enumerate(dates, 2):
        assert pd.Timestamp(wp.cell(row, 1).value) == date
        block = plan[plan.publication_date == str(date.date())]
        assert block.template_slot.tolist() == list(range(144))
        expected_times = pd.date_range(date + timedelta(minutes=10), periods=144, freq='10min')
        assert list(pd.to_datetime(block.interval_start)) == list(expected_times)
        for slot, q in enumerate(block.planned_kWh):
            label = original_headers[slot + 1].split('-')[0]
            hour, minute = map(int, label.split(':'))
            assert hour * 60 + minute == (10 * (slot + 1)) % 1440
            put(wp, row, slot + 2, float(q))
        put(wp, row, 146, float(block.planned_kWh.sum()))
        put(wp, row, 147, float((block.planned_kWh * block.price_yuan_kWh).sum()))

    # Expand the template's six-row example without changing the sheet structure.
    styles = [[copy(ws.cell(r, c)._style) for c in range(1, 7)] for r in range(2, 8)]
    heights = [ws.row_dimensions[r].height for r in range(2, 8)]
    ws.delete_rows(2, ws.max_row - 1)
    for day, date in enumerate(dates):
        block = actual[actual.date == str(date.date())]
        assert block.natural_slot.tolist() == list(range(144))
        assert list(pd.to_datetime(block.interval_start)) == list(pd.date_range(date, periods=144, freq='10min'))
        for period in range(6):
            row = 2 + day * 6 + period
            for col in range(1, 7):
                ws.cell(row, col)._style = copy(styles[period][col - 1])
            ws.row_dimensions[row].height = heights[period]
            if period == 0:
                put(ws, row, 1, date.to_pydatetime())
                ws.cell(row, 1).number_format = 'yyyy-mm-dd'
            put(ws, row, 2, f'{period * 4}:00-{(period + 1) * 4}:00')
            sub = block.iloc[period * 24:(period + 1) * 24]
            put(ws, row, 3, float(sub.charge_kWh.sum()))
            put(ws, row, 4, float(sub.discharge_kWh.sum()))
            if period < 2:
                put(ws, row, 5, time(0) if period == 0 else '24:00')
                put(ws, row, 6, float(block.state_start_kWh.iloc[0] if period == 0 else block.state_end_kWh.iloc[-1]))

    emergency_styles = [copy(we.cell(2, c)._style) for c in range(1, 4)]
    we.delete_rows(2, we.max_row - 1)
    emergency_rows = []
    row = 2
    for date in dates:
        values = actual.loc[actual.date == str(date.date()), 'emergency_kWh'].to_numpy()
        assert (values >= 0).all()
        starts = np.flatnonzero((values > 0) & np.r_[True, values[:-1] == 0])
        ends = np.flatnonzero((values > 0) & np.r_[values[1:] == 0, True]) + 1
        events = list(zip(starts, ends))
        if not events:
            events = [(0, 0)]
        for start, end in events:
            for col in range(1, 4):
                we.cell(row, col)._style = copy(emergency_styles[col - 1])
            label = f'{start // 6}:{start % 6 * 10:02d}-{end // 6}:{end % 6 * 10:02d}' if end else '无紧急购电'
            amount = float(values[start:end].sum())
            put(we, row, 1, date.to_pydatetime())
            we.cell(row, 1).number_format = 'yyyy-mm-dd'
            put(we, row, 2, label)
            put(we, row, 3, amount)
            emergency_rows.append(dict(date=str(date.date()), interval=label, emergency_kWh=amount))
            row += 1

    wp.cell(1, 147).comment = Comment('本列仅为本行144段计划购电费用，包含次日00:00—00:10，不含5倍价紧急购电费用。', '模型说明')
    wp.cell(1, 1).comment = Comment('冻结主模型：负载15列/光伏12列LightGBM，W28/q80，自由末态MILP，start_time_v1；来源29号N_free。按用户决定先填表，未完成独立审计。', '模型说明')
    ws.cell(1, 3).comment = Comment('自然日00:00—24:00实际执行的母线侧电量，单位kWh；不是日前名义充放电。充/放效率各0.9。', '模型说明')
    ws.cell(1, 6).comment = Comment('电池内部实际储电量(kWh)，每个自然日分别列0:00与24:00；不强制相等。', '模型说明')
    we.cell(1, 2).comment = Comment('按自然日列连续正应急电量区间，跨午夜分日列示；无应急日填0。费用按交付区间正常电价的5倍计算。', '模型说明')
    wp.freeze_panes = 'B2'
    ws.freeze_panes = we.freeze_panes = 'A2'
    temporary = OUT / 'result2_filled.xlsx'
    wb.save(temporary)
    check = load_workbook(temporary, data_only=True)
    max_error = 0.0
    for (name, r, c), value in expected.items():
        got = check[name].cell(r, c).value
        if isinstance(value, (float, np.floating)):
            max_error = max(max_error, abs(got - value))
        else:
            assert got == value, (name, r, c, got, value)
    assert max_error < 1e-6
    assert [c.value for c in check.worksheets[0][1]] == original_headers
    assert [check.worksheets[0].cell(r, 1).value for r in range(2, 336)] == original_dates
    assert check['充放电量'].max_row == 2005
    assert abs(sum(x['emergency_kWh'] for x in emergency_rows) - actual.emergency_kWh.sum()) < 1e-6
    assert source_hashes == {str(p.relative_to(ROOT)): sha(p) for p in paths}
    check.close()
    shutil.copy2(temporary, TARGET)
    assert sha(TARGET) == sha(temporary)
    summary = dict(model='15/12 LightGBM + W28/q80 + free terminal MILP',
                   independent_audit='deferred by user; delivery readback only',
                   workbook=str(TARGET), backup=str(backup), source_sha256=source_hashes,
                   output_sha256=sha(TARGET), script_sha256=sha(Path(__file__)),
                   days=334, planned_cells=48096, storage_periods=2004,
                   emergency_table_rows=len(emergency_rows), max_readback_error=max_error,
                   template_planned_kWh=float(plan.planned_kWh.sum()),
                   template_planned_cost_yuan=float((plan.planned_kWh * plan.price_yuan_kWh).sum()),
                   natural_emergency_kWh=float(actual.emergency_kWh.sum()),
                   natural_total_cost_yuan=float((actual.price_yuan_kWh * (actual.plan_kWh + 5 * actual.emergency_kWh)).sum()),
                   template_total_cost_yuan=float((plan.planned_cost_yuan + plan.emergency_cost_yuan).sum()),
                   all_readback_checks_passed=True)
    (OUT / 'fill_validation.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    pd.DataFrame(emergency_rows).to_csv(OUT / 'emergency_intervals.csv', index=False, encoding='utf-8-sig')
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
