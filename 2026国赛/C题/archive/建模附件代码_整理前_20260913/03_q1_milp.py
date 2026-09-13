"""Rerun Q1 MILP and export its own primal solution, binary modes and certificate."""
from pathlib import Path
import importlib.util
import hashlib
import json
import sys
from datetime import datetime, timezone

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.comments import Comment
import scipy

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'results/q1_milp'
spec = importlib.util.spec_from_file_location('q1_core', ROOT/'code/02_q1_baseline.py')
core = importlib.util.module_from_spec(spec)
spec.loader.exec_module(core)


def main():
    assert Path(sys.prefix).name == 'math_modeling', sys.prefix
    OUT.mkdir(parents=True, exist_ok=True)
    source = ROOT/'results/audit/q1_normalized.csv'
    audit = json.loads((ROOT/'results/audit/audit_summary.json').read_text(encoding='utf-8'))
    raw = ROOT/'附件/附件1.xlsx'
    assert hashlib.sha256(raw.read_bytes()).hexdigest() == audit['source_sha256']['附件\\附件1.xlsx']
    data = pd.read_csv(source)
    assert len(data) == 144
    assert np.array_equal(data.start_minute.to_numpy(), np.arange(0,1440,10))
    # Verify normalized numeric values directly against the source workbook.
    wb = openpyxl.load_workbook(raw, read_only=True, data_only=True)
    raw_values = np.array([r[1:] for r in list(wb.worksheets[0].values)[1:]], dtype=float)
    wb.close()
    assert np.allclose(data[['price_yuan_kWh','load_kW','pv_kW']],raw_values,rtol=0,atol=1e-10)
    load, pv, price = [data[k].to_numpy() for k in ['load_kW','pv_kW','price_yuan_kWh']]
    (q,c,d,w,state), summary = core.solve(load,pv,price,integer=True)
    lp = pd.read_csv(ROOT/'results/q1_baseline/dispatch_10min.csv')
    assert np.array_equal(lp.start_minute.to_numpy(),data.start_minute.to_numpy())
    lp_fee = float(price@lp.grid_kWh.to_numpy())
    no_fee = float(price@np.maximum(load-pv,0)/6)
    comparisons = dict(lp_cost_yuan=lp_fee, milp_minus_lp_cost_yuan=summary['cost_yuan']-lp_fee,
                       no_storage_cost_yuan=no_fee, savings_pct=100*(1-summary['cost_yuan']/no_fee))
    assert abs(comparisons['milp_minus_lp_cost_yuan']) < 1e-5
    assert abs(summary['objective_bound_gap_yuan']) < 1e-5
    data.insert(0,'interval',[f'{core.label(t)}-{core.label(t+10)}' for t in range(0,1440,10)])
    for key, value in dict(grid_kWh=q,charge_bus_kWh=c,discharge_bus_kWh=d,unused_kWh=w,
                           state_start_kWh=state[:-1],state_end_kWh=state[1:],cost_yuan=price*q,
                           binary_mode=np.rint(summary['binary_mode_raw']).astype(int)).items():
        data[key] = value
    data['operation'] = np.where(c>1e-6,'charge',np.where(d>1e-6,'discharge','idle'))
    for key in ['grid_kWh','charge_bus_kWh','discharge_bus_kWh','state_end_kWh']:
        comparisons[f'max_abs_difference_{key}'] = float(np.max(np.abs(data[key]-lp[key])))
    data.to_csv(OUT/'dispatch_10min.csv',index=False,encoding='utf-8-sig')
    diff = data.loc[np.abs(data.grid_kWh-lp.grid_kWh)>1e-6,['interval','price_yuan_kWh','grid_kWh']].copy()
    diff['lp_grid_kWh'] = lp.loc[diff.index,'grid_kWh']
    diff['milp_minus_lp_kWh'] = diff.grid_kWh-diff.lp_grid_kWh
    diff.to_csv(OUT/'lp_milp_differences.csv',index=False,encoding='utf-8-sig')
    groups = pd.DataFrame([dict(interval=f'{core.label(t*10)}-{core.label((t+24)*10)}',
                                 charge_kWh=float(c[t:t+24].sum()),discharge_kWh=float(d[t:t+24].sum()))
                          for t in range(0,144,24)])
    selected = data[data.start_minute.isin([600,720,840,960,1080,1200])][['interval','grid_kWh']]
    groups.to_csv(OUT/'charge_discharge_4hour.csv',index=False,encoding='utf-8-sig')
    selected.to_csv(OUT/'selected_intervals.csv',index=False,encoding='utf-8-sig')
    template = ROOT/'results/template_backups/result1_original.xlsx'
    if not template.exists():
        template = ROOT/'附件/附件5/result1.xlsx'
    wb = openpyxl.load_workbook(template)
    ws = wb['计划购电量']
    mapping = wb.create_sheet('时间映射')
    mapping.append(['原模板时段','计算时段','样本时刻_分钟'])
    for i,r in data.iterrows():
        mapping.append([ws.cell(i+2,1).value,r.interval,int(r.end_minute)])
        ws.cell(i+2,1,r.interval)
        ws.cell(i+2,2,float(r.grid_kWh)).number_format = '0.000000'
    ws['A1'].comment = Comment('右端点代表功率假设；原模板时段偏移已在计算副本修正，详见时间映射。','Modeling')
    ws = wb['充放电量']
    for i,r in groups.iterrows():
        ws.cell(i+2,2,float(r.charge_kWh))
        ws.cell(i+2,3,float(r.discharge_kWh))
    ws['E2'],ws['E3'] = float(state[0]),float(state[-1])
    detail = wb.create_sheet('MILP完整策略')
    detail.append(data.columns.tolist())
    for row in data.itertuples(index=False,name=None):
        detail.append(row)
    info = wb.create_sheet('求解与口径')
    for row in [
        ['状态','MILP重跑计算副本；时间和效率口径仍为显式假设'],
        ['求解器',summary['solver']],['最优购电费_元',summary['cost_yuan']],
        ['全天购电量_kWh',summary['grid_kWh']],['最优值下界_元',summary['mip_dual_bound_yuan']],
        ['MIP相对gap',summary['mip_gap']],['充放电效率','单向各0.9，往返0.81'],
        ['电量口径','充放电为母线侧kWh，储电量为电池内部kWh'],
        ['时间口径','样本为前10分钟区间代表功率；模板时段已显式修正'],
        ['binary_mode含义','1允许充电，0允许放电；闲置时可为任一值'],
    ]:
        info.append(row)
    target = OUT/'result1.xlsx'
    wb.save(target)
    wb.close()
    readback = openpyxl.load_workbook(target,data_only=True)
    assert np.allclose([r[1] for r in list(readback['计划购电量'].values)[1:]],q,rtol=0,atol=1e-9)
    assert np.allclose([[r[1],r[2]] for r in list(readback['充放电量'].values)[1:]],
                       groups[['charge_kWh','discharge_kWh']],rtol=0,atol=1e-9)
    assert readback['充放电量']['E2'].value == state[0]
    assert readback['充放电量']['E3'].value == state[-1]
    assert readback['计划购电量']['A2'].value == '00:00-00:10'
    assert readback['计划购电量']['A145'].value == '23:50-24:00'
    readback.close()
    metadata = dict(run_time_utc=datetime.now(timezone.utc).isoformat(),environment=sys.executable,
                    scipy_version=scipy.__version__,config=core.CONFIG,
                    normalized_input_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                    summary=summary,comparisons=comparisons,excel_readback_passed=True)
    (OUT/'validation_summary.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8')
    report = [
        '# 问题1 MILP重跑结果','',
        f"最优购电费 **{summary['cost_yuan']:.6f}元**，全天购电量 **{summary['grid_kWh']:.6f} kWh**。求解器返回最优，MIP gap为{summary['mip_gap']:.3e}；与既有LP费用差{comparisons['milp_minus_lp_cost_yuan']:.3e}元。",'',
        '## 问题分析','',
        '目标是在每段满足负载、日初日末储电量相同的条件下最小化全天购电费。本次保留统一物理参数，显式加入充放电互斥的二进制变量，保存MILP自己的完整解。', '',
        '## 数据预处理','',
        '输入附件1的144个时点。标准化数值逐项与原Excel核对，并核对源文件SHA256。采样点作为前10分钟区间代表值，功率乘以1/6小时转为电量，不平滑、不插补。原模板时间偏移在输出副本中修正，另表保存原标签。','',
        '## 模型建立','',
        r'$t=0,\ldots,143$，$\Delta t=1/6$小时。$L_t,V_t$为负载及光伏功率（kW），$p_t$为电价（元/kWh）；$q_t,c_t,d_t,w_t$分别为购电、母线侧充电、母线侧放电、未使用电量（kWh），$E_t$为时段起点电池内部储电量（kWh）。', '',
        r'$$\min\sum_{t=0}^{143}p_tq_t$$','',
        r'$$q_t+V_t\Delta t+d_t=L_t\Delta t+c_t+w_t,\quad E_{t+1}=E_t+0.9c_t-d_t/0.9,$$','',
        r'$$0\le c_t\le Mz_t,\quad0\le d_t\le M(1-z_t),\quad z_t\in\{0,1\},\quad M=5000\Delta t=833.333\ldots\ \mathrm{kWh},$$','',
        r'$$q_t,w_t\ge0,\quad1200\le E_t\le10800,\quad E_0=E_{144}=6000.$$','',
        '单向效率各90%，功率限制按母线侧解释；没有售电收入及额外损耗费用。z=1允许充电、z=0允许放电，闲置时z可取任一值，应结合实际充放电量判定动作。', '',
        '## 模型求解','',
        '结构性补充（2026-09-10）：LP与MILP最优费用相同可由“免费且无上限的弃余允许消去同段充放电”严格解释，并可重构为289个连续变量的精确LP。完整证明、条件及反例见[问题1_LP等价性与深入分析.md](问题1_LP等价性与深入分析.md)。本报告原MILP数值与已填写Excel保持锁定；下文所称独立LP对照删除了模式约束，比只放宽二进制变量的标准LP松弛更宽，但其下界有效，且在上述条件下仍精确。','',
        '721个连续变量、144个二进制变量；288个等式、288个互斥不等式及变量界。使用SciPy的HiGHS MILP求解，无人工初始解；相对gap目标1e-9，时限120秒。MILP最坏情形计算复杂度较高，本实例以求解器上下界证实最优。', '',
        "### 求解器与算法",
        "",
        "数学模型为混合整数线性规划（MILP）；求解器为HiGHS，通过Python的`scipy.optimize.milp`接口调用，运行环境为`math_modeling`，本次记录的SciPy版本为1.18.0。MILP是模型类型，HiGHS是求解器，二者不等同于具体求解算法。",
        "",
        "HiGHS的混合整数优化采用分支割（Branch-and-Cut）框架，结合预处理、线性规划松弛、割平面、整数可行解启发式和分支定界。其基本原理如下：",
        "",
        "1. 放宽二进制变量约束，令$0\\le z_t\\le1$，求解LP松弛，为最小化问题提供目标值下界。",
        "2. 通过预处理、有效不等式和分支搜索收紧松弛，寻找满足整数约束的可行解；当前最好整数可行解的费用提供上界。",
        "3. 根据上下界剪枝，迭代缩小最优性间隙；达到规定的相对间隙容差时判定满足精度要求，也可能因时间上限等条件终止。是否得到最优结果须结合返回状态和间隙判断。",
        "",
        "本题设定`mip_rel_gap=1e-9`、`time_limit=120`秒，未提供人工初始解。实际返回`HiGHS Status 7: Optimal`，MIP相对间隙为0，搜索节点数为1，表明本实例无需展开大规模分支搜索。以上描述的是求解器的一般算法框架；当前保存的运行记录不含逐次割平面和启发式日志，因此不将其描述为本次实际执行步骤的逐项记录。",
        "",
        "另外，通过`scipy.optimize.linprog(method='highs')`求解独立的LP松弛作为对照。LP最优费用与MILP最优费用均为35126.948589元，在数值容差内整数间隙为零，结合整数与物理约束校验验证MILP解的最优性。",
        "",
        "论文可表述为：本文通过SciPy接口调用HiGHS求解器，采用分支割框架求解含充放电互斥约束的混合整数线性规划模型，并以线性规划松弛提供最优值下界进行交叉验证。求解结果的MIP相对间隙为零，说明所得调度在给定模型与数值容差下达到全局最优。",
        "",
        f"求解用时{summary['elapsed_seconds']:.6f}秒（不含读取和导出），节点数{summary['mip_node_count']}；目标下界{summary['mip_dual_bound_yuan']:.6f}元。", '',
        '| 指标 | MILP结果 |','|---|---:|',
        f"| 全天购电费（元） | {summary['cost_yuan']:.6f} |",
        f"| 全天购电量（kWh） | {summary['grid_kWh']:.6f} |",
        f"| 充电量（kWh） | {summary['charge_kWh']:.6f} |",
        f"| 放电量（kWh） | {summary['discharge_kWh']:.6f} |",
        f"| 未使用电量（kWh） | {summary['unused_kWh']:.6f} |",
        f"| 初始/最终储电量（kWh） | {state[0]:.6f} / {state[-1]:.6f} |", '',
        '| 指定时段 | 购电量（kWh） |','|---|---:|',
    ]
    report += [f'| {r.interval} | {r.grid_kWh:.6f} |' for r in selected.itertuples()]
    report += ['', '| 时间段 | 充电量（kWh） | 放电量（kWh） |','|---|---:|---:|']
    report += [f'| {r.interval} | {r.charge_kWh:.6f} | {r.discharge_kWh:.6f} |' for r in groups.itertuples()]
    report += ['', '### 验证与比较','']
    report += [f'- {k}：{v:.6e}。' for k,v in summary['checks'].items()]
    report += ['',f"相对无储能对照节省{comparisons['savings_pct']:.4f}%。LP与MILP逐段购电最大差{comparisons['max_abs_difference_grid_kWh']:.6e} kWh，末态轨迹最大差{comparisons['max_abs_difference_state_end_kWh']:.6e} kWh。即使费用相同，也不据此声称策略唯一。",'',
               'Excel购电、六段充放电汇总、初末状态和时间标签均经重读核验。此次新增整数变量不产生额外节费：已有LP解已满足互斥，MILP明确表达了这项物理条件。结果仍以时间、效率及初始状态假设为条件。', '',
               '### 文件与复现','',
               '- `results/q1_milp/result1.xlsx`：MILP策略、二进制状态、映射和求解说明。',
               '- `results/q1_milp/dispatch_10min.csv`：144段完整解；另有4小时汇总与指定时段CSV。',
               '- `results/q1_milp/validation_summary.json`：环境、参数、最优上下界、整数及物理检查、LP比较。',
               '- 赛题目录执行 `conda run -n math_modeling python code/03_q1_milp.py`。','']
    report += ['### LP与MILP时段差异','',
               '| 时段 | 电价（元/kWh） | LP购电（kWh） | MILP购电（kWh） | 差值（kWh） |',
               '|---|---:|---:|---:|---:|']
    report += [f'| {r.interval} | {r.price_yuan_kWh:.6f} | {r.lp_grid_kWh:.6f} | {r.grid_kWh:.6f} | {r.milp_minus_lp_kWh:.6f} |' for r in diff.itertuples()]
    report += ['', '本次差异是等价时段之间充电安排的变化：23:10—23:20和23:30—23:40电价同为0.424元/kWh，333.333333 kWh充电由前者移至后者，费用不变。两者均通过约束校验，说明在数值容差内存在不同的最优调度。指定的六个10分钟结果及六段4小时汇总均相同。差异数据另存`results/q1_milp/lp_milp_differences.csv`。','']
    (ROOT/'reports/问题一/问题1_MILP重跑结果.md').write_text('\n'.join(report),encoding='utf-8')
    print(json.dumps(dict(summary={k:v for k,v in summary.items() if k!='binary_mode_raw'},
                          comparisons=comparisons),ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
